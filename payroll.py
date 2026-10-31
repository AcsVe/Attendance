"""
الرواتب (بديل MenaPAY) + الربط مع MenaPAY / MenaITech — من بيانات الحضور نفسها.

• رواتب الموظفين: الرقم الوظيفي (رمز الموظف في MenaPAY)، الراتب الأساسي، العلاوات الثابتة، الضمان الاجتماعي،
  الضريبة الشهرية، البنك وIBAN وطريقة الدفع.
• قواعد الاحتساب (كل قاعدة تُفعَّل أو تُعطَّل، وكلها معطّلة افتراضياً عدا الإجازة بلا راتب):
  أساس أجر اليوم (30 يوماً أو أيام العمل الفعلية)، التأخير (بالدقيقة أو بالشرائح لكل يوم أو كل N مرات = يوم)
  مع سماح شهري، المغادرة المبكرة، الغياب بلا عذر، الإجازة بلا راتب، نسيان الانصراف، العمل الإضافي (أيام العمل
  والعطلة)، حدّ أقصى للخصومات، الضمان الاجتماعي (حصة الموظف والمؤسسة).
• مسيرة الشهر: مسودة تُحسب من «الكشف الشهري» ← مراجعة وتعديلات يدوية (مكافأة/خصم بسبب) ← اعتماد (يُقفل الشهر)
  ← فتحه يحتاج سبباً يُسجَّل. لا يُعتمد خصم تلقائياً أبداً.
• المخرجات: كشف الرواتب، ملف البنك، قسائم الرواتب للطباعة، وقسيمة الموظف في تطبيقه بعد الاعتماد (خيار).
• الربط: تصدير متغيرات الحضور (أيام/ساعات/دقائق/مبالغ) بصيغة استيراد MenaPAY القابلة للتخصيص (صف لكل موظف
  أو صف لكل بند برمز البند)، واستيراد بيانات الموظفين والرواتب من ملف MenaITech/MenaPAY.
• الصلاحية: المدير، ومن يمنحه المدير «الرواتب» من هذه الصفحة. كل تغيير في سجل التدقيق.
"""
import csv
import io
import json
import re
from datetime import date as _date, datetime, timedelta

from flask import Blueprint, jsonify, request, send_file, session, Response

import db

bp = Blueprint('payroll', __name__)

DEFAULT_RULES = {
    'currency': 'JOD', 'decimals': 3,
    'day_basis': 'thirty',            # thirty | workdays
    'hours_per_day': 8,
    'late': {'on': False, 'mode': 'minutes', 'grace_minutes': 0, 'factor': 1,
             'tiers': [{'from': 1, 'to': 15, 'day': 0}, {'from': 16, 'to': 30, 'day': 0.25},
                       {'from': 31, 'to': 60, 'day': 0.5}, {'from': 61, 'to': 9999, 'day': 1}],
             'every_n': 3, 'every_day': 1},
    'early': {'on': False, 'grace_minutes': 0, 'factor': 1},
    'absence': {'on': False, 'factor': 1},
    'unpaid_leave': {'on': True, 'factor': 1},
    'no_out': {'on': False, 'free': 2, 'day': 0.25},
    'overtime': {'on': False, 'factor': 1.25, 'weekend_factor': 1.5},
    'cap_percent': 0,
    'ss': {'on': False, 'employee_rate': 7.5, 'employer_rate': 14.25},
    'tax': {'on': False},
    'payslip_app': False,
}

MENAPAY_ITEMS = [
    # key, Arabic label, default code, unit
    ('present', 'أيام الحضور', 'WD', 'أيام'),
    ('absent', 'أيام الغياب', 'ABS', 'أيام'),
    ('unpaid_leave', 'أيام إجازة بلا راتب', 'UPL', 'أيام'),
    ('vacation_days', 'أيام الإجازات', 'VAC', 'أيام'),
    ('late_minutes', 'دقائق التأخير', 'LATE', 'دقائق'),
    ('late_hours', 'ساعات التأخير', 'LATEH', 'ساعات'),
    ('early_minutes', 'دقائق المغادرة المبكرة', 'EARLY', 'دقائق'),
    ('no_out', 'مرات بدون سجل انصراف', 'NOOUT', 'مرات'),
    ('ot_hours', 'ساعات العمل الإضافي', 'OT', 'ساعات'),
    ('weekend_work', 'أيام عمل في العطلة', 'OTW', 'أيام'),
    ('deduct_days', 'أيام الخصم (حسب القواعد)', 'DED', 'أيام'),
    ('deduct_amount', 'مبلغ خصومات الحضور', 'DEDA', 'مبلغ'),
    ('ot_amount', 'مبلغ العمل الإضافي', 'OTA', 'مبلغ'),
    ('net', 'صافي الراتب', 'NET', 'مبلغ'),
]
DEFAULT_MENAPAY = {'layout': 'long', 'code_field': 'employee_no', 'period_format': '%Y%m',
                   'items': [{'key': k, 'code': c, 'on': k in ('absent', 'unpaid_leave', 'late_minutes', 'early_minutes', 'ot_hours')}
                             for k, _, c, _ in MENAPAY_ITEMS]}
ITEM_LABEL = {k: l for k, l, _, _ in MENAPAY_ITEMS}
STATUS_LABEL = {'draft': 'مسودة', 'approved': 'معتمدة'}


# ---------------------------------------------------------------
#   الجداول
# ---------------------------------------------------------------
def init_tables(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS payroll_profiles (
        employee_id INTEGER PRIMARY KEY,
        employee_no TEXT,
        basic REAL NOT NULL DEFAULT 0,
        allowances TEXT NOT NULL DEFAULT '[]',
        ss_subject INTEGER NOT NULL DEFAULT 1,
        ss_salary REAL,
        tax_monthly REAL NOT NULL DEFAULT 0,
        payment TEXT NOT NULL DEFAULT 'bank',
        bank_name TEXT,
        iban TEXT,
        notes TEXT,
        updated_at TEXT
    )''')
    conn.execute('''CREATE TABLE IF NOT EXISTS payroll_runs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        month TEXT UNIQUE NOT NULL,
        status TEXT NOT NULL DEFAULT 'draft',
        rules TEXT,
        created_at TEXT, created_by TEXT,
        approved_at TEXT, approved_by TEXT,
        reopen_note TEXT
    )''')
    conn.execute('''CREATE TABLE IF NOT EXISTS payroll_lines (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        run_id INTEGER NOT NULL,
        employee_id INTEGER NOT NULL,
        name TEXT, dept TEXT, employee_no TEXT,
        data TEXT NOT NULL,
        gross REAL NOT NULL DEFAULT 0, deductions REAL NOT NULL DEFAULT 0, net REAL NOT NULL DEFAULT 0,
        UNIQUE (run_id, employee_id)
    )''')
    conn.execute('''CREATE TABLE IF NOT EXISTS payroll_adjustments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        run_id INTEGER NOT NULL,
        employee_id INTEGER NOT NULL,
        kind TEXT NOT NULL,
        amount REAL NOT NULL,
        reason TEXT,
        created_by TEXT, created_at TEXT
    )''')
    conn.execute("CREATE TABLE IF NOT EXISTS payroll_access (user_id INTEGER PRIMARY KEY)")


# ---------------------------------------------------------------
#   الصلاحية والإعدادات
# ---------------------------------------------------------------
def can_payroll(conn=None):
    if not session.get('user_id'):
        return False
    if session.get('role') == 'admin':
        return True
    own = conn is None
    conn = conn or db.get_db()
    try:
        return bool(conn.execute("SELECT 1 FROM payroll_access WHERE user_id = ?", (session['user_id'],)).fetchone())
    finally:
        if own:
            conn.close()


def _deny():
    if not session.get('user_id'):
        return jsonify({'ok': False, 'error': 'يرجى تسجيل الدخول أولاً'}), 401
    return jsonify({'ok': False, 'error': 'ليست لديك صلاحية الرواتب'}), 403


def _merge(base, over):
    out = json.loads(json.dumps(base))
    for k, v in (over or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _merge(out[k], v)
        elif k in out:
            out[k] = v
    return out


def rules():
    try:
        return _merge(DEFAULT_RULES, json.loads(db.get_setting('payroll_rules_json', '') or '{}'))
    except ValueError:
        return json.loads(json.dumps(DEFAULT_RULES))


def menapay_cfg():
    try:
        saved = json.loads(db.get_setting('payroll_menapay_json', '') or '{}')
    except ValueError:
        saved = {}
    cfg = _merge(DEFAULT_MENAPAY, {k: v for k, v in saved.items() if k != 'items'})
    by_key = {i['key']: i for i in saved.get('items') or [] if isinstance(i, dict)}
    cfg['items'] = [{**it, **{k: by_key[it['key']][k] for k in ('code', 'on') if it['key'] in by_key and k in by_key[it['key']]}}
                    for it in cfg['items']]
    return cfg


def _audit(conn, action, value, name=None):
    try:
        import app as app_module
        app_module.log_audit(conn, action, 'payroll', employee_name=name, new_value=str(value)[:500])
    except Exception:   # noqa: BLE001
        pass


def _num(v, default=0.0):
    try:
        return float(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return default


def _round(v, r):
    return round(v + 0.0, int(r.get('decimals', 3)))


# ---------------------------------------------------------------
#   الاحتساب
# ---------------------------------------------------------------
def _month_bounds(month):
    y, m = int(month[:4]), int(month[5:7])
    first = _date(y, m, 1)
    last = (_date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1))
    return first, last


def _attendance(conn, month):
    """مجاميع حضور كل موظف للشهر (من الكشف الشهري) + أيام التأخير بدقائقها (للشرائح) + الإجازات بلا راتب"""
    import app as app_module
    _, _, summary, _ = app_module._monthly_summary({'month': month})
    by_name = {x['name']: dict(x) for x in summary}
    unpaid = {r['name'] for r in conn.execute("SELECT name FROM vacation_types WHERE paid = 0")} \
        if conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'vacation_types'").fetchone() else set()
    for r in conn.execute("SELECT name, late_minutes, day_note FROM report_cache WHERE date LIKE ?", (month + '%',)):
        x = by_name.get(r['name'])
        if not x:
            continue
        if int(r['late_minutes'] or 0) > 0:
            x.setdefault('late_list', []).append(int(r['late_minutes']))
        if (r['day_note'] or '') in unpaid:
            x['unpaid_leave'] = x.get('unpaid_leave', 0) + 1
    return by_name


def _late_days(x, rl):
    """أيام الخصم بسبب التأخير حسب الطريقة"""
    lt = rl['late']
    if lt['mode'] == 'tiers':
        days = 0.0
        for m in x.get('late_list', []):
            for t in lt.get('tiers') or []:
                if _num(t.get('from')) <= m <= _num(t.get('to'), 10 ** 9):
                    days += _num(t.get('day'))
                    break
        return days, None
    if lt['mode'] == 'count':
        n = max(1, int(_num(lt.get('every_n'), 3)))
        return (len(x.get('late_list', [])) // n) * _num(lt.get('every_day'), 1), None
    minutes = max(0, int(x.get('late_minutes', 0)) - int(_num(lt.get('grace_minutes'))))
    return None, minutes


def compute_line(profile, x, rl, workdays_count):
    """سطر موظف واحد: الاستحقاقات، الخصومات (مع السبب)، الصافي"""
    basic = _num(profile.get('basic'))
    allowances = [a for a in (profile.get('allowances') or []) if _num(a.get('amount'))]
    fixed = basic + sum(_num(a['amount']) for a in allowances)
    divisor = workdays_count if rl['day_basis'] == 'workdays' and workdays_count else 30
    day_wage = basic / divisor if divisor else 0
    minute_wage = day_wage / (max(1, _num(rl['hours_per_day'], 8)) * 60)
    earn = [{'key': 'basic', 'label': 'الراتب الأساسي', 'amount': basic}]
    earn += [{'key': 'allowance', 'label': a.get('name') or 'علاوة', 'amount': _num(a['amount'])} for a in allowances]
    ded = []
    x = x or {}

    def add_ded(key, label, qty, unit, amount):
        if amount > 0:
            ded.append({'key': key, 'label': label, 'qty': round(qty, 2), 'unit': unit, 'amount': amount, 'attendance': True})

    deduct_days = 0.0
    if rl['late']['on']:
        days, minutes = _late_days(x, rl)
        if days is not None:
            deduct_days += days
            add_ded('late', 'التأخير', days, 'يوم', days * day_wage)
        else:
            add_ded('late', 'التأخير', minutes, 'دقيقة', minutes * minute_wage * _num(rl['late']['factor'], 1))
    if rl['early']['on']:
        m = max(0, int(x.get('early_minutes', 0)) - int(_num(rl['early'].get('grace_minutes'))))
        add_ded('early', 'المغادرة المبكرة', m, 'دقيقة', m * minute_wage * _num(rl['early']['factor'], 1))
    if rl['absence']['on']:
        d = float(x.get('absent', 0))
        deduct_days += d
        add_ded('absent', 'الغياب بلا عذر', d, 'يوم', d * day_wage * _num(rl['absence']['factor'], 1))
    if rl['unpaid_leave']['on']:
        d = float(x.get('unpaid_leave', 0))
        deduct_days += d
        add_ded('unpaid_leave', 'إجازة بلا راتب', d, 'يوم', d * day_wage * _num(rl['unpaid_leave']['factor'], 1))
    if rl['no_out']['on']:
        n = max(0, int(x.get('no_out', 0)) - int(_num(rl['no_out'].get('free'))))
        d = n * _num(rl['no_out'].get('day'))
        deduct_days += d
        add_ded('no_out', 'نسيان الانصراف', n, 'مرة', d * day_wage)
    att_total = sum(d['amount'] for d in ded)
    cap = _num(rl.get('cap_percent'))
    if cap > 0 and att_total > fixed * cap / 100:
        factor = (fixed * cap / 100) / att_total if att_total else 0
        for d in ded:
            d['amount'] *= factor
        ded.append({'key': 'cap', 'label': f'سقف الخصومات {cap:g}% (خُفّضت)', 'qty': 0, 'unit': '', 'amount': 0, 'attendance': False})
    ot_amount = 0.0
    if rl['overtime']['on']:
        ot_h = int(x.get('ot_minutes', 0)) / 60
        ot_amount += ot_h * minute_wage * 60 * _num(rl['overtime']['factor'], 1.25)
        wk = int(x.get('weekend_work', 0))
        ot_amount += wk * day_wage * _num(rl['overtime']['weekend_factor'], 1.5)
        if ot_amount:
            earn.append({'key': 'overtime', 'label': 'العمل الإضافي', 'qty': round(ot_h, 2), 'weekend_days': wk, 'amount': ot_amount})
    employer_ss = 0.0
    if rl['ss']['on'] and profile.get('ss_subject', 1):
        insurable = _num(profile.get('ss_salary')) or fixed
        ded.append({'key': 'ss', 'label': 'الضمان الاجتماعي (حصة الموظف)', 'qty': _num(rl['ss']['employee_rate']), 'unit': '%',
                    'amount': insurable * _num(rl['ss']['employee_rate']) / 100, 'attendance': False})
        employer_ss = insurable * _num(rl['ss']['employer_rate']) / 100
    if rl['tax']['on'] and _num(profile.get('tax_monthly')):
        ded.append({'key': 'tax', 'label': 'ضريبة الدخل', 'qty': 0, 'unit': '', 'amount': _num(profile.get('tax_monthly')),
                    'attendance': False})
    for d in ded:
        d['amount'] = _round(d['amount'], rl)
    for e in earn:
        e['amount'] = _round(e['amount'], rl)
    return {'earnings': earn, 'deductions': ded, 'day_wage': _round(day_wage, rl), 'deduct_days': round(deduct_days, 2),
            'attendance': {k: x.get(k, 0) for k in ('present', 'absent', 'late_days', 'late_minutes', 'early_minutes', 'no_out',
                                                     'vacation_days', 'unpaid_leave', 'ot_minutes', 'weekend_work')},
            'employer_ss': _round(employer_ss, rl), 'ot_amount': _round(ot_amount, rl)}


def _totals(data, adjustments, rl):
    gross = sum(e['amount'] for e in data['earnings']) + sum(a['amount'] for a in adjustments if a['kind'] == 'add')
    deds = sum(d['amount'] for d in data['deductions']) + sum(a['amount'] for a in adjustments if a['kind'] == 'deduct')
    return _round(gross, rl), _round(deds, rl), _round(gross - deds, rl)


def _profiles(conn):
    out = {}
    for r in conn.execute("""SELECT e.id, e.name, e.dept, e.device_uid, e.active, p.* FROM employees e
                             LEFT JOIN payroll_profiles p ON p.employee_id = e.id"""):
        d = dict(r)
        d['id'] = r['id']
        try:
            d['allowances'] = json.loads(r['allowances'] or '[]')
        except ValueError:
            d['allowances'] = []
        out[r['id']] = d
    return out


def build_run(conn, month, user):
    """يحسب مسودة الشهر من جديد (يحفظ التعديلات اليدوية). يرفض إن كانت معتمدة."""
    run = conn.execute("SELECT * FROM payroll_runs WHERE month = ?", (month,)).fetchone()
    if run and run['status'] == 'approved':
        raise ValueError('مسيرة هذا الشهر معتمدة ومقفلة — افتحها أولاً لإعادة الاحتساب.')
    rl = rules()
    if not run:
        conn.execute("INSERT INTO payroll_runs (month, status, created_at, created_by) VALUES (?, 'draft', ?, ?)",
                     (month, _now(), user))
        run = conn.execute("SELECT * FROM payroll_runs WHERE month = ?", (month,)).fetchone()
    run_id = run['id']
    conn.execute("UPDATE payroll_runs SET rules = ? WHERE id = ?", (json.dumps(rl, ensure_ascii=False), run_id))
    first, last = _month_bounds(month)
    import vacations
    wd = len(vacations.work_days(conn, first, last))
    att = _attendance(conn, month)
    adj = {}
    for a in conn.execute("SELECT * FROM payroll_adjustments WHERE run_id = ?", (run_id,)):
        adj.setdefault(a['employee_id'], []).append(dict(a))
    conn.execute("DELETE FROM payroll_lines WHERE run_id = ?", (run_id,))
    n = 0
    for p in _profiles(conn).values():
        if not p['active'] or _num(p.get('basic')) <= 0:
            continue
        data = compute_line(p, att.get(p['name']), rl, wd)
        data['adjustments'] = [{'id': a['id'], 'kind': a['kind'], 'amount': a['amount'], 'reason': a['reason'] or ''}
                               for a in adj.get(p['id'], [])]
        data['workdays'] = wd
        gross, deds, net = _totals(data, data['adjustments'], rl)
        conn.execute("""INSERT INTO payroll_lines (run_id, employee_id, name, dept, employee_no, data, gross, deductions, net)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                     (run_id, p['id'], p['name'], p['dept'] or '', p.get('employee_no') or '',
                      json.dumps(data, ensure_ascii=False), gross, deds, net))
        n += 1
    _audit(conn, 'payroll_compute', f'{month}: {n} موظف')
    return run_id, n


def _now():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def _run_payload(conn, run):
    lines = []
    for r in conn.execute("SELECT * FROM payroll_lines WHERE run_id = ? ORDER BY dept, name", (run['id'],)):
        d = dict(r)
        d['data'] = json.loads(r['data'])
        lines.append(d)
    rl = json.loads(run['rules'] or '{}') or rules()
    tot = {k: _round(sum(x[k] for x in lines), rl) for k in ('gross', 'deductions', 'net')}
    tot['employer_ss'] = _round(sum(x['data'].get('employer_ss', 0) for x in lines), rl)
    return {'run': {**dict(run), 'status_label': STATUS_LABEL.get(run['status'], run['status']), 'rules': None},
            'lines': lines, 'totals': tot, 'currency': rl.get('currency', 'JOD')}


# ---------------------------------------------------------------
#   واجهة الإدارة
# ---------------------------------------------------------------
def _guard():
    return None if can_payroll() else _deny()


def _month_arg(v):
    v = (v or '')[:7]
    try:
        datetime.strptime(v, '%Y-%m')
        return v
    except ValueError:
        return None


@bp.route('/api/payroll/overview')
def api_overview():
    if not can_payroll():
        return _deny()
    conn = db.get_db()
    try:
        runs = [{**dict(r), 'status_label': STATUS_LABEL.get(r['status'], r['status']), 'rules': None,
                 'count': conn.execute("SELECT COUNT(*) FROM payroll_lines WHERE run_id = ?", (r['id'],)).fetchone()[0],
                 'net': conn.execute("SELECT COALESCE(SUM(net), 0) FROM payroll_lines WHERE run_id = ?", (r['id'],)).fetchone()[0]}
                for r in conn.execute("SELECT * FROM payroll_runs ORDER BY month DESC")]
        out = {'ok': True, 'rules': rules(), 'menapay': menapay_cfg(), 'runs': runs,
               'items': [{'key': k, 'label': l, 'unit': u} for k, l, _, u in MENAPAY_ITEMS],
               'is_admin': session.get('role') == 'admin'}
        if session.get('role') == 'admin':
            have = {r['user_id'] for r in conn.execute("SELECT user_id FROM payroll_access")}
            out['users'] = [{'id': r['id'], 'username': r['username'], 'full_name': r['full_name'] or '', 'role': r['role'],
                             'access': r['role'] == 'admin' or r['id'] in have}
                            for r in conn.execute("SELECT id, username, full_name, role FROM users ORDER BY username")]
        return jsonify(out)
    finally:
        conn.close()


@bp.route('/api/payroll/profiles')
def api_profiles():
    if not can_payroll():
        return _deny()
    conn = db.get_db()
    try:
        rows = [{'id': p['id'], 'name': p['name'], 'dept': p['dept'] or '', 'device_uid': p['device_uid'] or '',
                 'active': bool(p['active']), 'employee_no': p.get('employee_no') or '', 'basic': _num(p.get('basic')),
                 'allowances': p['allowances'], 'ss_subject': p.get('ss_subject', 1) if p.get('ss_subject') is not None else 1,
                 'ss_salary': p.get('ss_salary'), 'tax_monthly': _num(p.get('tax_monthly')), 'payment': p.get('payment') or 'bank',
                 'bank_name': p.get('bank_name') or '', 'iban': p.get('iban') or '', 'notes': p.get('notes') or ''}
                for p in sorted(_profiles(conn).values(), key=lambda x: (x['dept'] or '', x['name']))]
        return jsonify({'ok': True, 'rows': rows})
    finally:
        conn.close()


def _clean_profile(d):
    allowances = []
    for a in d.get('allowances') or []:
        name = str(a.get('name') or '').strip()[:60]
        amt = _num(a.get('amount'))
        if name and amt:
            allowances.append({'name': name, 'amount': amt})
    iban = re.sub(r'\s+', '', str(d.get('iban') or '')).upper()[:40]
    if iban and not re.fullmatch(r'[A-Z]{2}\d{2}[A-Z0-9]{8,30}', iban):
        raise ValueError('رقم IBAN غير صالح (مثال: JO94CBJO0010000000000131000302)')
    basic = _num(d.get('basic'))
    if basic < 0:
        raise ValueError('الراتب الأساسي غير صالح')
    return {'employee_no': str(d.get('employee_no') or '').strip()[:40], 'basic': basic,
            'allowances': json.dumps(allowances, ensure_ascii=False),
            'ss_subject': 1 if d.get('ss_subject', True) in (True, 1, '1', 'true', 'on') else 0,
            'ss_salary': _num(d.get('ss_salary')) or None, 'tax_monthly': max(0.0, _num(d.get('tax_monthly'))),
            'payment': 'cash' if d.get('payment') == 'cash' else 'bank',
            'bank_name': str(d.get('bank_name') or '').strip()[:80], 'iban': iban,
            'notes': str(d.get('notes') or '').strip()[:300]}


def _save_profile(conn, emp_id, row):
    row = {**row, 'updated_at': _now()}
    cols = ', '.join(row)
    conn.execute(f"INSERT INTO payroll_profiles (employee_id, {cols}) VALUES (?, {', '.join('?' * len(row))}) "
                 f"ON CONFLICT(employee_id) DO UPDATE SET {', '.join(f'{k} = excluded.{k}' for k in row)}",
                 [emp_id, *row.values()])


@bp.route('/api/payroll/profiles/<int:emp_id>', methods=['POST'])
def api_profile_save(emp_id):
    if not can_payroll():
        return _deny()
    d = request.get_json(silent=True) or {}
    try:
        row = _clean_profile(d)
    except ValueError as e:
        return jsonify({'ok': False, 'error': str(e)}), 400
    conn = db.get_db()
    try:
        emp = conn.execute("SELECT name FROM employees WHERE id = ?", (emp_id,)).fetchone()
        if not emp:
            return jsonify({'ok': False, 'error': 'الموظف غير موجود'}), 404
        old = conn.execute("SELECT basic FROM payroll_profiles WHERE employee_id = ?", (emp_id,)).fetchone()
        _save_profile(conn, emp_id, row)
        _audit(conn, 'payroll_profile', f"الأساسي {old['basic'] if old else 0} ← {row['basic']}", emp['name'])
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


@bp.route('/api/payroll/profiles/bulk', methods=['POST'])
def api_profiles_bulk():
    """تعديل جماعي: زيادة نسبة/مبلغ على الأساسي، أو ضبط الضمان/طريقة الدفع/البنك للمحددين"""
    if not can_payroll():
        return _deny()
    d = request.get_json(silent=True) or {}
    ids = [int(i) for i in d.get('ids') or [] if str(i).isdigit()]
    action = d.get('action')
    if not ids or action not in ('raise_percent', 'raise_amount', 'ss_on', 'ss_off', 'pay_bank', 'pay_cash', 'bank_name'):
        return jsonify({'ok': False, 'error': 'حدّد موظفاً واحداً على الأقل واختر الإجراء'}), 400
    val = d.get('value')
    conn = db.get_db()
    try:
        profiles = _profiles(conn)
        n = 0
        for i in ids:
            p = profiles.get(i)
            if not p:
                continue
            row = _clean_profile({**p, 'allowances': p['allowances']})
            if action == 'raise_percent':
                row['basic'] = round(row['basic'] * (1 + _num(val) / 100), 3)
            elif action == 'raise_amount':
                row['basic'] = round(row['basic'] + _num(val), 3)
            elif action in ('ss_on', 'ss_off'):
                row['ss_subject'] = 1 if action == 'ss_on' else 0
            elif action in ('pay_bank', 'pay_cash'):
                row['payment'] = 'bank' if action == 'pay_bank' else 'cash'
            elif action == 'bank_name':
                row['bank_name'] = str(val or '').strip()[:80]
            _save_profile(conn, i, row)
            n += 1
        _audit(conn, f'payroll_bulk_{action}', f'{n} موظف: {val if val is not None else ""}')
        conn.commit()
        return jsonify({'ok': True, 'updated': n})
    finally:
        conn.close()


@bp.route('/api/payroll/rules', methods=['POST'])
def api_rules():
    if not can_payroll():
        return _deny()
    d = request.get_json(silent=True) or {}
    new = _merge(DEFAULT_RULES, d)
    if new['day_basis'] not in ('thirty', 'workdays') or new['late']['mode'] not in ('minutes', 'tiers', 'count'):
        return jsonify({'ok': False, 'error': 'قيمة غير صالحة'}), 400
    new['decimals'] = max(0, min(3, int(_num(new.get('decimals'), 3))))
    new['currency'] = str(new.get('currency') or 'JOD')[:8]
    db.set_setting('payroll_rules_json', json.dumps(new, ensure_ascii=False))
    conn = db.get_db()
    try:
        _audit(conn, 'payroll_rules', 'تعديل قواعد الاحتساب')
        conn.commit()
    finally:
        conn.close()
    return jsonify({'ok': True, 'rules': new})


@bp.route('/api/payroll/access', methods=['POST'])
def api_access():
    if session.get('role') != 'admin':
        return _deny()
    d = request.get_json(silent=True) or {}
    ids = [int(i) for i in d.get('ids') or [] if str(i).isdigit()]
    grant = bool(d.get('grant'))
    conn = db.get_db()
    try:
        for i in ids:
            if grant:
                conn.execute("INSERT OR IGNORE INTO payroll_access (user_id) VALUES (?)", (i,))
            else:
                conn.execute("DELETE FROM payroll_access WHERE user_id = ?", (i,))
        _audit(conn, 'payroll_access_' + ('grant' if grant else 'revoke'), ids)
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


@bp.route('/api/payroll/runs', methods=['POST'])
def api_run_compute():
    if not can_payroll():
        return _deny()
    month = _month_arg((request.get_json(silent=True) or {}).get('month'))
    if not month:
        return jsonify({'ok': False, 'error': 'اختر الشهر'}), 400
    conn = db.get_db()
    try:
        try:
            run_id, n = build_run(conn, month, session.get('username') or '')
        except ValueError as e:
            return jsonify({'ok': False, 'error': str(e)}), 400
        conn.commit()
        if not n:
            return jsonify({'ok': True, 'id': run_id, 'count': 0,
                            'warning': 'لا يوجد موظف نشط له راتب أساسي — أدخل الرواتب من «رواتب الموظفين» أولاً.'})
        return jsonify({'ok': True, 'id': run_id, 'count': n})
    finally:
        conn.close()


@bp.route('/api/payroll/runs/<int:run_id>')
def api_run(run_id):
    if not can_payroll():
        return _deny()
    conn = db.get_db()
    try:
        run = conn.execute("SELECT * FROM payroll_runs WHERE id = ?", (run_id,)).fetchone()
        if not run:
            return jsonify({'ok': False, 'error': 'المسيرة غير موجودة'}), 404
        return jsonify({'ok': True, **_run_payload(conn, run)})
    finally:
        conn.close()


@bp.route('/api/payroll/runs/<int:run_id>/action', methods=['POST'])
def api_run_action(run_id):
    if not can_payroll():
        return _deny()
    d = request.get_json(silent=True) or {}
    action = d.get('action')
    conn = db.get_db()
    try:
        run = conn.execute("SELECT * FROM payroll_runs WHERE id = ?", (run_id,)).fetchone()
        if not run:
            return jsonify({'ok': False, 'error': 'المسيرة غير موجودة'}), 404
        if action == 'approve':
            if run['status'] == 'approved':
                return jsonify({'ok': True})
            conn.execute("UPDATE payroll_runs SET status = 'approved', approved_at = ?, approved_by = ? WHERE id = ?",
                         (_now(), session.get('username') or '', run_id))
            _audit(conn, 'payroll_approve', run['month'])
        elif action == 'reopen':
            note = (d.get('note') or '').strip()[:300]
            if not note:
                return jsonify({'ok': False, 'error': 'اكتب سبب فتح المسيرة'}), 400
            conn.execute("UPDATE payroll_runs SET status = 'draft', reopen_note = ? WHERE id = ?", (note, run_id))
            _audit(conn, 'payroll_reopen', f"{run['month']}: {note}")
        elif action == 'delete':
            if run['status'] == 'approved':
                return jsonify({'ok': False, 'error': 'لا تُحذف مسيرة معتمدة — افتحها أولاً'}), 400
            for t in ('payroll_lines', 'payroll_adjustments'):
                conn.execute(f"DELETE FROM {t} WHERE run_id = ?", (run_id,))
            conn.execute("DELETE FROM payroll_runs WHERE id = ?", (run_id,))
            _audit(conn, 'payroll_delete_draft', run['month'])
        else:
            return jsonify({'ok': False, 'error': 'طلب غير صالح'}), 400
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


@bp.route('/api/payroll/runs/<int:run_id>/adjust', methods=['POST'])
def api_adjust(run_id):
    """تعديل يدوي على موظف في مسودة: مكافأة أو خصم بمبلغ وسبب، أو حذف تعديل"""
    if not can_payroll():
        return _deny()
    d = request.get_json(silent=True) or {}
    conn = db.get_db()
    try:
        run = conn.execute("SELECT * FROM payroll_runs WHERE id = ?", (run_id,)).fetchone()
        if not run:
            return jsonify({'ok': False, 'error': 'المسيرة غير موجودة'}), 404
        if run['status'] == 'approved':
            return jsonify({'ok': False, 'error': 'المسيرة معتمدة ومقفلة'}), 400
        if d.get('delete_id'):
            row = conn.execute("SELECT * FROM payroll_adjustments WHERE id = ? AND run_id = ?", (int(d['delete_id']), run_id)).fetchone()
            conn.execute("DELETE FROM payroll_adjustments WHERE id = ? AND run_id = ?", (int(d['delete_id']), run_id))
            if row:
                _audit(conn, 'payroll_adjust_delete', f"{run['month']}: {row['kind']} {row['amount']} — {row['reason']}")
        else:
            ids = [int(i) for i in d.get('employee_ids') or [] if str(i).isdigit()]
            kind = d.get('kind')
            amount = _num(d.get('amount'))
            reason = (d.get('reason') or '').strip()[:200]
            if not ids or kind not in ('add', 'deduct') or amount <= 0 or not reason:
                return jsonify({'ok': False, 'error': 'حدّد الموظفين والنوع والمبلغ والسبب'}), 400
            for i in ids:
                conn.execute("INSERT INTO payroll_adjustments (run_id, employee_id, kind, amount, reason, created_by, created_at) "
                             "VALUES (?, ?, ?, ?, ?, ?, ?)", (run_id, i, kind, amount, reason, session.get('username') or '', _now()))
            _audit(conn, 'payroll_adjust', f"{run['month']}: {'مكافأة' if kind == 'add' else 'خصم'} {amount} × {len(ids)} — {reason}")
        build_run(conn, run['month'], session.get('username') or '')
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


# ---------------------------------------------------------------
#   المخرجات: كشف الرواتب، ملف البنك، القسائم
# ---------------------------------------------------------------
def _xlsx(title, header, rows, rtl=True, money_cols=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = rtl
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='0847B0')
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    for r in rows:
        ws.append(r)
    for i in range(1, len(header) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 16
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                c.number_format = '#,##0.000'
    ws.freeze_panes = 'A2'
    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return mem


def _send_xlsx(mem, name):
    return send_file(mem, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=name)


def _load_run(conn, run_id):
    run = conn.execute("SELECT * FROM payroll_runs WHERE id = ?", (run_id,)).fetchone()
    return (run, _run_payload(conn, run)) if run else (None, None)


@bp.route('/api/payroll/runs/<int:run_id>/register.xlsx')
def api_register(run_id):
    if not can_payroll():
        return _deny()
    conn = db.get_db()
    try:
        run, p = _load_run(conn, run_id)
    finally:
        conn.close()
    if not run:
        return jsonify({'ok': False, 'error': 'المسيرة غير موجودة'}), 404
    keys = []
    for ln in p['lines']:
        for e in ln['data']['earnings']:
            if e['label'] not in keys:
                keys.append(e['label'])
    dkeys = []
    for ln in p['lines']:
        for e in ln['data']['deductions']:
            if e['amount'] and e['label'] not in dkeys:
                dkeys.append(e['label'])
    header = ['الرقم الوظيفي', 'الموظف', 'القسم'] + keys + ['مكافآت يدوية', 'الإجمالي'] + dkeys + \
        ['خصومات يدوية', 'مجموع الخصومات', 'الصافي', 'ضمان المؤسسة', 'طريقة الدفع']
    rows = []
    conn = db.get_db()
    try:
        pay = {r['employee_id']: r['payment'] for r in conn.execute("SELECT employee_id, payment FROM payroll_profiles")}
    finally:
        conn.close()
    for ln in p['lines']:
        e = {x['label']: x['amount'] for x in ln['data']['earnings']}
        dd = {}
        for x in ln['data']['deductions']:
            dd[x['label']] = dd.get(x['label'], 0) + x['amount']
        adj = ln['data'].get('adjustments') or []
        rows.append([ln['employee_no'], ln['name'], ln['dept']] + [e.get(k, 0) for k in keys]
                    + [sum(a['amount'] for a in adj if a['kind'] == 'add'), ln['gross']] + [dd.get(k, 0) for k in dkeys]
                    + [sum(a['amount'] for a in adj if a['kind'] == 'deduct'), ln['deductions'], ln['net'],
                       ln['data'].get('employer_ss', 0), 'نقداً' if pay.get(ln['employee_id']) == 'cash' else 'بنك'])
    money = range(4, len(header))
    return _send_xlsx(_xlsx(f"رواتب {run['month']}", header, rows, money_cols=money), f"payroll-{run['month']}.xlsx")


@bp.route('/api/payroll/runs/<int:run_id>/bank.<fmt>')
def api_bank(run_id, fmt):
    if not can_payroll():
        return _deny()
    conn = db.get_db()
    try:
        run, p = _load_run(conn, run_id)
        prof = {r['employee_id']: dict(r) for r in conn.execute("SELECT * FROM payroll_profiles")}
    finally:
        conn.close()
    if not run:
        return jsonify({'ok': False, 'error': 'المسيرة غير موجودة'}), 404
    header = ['الرقم الوظيفي', 'الموظف', 'البنك', 'IBAN', 'الصافي', 'العملة', 'الشهر']
    rows = [[ln['employee_no'], ln['name'], prof.get(ln['employee_id'], {}).get('bank_name') or '',
             prof.get(ln['employee_id'], {}).get('iban') or '', ln['net'], p['currency'], run['month']]
            for ln in p['lines'] if prof.get(ln['employee_id'], {}).get('payment', 'bank') != 'cash' and ln['net'] > 0]
    if fmt == 'csv':
        buf = io.StringIO()
        buf.write('﻿')
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(rows)
        return send_file(io.BytesIO(buf.getvalue().encode('utf-8')), mimetype='text/csv', as_attachment=True,
                         download_name=f"bank-{run['month']}.csv")
    return _send_xlsx(_xlsx(f"بنك {run['month']}", header, rows, money_cols=(5,)), f"bank-{run['month']}.xlsx")


def _i18n_tags():
    """ترجمة القسيمة للإنجليزية بالقاموس نفسه حين تكون لغة الواجهة الإنجليزية (كوكي lang)"""
    try:
        from flask import url_for
        if request.cookies.get('lang') != 'en':
            return ''
        return (f'<script src="{url_for("static", filename="i18n/en.js")}"></script>'
                f'<script src="{url_for("static", filename="js/i18n.js")}"></script>')
    except Exception:   # noqa: BLE001
        return ''


def payslip_html(run, lines, org, currency, logo=None):
    import html as H
    esc = H.escape

    def money(v):
        return f'{v:,.3f}'
    parts = []
    for ln in lines:
        d = ln['data']
        earn = ''.join(f"<tr><td>{esc(e['label'])}</td><td class=n>{money(e['amount'])}</td></tr>" for e in d['earnings'])
        earn += ''.join(f"<tr><td>{esc(a['reason'] or 'مكافأة')}</td><td class=n>{money(a['amount'])}</td></tr>"
                        for a in d.get('adjustments') or [] if a['kind'] == 'add')
        ded = ''.join(f"<tr><td>{esc(x['label'])}{f' <small>({x['qty']:g} {esc(x['unit'])})</small>' if x.get('qty') and x.get('unit') else ''}</td>"
                      f"<td class=n>{money(x['amount'])}</td></tr>" for x in d['deductions'] if x['amount'])
        ded += ''.join(f"<tr><td>{esc(a['reason'] or 'خصم')}</td><td class=n>{money(a['amount'])}</td></tr>"
                       for a in d.get('adjustments') or [] if a['kind'] == 'deduct')
        att = d.get('attendance') or {}
        parts.append(f"""<section class=slip>
  <header><div><b>{esc(org)}</b><div class=m>قسيمة راتب — {esc(run['month'])}</div></div>
  <div class=who><b>{esc(ln['name'])}</b><div class=m>{esc(ln['dept'] or '')}{(' · ' + esc(ln['employee_no'])) if ln['employee_no'] else ''}</div></div></header>
  <div class=cols>
    <table><thead><tr><th>الاستحقاقات</th><th class=n>{esc(currency)}</th></tr></thead><tbody>{earn}</tbody>
      <tfoot><tr><td>الإجمالي</td><td class=n>{money(ln['gross'])}</td></tr></tfoot></table>
    <table><thead><tr><th>الاستقطاعات</th><th class=n>{esc(currency)}</th></tr></thead><tbody>{ded or '<tr><td>لا يوجد</td><td></td></tr>'}</tbody>
      <tfoot><tr><td>المجموع</td><td class=n>{money(ln['deductions'])}</td></tr></tfoot></table>
  </div>
  <div class=net>صافي الراتب: <b>{money(ln['net'])} {esc(currency)}</b></div>
  <div class=att>الحضور: {att.get('present', 0)} يوم · الغياب: {att.get('absent', 0)} · التأخير: {att.get('late_minutes', 0)} دقيقة ·
    المغادرة المبكرة: {att.get('early_minutes', 0)} دقيقة · الإجازات: {att.get('vacation_days', 0)} · العمل الإضافي: {round((att.get('ot_minutes') or 0) / 60, 2)} ساعة</div>
</section>""")
    return f"""<!DOCTYPE html><html lang=ar dir=rtl><head><meta charset=utf-8><meta name=viewport content="width=device-width,initial-scale=1">
<title>قسائم الرواتب {esc(run['month'])}</title><style>
body{{font-family:Tajawal,Tahoma,Arial,sans-serif;background:#eef2f7;margin:0;padding:16px;color:#1d2939}}
.slip{{background:#fff;max-width:780px;margin:0 auto 16px;border-radius:12px;padding:18px 20px;box-shadow:0 2px 10px rgba(0,0,0,.06);page-break-after:always}}
header{{display:flex;justify-content:space-between;gap:12px;border-bottom:2px solid #0847B0;padding-bottom:10px;margin-bottom:12px}}
.m{{color:#667085;font-size:.85rem}} .who{{text-align:left}}
.cols{{display:grid;grid-template-columns:1fr 1fr;gap:14px}} table{{width:100%;border-collapse:collapse;font-size:.9rem}}
th,td{{padding:6px 8px;border-bottom:1px solid #eaecf0;text-align:right}} th{{background:#f2f6fc;color:#0847B0}}
.n{{text-align:left;direction:ltr;white-space:nowrap}} tfoot td{{font-weight:700;border-top:2px solid #d0d5dd}}
.net{{margin-top:12px;padding:10px 12px;background:#ecfdf3;border-radius:8px;font-size:1.05rem}} .att{{margin-top:8px;color:#667085;font-size:.8rem}}
@media (max-width:640px){{.cols{{grid-template-columns:1fr}}}}
@media print{{body{{background:#fff;padding:0}} .slip{{box-shadow:none;margin:0 auto}} .noprint{{display:none}}}}
</style></head><body>
<div class=noprint style="max-width:780px;margin:0 auto 12px;text-align:left"><button onclick="window.print()" style="padding:8px 16px;border:0;border-radius:8px;background:#0847B0;color:#fff;font:inherit;cursor:pointer">طباعة</button></div>
{''.join(parts) or '<p>لا توجد قسائم</p>'}{_i18n_tags()}</body></html>"""


@bp.route('/api/payroll/runs/<int:run_id>/payslips')
def api_payslips(run_id):
    if not can_payroll():
        return _deny()
    ids = {int(i) for i in (request.args.get('ids') or '').split(',') if i.strip().isdigit()}
    conn = db.get_db()
    try:
        run, p = _load_run(conn, run_id)
    finally:
        conn.close()
    if not run:
        return jsonify({'ok': False, 'error': 'المسيرة غير موجودة'}), 404
    lines = [ln for ln in p['lines'] if not ids or ln['employee_id'] in ids]
    resp = Response(payslip_html(run, lines, db.get_setting('org_name', '') or '', p['currency']), mimetype='text/html')
    resp.headers['Cache-Control'] = 'no-store'
    return resp


# ---------------------------------------------------------------
#   الربط مع MenaPAY / MenaITech
# ---------------------------------------------------------------
def _item_values(x, line, rl):
    """قيم بنود التصدير لموظف: من الحضور، ومن سطر المسيرة إن وُجد (المبالغ)"""
    att = x or {}
    data = (line or {}).get('data') or {}
    ded_att = sum(d['amount'] for d in data.get('deductions', []) if d.get('attendance'))
    return {
        'present': att.get('present', 0), 'absent': att.get('absent', 0), 'unpaid_leave': att.get('unpaid_leave', 0),
        'vacation_days': att.get('vacation_days', 0), 'late_minutes': att.get('late_minutes', 0),
        'late_hours': round(att.get('late_minutes', 0) / 60, 2), 'early_minutes': att.get('early_minutes', 0),
        'no_out': att.get('no_out', 0), 'ot_hours': round(att.get('ot_minutes', 0) / 60, 2),
        'weekend_work': att.get('weekend_work', 0), 'deduct_days': data.get('deduct_days', 0),
        'deduct_amount': _round(ded_att, rl), 'ot_amount': data.get('ot_amount', 0), 'net': (line or {}).get('net', 0),
    }


@bp.route('/api/payroll/menapay', methods=['POST'])
def api_menapay_save():
    if not can_payroll():
        return _deny()
    d = request.get_json(silent=True) or {}
    cfg = menapay_cfg()
    if d.get('layout') in ('long', 'wide'):
        cfg['layout'] = d['layout']
    if d.get('code_field') in ('employee_no', 'device_uid', 'name'):
        cfg['code_field'] = d['code_field']
    if d.get('period_format') in ('%Y%m', '%Y-%m', '%m/%Y', '%m%Y'):
        cfg['period_format'] = d['period_format']
    by_key = {i.get('key'): i for i in d.get('items') or [] if isinstance(i, dict)}
    for it in cfg['items']:
        if it['key'] in by_key:
            it['on'] = bool(by_key[it['key']].get('on'))
            it['code'] = str(by_key[it['key']].get('code') or it['code']).strip()[:30]
    db.set_setting('payroll_menapay_json', json.dumps(cfg, ensure_ascii=False))
    return jsonify({'ok': True, 'menapay': cfg})


@bp.route('/api/payroll/menapay/export.<fmt>')
def api_menapay_export(fmt):
    """متغيرات الحضور للشهر بصيغة استيراد MenaPAY: صف لكل بند (رمز الموظف، رمز البند، القيمة، الفترة) أو صف لكل موظف"""
    if not can_payroll():
        return _deny()
    month = _month_arg(request.args.get('month'))
    if not month or fmt not in ('xlsx', 'csv'):
        return jsonify({'ok': False, 'error': 'اختر الشهر'}), 400
    cfg = menapay_cfg()
    rl = rules()
    conn = db.get_db()
    try:
        att = _attendance(conn, month)
        prof = _profiles(conn)
        run = conn.execute("SELECT id FROM payroll_runs WHERE month = ?", (month,)).fetchone()
        lines = {r['name']: {'net': r['net'], 'data': json.loads(r['data'])}
                 for r in conn.execute("SELECT * FROM payroll_lines WHERE run_id = ?", (run['id'],))} if run else {}
    finally:
        conn.close()
    items = [i for i in cfg['items'] if i['on']]
    if not items:
        return jsonify({'ok': False, 'error': 'اختر بنداً واحداً على الأقل للتصدير'}), 400
    period = datetime.strptime(month, '%Y-%m').strftime(cfg['period_format'])
    emps = sorted((p for p in prof.values() if p['active']), key=lambda p: (p['dept'] or '', p['name']))
    missing = []

    def code(p):
        v = p.get('employee_no') if cfg['code_field'] == 'employee_no' else p.get('device_uid') if cfg['code_field'] == 'device_uid' else p['name']
        if not v:
            missing.append(p['name'])
        return v or ''
    if cfg['layout'] == 'long':
        header = ['EmployeeCode', 'EmployeeName', 'ItemCode', 'Value', 'Period']
        rows = []
        for p in emps:
            vals = _item_values(att.get(p['name']), lines.get(p['name']), rl)
            c = code(p)
            for it in items:
                v = vals.get(it['key'], 0)
                if v:
                    rows.append([c, p['name'], it['code'], v, period])
    else:
        header = ['EmployeeCode', 'EmployeeName'] + [it['code'] for it in items] + ['Period']
        rows = []
        for p in emps:
            vals = _item_values(att.get(p['name']), lines.get(p['name']), rl)
            rows.append([code(p), p['name']] + [vals.get(it['key'], 0) for it in items] + [period])
    name = f"menapay-{month}"
    if fmt == 'csv':
        buf = io.StringIO()
        buf.write('﻿')
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(rows)
        resp = send_file(io.BytesIO(buf.getvalue().encode('utf-8')), mimetype='text/csv', as_attachment=True, download_name=name + '.csv')
    else:
        resp = _send_xlsx(_xlsx('MenaPAY', header, rows, rtl=False), name + '.xlsx')
    if missing:
        resp.headers['X-Missing-Codes'] = str(len(set(missing)))
    return resp


_COLS = {
    'employee_no': ('employee code', 'employee no', 'emp code', 'emp no', 'employee id', 'code', 'رقم الموظف', 'الرقم الوظيفي', 'رمز الموظف', 'رقم وظيفي'),
    'name': ('employee name', 'name', 'full name', 'arabic name', 'الاسم', 'اسم الموظف', 'الموظف', 'الاسم الكامل'),
    'dept': ('department', 'dept', 'section', 'القسم', 'الدائرة'),
    'basic': ('basic salary', 'basic', 'salary', 'الراتب الأساسي', 'الأساسي', 'الراتب'),
    'iban': ('iban', 'account', 'bank account', 'account no', 'رقم الحساب', 'الآيبان', 'ايبان'),
    'bank_name': ('bank', 'bank name', 'البنك', 'اسم البنك'),
    'ss_salary': ('ss salary', 'social security salary', 'insurable salary', 'راتب الضمان', 'الراتب الخاضع للضمان'),
}


def _norm(h):
    return re.sub(r'[\s_\-\.]+', ' ', str(h or '')).strip().lower()


def _read_table(filename, data):
    name = (filename or '').lower()
    if name.endswith('.csv') or name.endswith('.txt'):
        text = data.decode('utf-8-sig', errors='replace')
        return [r for r in csv.reader(io.StringIO(text))]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [[('' if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]


@bp.route('/api/payroll/import', methods=['POST'])
def api_import():
    """بيانات الموظفين والرواتب من ملف MenaITech/MenaPAY (xlsx أو csv): الرقم الوظيفي، الاسم، القسم، الأساسي، البنك، IBAN"""
    if not can_payroll():
        return _deny()
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': 'اختر الملف أولاً'}), 400
    create = request.form.get('create') == '1'
    try:
        table = _read_table(f.filename, f.read())
    except Exception as e:   # noqa: BLE001
        return jsonify({'ok': False, 'error': f'تعذّرت قراءة الملف ({type(e).__name__})'}), 400
    hdr_i, cols = None, {}
    for i, row in enumerate(table[:15]):
        found = {}
        for j, cell in enumerate(row):
            h = _norm(cell)
            for key, names in _COLS.items():
                if key not in found and h in names:
                    found[key] = j
        if 'name' in found or 'employee_no' in found:
            hdr_i, cols = i, found
            break
    if hdr_i is None:
        return jsonify({'ok': False, 'error': 'لم أجد عمود الاسم أو الرقم الوظيفي في أول 15 صفاً'}), 400
    conn = db.get_db()
    try:
        emps = {r['name']: r['id'] for r in conn.execute("SELECT id, name FROM employees")}
        by_no = {r['employee_no']: r['employee_id'] for r in conn.execute(
            "SELECT employee_id, employee_no FROM payroll_profiles WHERE employee_no IS NOT NULL AND employee_no != ''")}
        prof = _profiles(conn)
        updated = created = skipped = 0
        not_found = []
        for row in table[hdr_i + 1:]:
            def cell(k):
                j = cols.get(k)
                return str(row[j]).strip() if j is not None and j < len(row) and row[j] not in (None, '') else ''
            name, no = cell('name'), cell('employee_no')
            if re.fullmatch(r'\d+\.0', no):
                no = no[:-2]
            if not name and not no:
                continue
            emp_id = by_no.get(no) if no else None
            emp_id = emp_id or (emps.get(name) if name else None)
            if not emp_id:
                if create and name:
                    cur = conn.execute("INSERT INTO employees (name, dept, active) VALUES (?, ?, 1)", (name, cell('dept')))
                    emp_id = cur.lastrowid
                    emps[name] = emp_id
                    created += 1
                else:
                    if len(not_found) < 30:
                        not_found.append(name or no)
                    skipped += 1
                    continue
            p = prof.get(emp_id) or {'allowances': []}
            base = {k: p.get(k) for k in ('employee_no', 'basic', 'ss_subject', 'ss_salary', 'tax_monthly', 'payment', 'bank_name', 'iban', 'notes')}
            base['allowances'] = p.get('allowances') or []
            if no:
                base['employee_no'] = no
            if cell('basic'):
                base['basic'] = _num(cell('basic'))
            if cell('bank_name'):
                base['bank_name'] = cell('bank_name')
            if cell('iban'):
                base['iban'] = cell('iban')
            if cell('ss_salary'):
                base['ss_salary'] = _num(cell('ss_salary'))
            if base.get('ss_subject') is None:
                base['ss_subject'] = 1
            try:
                row_clean = _clean_profile(base)
            except ValueError:
                base['iban'] = ''
                row_clean = _clean_profile(base)
            _save_profile(conn, emp_id, row_clean)
            if cell('dept'):
                conn.execute("UPDATE employees SET dept = ? WHERE id = ? AND (dept IS NULL OR dept = '')", (cell('dept'), emp_id))
            updated += 1
        _audit(conn, 'payroll_import', f"{f.filename}: {updated} محدَّث، {created} جديد، {skipped} غير موجود")
        conn.commit()
        return jsonify({'ok': True, 'updated': updated, 'created': created, 'skipped': skipped, 'not_found': not_found,
                        'columns': {k: str(table[hdr_i][j]) for k, j in cols.items()}})
    finally:
        conn.close()


# ---------------------------------------------------------------
#   قسيمة الموظف في تطبيقه (بعد الاعتماد، إن فُعّل الخيار)
# ---------------------------------------------------------------
def register_me_routes():
    from mobile import _me_json

    @bp.route('/me/api/payslips')
    @_me_json
    def me_payslips(conn, acc, emp, cfg):
        rl = rules()
        if not rl.get('payslip_app'):
            return jsonify({'ok': True, 'enabled': False})
        rows = [{'run_id': r['run_id'], 'month': r['month'], 'net': r['net']} for r in conn.execute(
            """SELECT l.run_id, r.month, l.net FROM payroll_lines l JOIN payroll_runs r ON r.id = l.run_id
               WHERE l.employee_id = ? AND r.status = 'approved' ORDER BY r.month DESC LIMIT 24""", (emp['id'],))]
        return jsonify({'ok': True, 'enabled': True, 'rows': rows, 'currency': rl.get('currency', 'JOD')})

    @bp.route('/me/api/payslips/<int:run_id>')
    @_me_json
    def me_payslip(conn, acc, emp, cfg, run_id):
        if not rules().get('payslip_app'):
            return jsonify({'ok': False, 'error': 'غير متاح'}), 404
        run = conn.execute("SELECT * FROM payroll_runs WHERE id = ? AND status = 'approved'", (run_id,)).fetchone()
        line = conn.execute("SELECT * FROM payroll_lines WHERE run_id = ? AND employee_id = ?", (run_id, emp['id'])).fetchone() if run else None
        if not line:
            return jsonify({'ok': False, 'error': 'غير موجودة'}), 404
        ln = {**dict(line), 'data': json.loads(line['data'])}
        html = payslip_html(run, [ln], db.get_setting('org_name', '') or '', rules().get('currency', 'JOD'))
        return jsonify({'ok': True, 'html': html})


register_me_routes()
