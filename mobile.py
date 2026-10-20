"""
تطبيق الموظفين على الهاتف — تسجيل الحضور والانصراف والخروج المؤقت من أي مكان ضمن مواقع العمل المعتمدة.

• كل موظف له اسم مستخدم وكلمة مرور (منفصلة عن حسابات الإدارة) — صفحة الدخول /me/login
• ربط الحساب بهاتف واحد: أول دخول يربط الهاتف، ولا يمكن الدخول أو التسجيل من هاتف آخر حتى تفكّ الإدارة الربط
• Geofence: مواقع عمل (مركز + نصف قطر) لكل الأقسام أو لأقسام محددة — يرفض الخادم أي حركة من خارجها
  (أو بدقة موقع ضعيفة)، وتُسجَّل المحاولة المرفوضة مع السبب للمتابعة
• الحضور/الانصراف تدخل جدول البصمات نفسه (مع بصمات الأجهزة: أبكر حضور وآخر انصراف في اليوم) فتظهر في كل
  التقارير والتحليل؛ والخروج المؤقت يُحسب إذناً للمغادرة ويُخصم من ساعات العمل (قابل للإيقاف من الإعدادات)
• للإدارة: حسابات الموظفين (إنشاء جماعي)، مواقع العمل على الخريطة، سجل الحركات المباشر، والتصدير
"""
import csv
import hashlib
import io
import json
import math
import os
import secrets
import threading
import time
from datetime import datetime, timedelta
from functools import wraps

from flask import Blueprint, Response, jsonify, redirect, render_template, request, send_file, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash

import db

bp = Blueprint('mobile', __name__)

KINDS = {'in': 'حضور', 'out': 'انصراف', 'leave_out': 'خروج مؤقت', 'leave_back': 'عودة من الخروج'}
NEXT = {   # الحركة المسموحة حسب آخر حركة اليوم
    None: ['in'],
    'in': ['out', 'leave_out'],
    'leave_out': ['leave_back', 'out'],
    'leave_back': ['out', 'leave_out'],
    'out': ['in'],
}
LEAVE_REASON = 'خروج مؤقت عبر تطبيق الهاتف'
# مفتاح توقيع تطبيق الموظف (APK) — نفس القيمة في android/app.properties (يمكن تغييره بـ MOBILE_APP_KEY في .env + إعادة بناء التطبيق)
APP_KEY = os.environ.get('MOBILE_APP_KEY') or '8aGskrDZUF1q0VN5_Eny0xQVPLl4a2lYSOzl63c93_A'


# =====================================================
#   أدوات
# =====================================================
def now_local():
    try:
        from zoneinfo import ZoneInfo
        return datetime.now(ZoneInfo(os.environ.get('ATT_TZ', 'Asia/Amman'))).replace(tzinfo=None)
    except Exception:
        return datetime.now()


def haversine_m(lat1, lng1, lat2, lng2):
    r = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def _hash(token):
    return hashlib.sha256(('att-dev:' + (token or '')).encode()).hexdigest()


def _ip():
    return (request.headers.get('CF-Connecting-IP') or request.headers.get('X-Forwarded-For', '').split(',')[0]
            or request.remote_addr or '')[:45]


def _device_label():
    ua = request.headers.get('User-Agent', '')
    for key, label in (('iPhone', 'iPhone'), ('iPad', 'iPad'), ('Android', 'Android'), ('Windows', 'Windows'), ('Mac', 'Mac'), ('Linux', 'Linux')):
        if key in ua:
            model = ''
            if key == 'Android' and ';' in ua:
                try:
                    model = ua.split('Android')[1].split(';')[1].split(')')[0].split('Build')[0].strip()
                except IndexError:
                    model = ''
            return (label + (' — ' + model if model else ''))[:80]
    return ua[:80] or 'غير معروف'


def settings():
    return db.get_all_settings()


def _admin():
    """طلبات لوحة الإدارة — نفس صلاحيات app.py"""
    import app as app_module
    return app_module


def admin_required(view):
    @wraps(view)
    def wrapped(*a, **kw):
        if not session.get('user_id'):
            return jsonify({'ok': False, 'error': 'يرجى تسجيل الدخول أولاً'}), 401
        return view(*a, **kw)
    return wrapped


def _employee_scope_ok(dept):
    allowed = session.get('allowed_depts') or []
    return not allowed or (dept or '') in allowed


# =====================================================
#   مواقع العمل
# =====================================================
def site_rows(conn, active_only=False):
    q = "SELECT * FROM work_sites" + (" WHERE active = 1" if active_only else "") + " ORDER BY name"
    out = []
    for r in conn.execute(q).fetchall():
        d = dict(r)
        try:
            d['depts'] = json.loads(d.get('depts_json') or '[]')
        except ValueError:
            d['depts'] = []
        out.append(d)
    return out


def sites_for(conn, dept):
    return [s for s in site_rows(conn, True) if not s['depts'] or (dept or '') in s['depts']]


def active_override(conn, emp_id):
    """استثناء Geofence ساري الآن لهذا الموظف (لحالات نادرة: مهمة خارجية، عطل GPS…)"""
    now = now_local().strftime('%Y-%m-%d %H:%M')
    r = conn.execute("SELECT * FROM geofence_overrides WHERE employee_id = ? AND revoked = 0 AND starts_at <= ? AND ends_at >= ? "
                     "ORDER BY no_location DESC, ends_at DESC LIMIT 1", (emp_id, now, now)).fetchone()
    return dict(r) if r else None


# ===== طرق توثيق الدوام: أي المصادر تُحتسب لكل موظف =====
#   device = جهاز البصمة، phone = تطبيق الهاتف بالموقع الجغرافي، qr = رمز QR على جهاز المدخل
MODE_SOURCES = {
    'both': {'device', 'phone', 'qr'},
    'device': {'device'},
    'mobile': {'phone'},
    'qr': {'qr'},
    'device_qr': {'device', 'qr'},
    'mobile_qr': {'phone', 'qr'},
}
MODE_TEXT = {
    'both': 'جهاز البصمة + الهاتف + رمز QR', 'device': 'جهاز البصمة فقط', 'mobile': 'الهاتف فقط',
    'qr': 'رمز QR على جهاز المدخل فقط', 'device_qr': 'جهاز البصمة + رمز QR', 'mobile_qr': 'الهاتف + رمز QR',
}


def mode_allows(emp, source):
    return source in MODE_SOURCES.get((emp['attendance_mode'] if emp else None) or 'both', MODE_SOURCES['both'])


def check_location(conn, emp, acc, lat, lng, accuracy, cfg, override=None):
    """يرجع (ok, reason, site, distance)"""
    if override and override.get('no_location') and (lat is None or lng is None):
        return True, '', None, None
    if lat is None or lng is None:
        return False, 'لم يصل موقع الهاتف — فعّل خدمة الموقع (GPS) واسمح للمتصفح باستخدامه', None, None
    if override:
        best, best_d = None, None
        for s in sites_for(conn, emp['dept']):
            d = haversine_m(lat, lng, s['lat'], s['lng'])
            if best_d is None or d < best_d:
                best, best_d = s, d
        return True, '', best, best_d
    max_acc = float(cfg.get('mobile_max_accuracy') or 150)
    if accuracy is not None and accuracy > max_acc and accuracy >= 1000:
        return False, 'الموقع التقريبي لا يكفي لتسجيل الدوام — فعّل «الموقع الدقيق» من أذونات الموقع ثم أعد المحاولة', None, None
    if accuracy is not None and accuracy > max_acc:
        return False, f'دقة الموقع ضعيفة ({int(accuracy)} م) — انتظر قليلاً في مكان مفتوح ثم أعد المحاولة', None, None
    sites = sites_for(conn, emp['dept'])
    best, best_d = None, None
    for s in sites:
        d = haversine_m(lat, lng, s['lat'], s['lng'])
        if best_d is None or d - s['radius_m'] < best_d - best['radius_m']:
            best, best_d = s, d
    if acc['anywhere'] or cfg.get('mobile_require_geofence', '1') != '1':
        return True, '', best, best_d
    if not sites:
        return False, 'لم تُحدَّد مواقع عمل لقسمك بعد — راجع الإدارة', None, None
    slack = min(accuracy or 0, 30)   # هامش صغير لخطأ GPS
    if best_d - slack <= best['radius_m']:
        return True, '', best, best_d
    return False, f'أنت خارج موقع العمل ({int(best_d)} م عن «{best["name"]}»، المسموح {best["radius_m"]} م)', best, best_d


def _ip_in(ip, ips_text):
    import ipaddress
    try:
        addr = ipaddress.ip_address(ip)
    except ValueError:
        return False
    for part in (ips_text or '').replace('\n', ',').split(','):
        part = part.strip()
        if not part:
            continue
        try:
            if addr in ipaddress.ip_network(part, strict=False):
                return True
        except ValueError:
            continue
    return False


def is_android_browser():
    """هاتف أندرويد يفتح صفحة الموظف من المتصفح (لا من تطبيق Check-in)"""
    ua = request.headers.get('User-Agent', '')
    return 'Android' in ua and 'AttendanceApp' not in ua and not request.headers.get('X-App')


def require_app(cfg):
    """هل يلزم هذا الطلب أن يأتي من تطبيق أندرويد؟ (الكل، أو هواتف أندرويد فقط — الآيفون يبقى من الصفحة)"""
    v = cfg.get('mobile_require_app', '0')
    return v == '1' or (v == 'android' and is_android_browser())


def check_integrity(conn, emp, kind, data, lat, lng, accuracy, site, cfg):
    """حماية من تطبيقات الموقع الوهمي (Fake GPS).
    يرجع (سبب الرفض أو '', [علامات اشتباه], المصدر app/browser)."""
    import hashlib as _h
    import hmac
    flags = []
    source = 'browser'
    if request.headers.get('X-App'):
        ts = request.headers.get('X-App-Ts', '')
        js = lambda v: 'null' if v is None else v   # noqa: E731 — نفس تمثيل الجافاسكربت
        msg = f"{kind}|{js(data.get('lat'))}|{js(data.get('lng'))}|{ts}|{1 if data.get('bio') else 0}".encode()
        good = hmac.new(APP_KEY.encode(), msg, _h.sha256).hexdigest()
        try:
            fresh = abs(time.time() * 1000 - float(ts)) < 180000
        except ValueError:
            fresh = False
        if fresh and hmac.compare_digest(good, request.headers.get('X-App-Sig', '')):
            source = 'app'
        else:
            flags.append('توقيع تطبيق الموظف غير صحيح')
    # 1) التطبيق يكشف الموقع الوهمي من نظام أندرويد مباشرة
    if source == 'app' and data.get('mock'):
        return 'تم اكتشاف تطبيق موقع وهمي (Fake GPS) على هاتفك — أغلقه وألغِ «تطبيق الموقع الوهمي» من خيارات المطوّر ثم أعد المحاولة', flags, source
    if source == 'app' and data.get('rooted'):
        flags.append('الهاتف معدّل (root)')
    if source != 'app' and require_app(cfg):
        return 'على هواتف أندرويد يتم التسجيل من تطبيق Check-in فقط — نزّله من صفحة التطبيقات ثم سجّل منه', flags, source
    # 2) شبكة العمل: الموقع قد يشترط الاتصال بشبكة Wi-Fi المؤسسة (لا يزوّرها تطبيق الموقع الوهمي)
    if site and (site.get('ips') or '').strip():
        if not _ip_in(_ip(), site['ips']):
            if site.get('ip_required'):
                return f'يجب أن تكون متصلاً بشبكة Wi-Fi العمل في «{site["name"]}» للتسجيل', flags, source
            flags.append('ليس على شبكة العمل')
    # 3) إحداثيات ثابتة تماماً لعدة قراءات (GPS الحقيقي يتذبذب دائماً)
    samples = [x for x in (data.get('samples') or []) if isinstance(x, dict)][:20]
    try:
        if len(samples) >= 3:
            span = (float(samples[-1].get('t', 0)) - float(samples[0].get('t', 0))) / 1000
            same = len({(round(float(x['lat']), 7), round(float(x['lng']), 7), round(float(x.get('acc') or 0), 1)) for x in samples}) == 1
            if span >= 8 and same:
                flags.append('إحداثيات ثابتة تماماً (اشتباه موقع وهمي)')
    except (KeyError, TypeError, ValueError):
        pass
    if accuracy is not None and accuracy in (1.0, 3.0, 5.0) and source == 'browser':
        flags.append(f'دقة مثالية غير معتادة ({accuracy:g} م)')
    # 4) قفزة غير منطقية عن آخر موقع مسجّل
    prev = conn.execute("SELECT lat, lng, date, time FROM mobile_punches WHERE employee_id = ? AND lat IS NOT NULL "
                        "ORDER BY id DESC LIMIT 1", (emp['id'],)).fetchone()
    if prev and lat is not None:
        try:
            t_prev = datetime.strptime(f"{prev['date']} {prev['time']}", '%Y-%m-%d %H:%M:%S')
            hours = max((now_local() - t_prev).total_seconds() / 3600, 1 / 60)
            km = haversine_m(lat, lng, prev['lat'], prev['lng']) / 1000
            if hours < 6 and km > 2 and km / hours > 250:
                flags.append(f'قفزة غير منطقية في الموقع ({km:.0f} كم خلال {hours * 60:.0f} دقيقة)')
        except ValueError:
            pass
    if flags and cfg.get('mobile_suspicious', 'flag') == 'reject':
        return 'لم يُقبل التسجيل لاشتباه في موقع غير حقيقي: ' + '؛ '.join(flags), flags, source
    return '', flags, source


# =====================================================
#   حالة اليوم ودمج الحركات في البصمات
# =====================================================
def today_moves(conn, emp_id, date_str, source=None):
    """حركات اليوم المقبولة. source: None = الكل، 'kiosk' = جهاز المدخل فقط، '!kiosk' = الهاتف فقط"""
    extra = ''
    if source == 'kiosk':
        extra = " AND source = 'kiosk'"
    elif source == '!kiosk':
        extra = " AND COALESCE(source, '') != 'kiosk'"
    return [dict(r) for r in conn.execute(
        "SELECT kind, time, site_name, source FROM mobile_punches WHERE employee_id = ? AND date = ? AND status = 'accepted'" + extra +
        " ORDER BY time, id", (emp_id, date_str)).fetchall()]


def _merge_punch(conn, emp, date_str, time_str, kind, batch='mobile'):
    """الحضور/الانصراف في جدول البصمات (سجل واحد لليوم: أبكر حضور وآخر انصراف)"""
    # سجل الهاتف (mobile) وسجل جهاز المدخل (kiosk) منفصلان عن سجلات أجهزة البصمة — التحليل يدمجها حسب طريقة توثيق الموظف
    row = conn.execute("SELECT id, time_in, time_out FROM punches WHERE name = ? AND date = ? AND batch_id = ? "
                       "ORDER BY id LIMIT 1", (emp['name'], date_str, batch)).fetchone()
    if row:
        tin, tout = row['time_in'] or '', row['time_out'] or ''
        if kind == 'in':
            tin = min(tin, time_str) if tin else time_str
            tout = tout or time_str
        else:
            tout = max(tout, time_str) if tout else time_str
            tin = tin or time_str
        conn.execute("UPDATE punches SET time_in = ?, time_out = ?, synced = 0 WHERE id = ?", (tin, tout, row['id']))
    else:
        conn.execute("INSERT INTO punches (name, date, dept, time_in, time_out, batch_id) VALUES (?, ?, ?, ?, ?, ?)",
                     (emp['name'], date_str, emp['dept'] or '', time_str, time_str, batch))


def _apply_leave(conn, emp, date_str):
    """مجموع الخروج المؤقت في اليوم ← إذن مغادرة واحد (يخصمه التحليل من ساعات العمل)"""
    moves = today_moves(conn, emp['id'], date_str)
    total, first_out, out_at = 0, None, None
    for m in moves:
        if m['kind'] == 'leave_out':
            out_at = m['time']
            first_out = first_out or m['time']
        elif m['kind'] == 'leave_back' and out_at:
            t1 = datetime.strptime(out_at[:5], '%H:%M')
            t2 = datetime.strptime(m['time'][:5], '%H:%M')
            total += max(0, int((t2 - t1).total_seconds() // 60))
            out_at = None
    if not first_out or total <= 0:
        return
    ret = (datetime.strptime(first_out[:5], '%H:%M') + timedelta(minutes=total)).strftime('%H:%M')
    ex = conn.execute("SELECT id, reason FROM departure_permissions WHERE name = ? AND date = ?", (emp['name'], date_str)).fetchone()
    if ex and (ex['reason'] or '') != LEAVE_REASON:
        return   # إذن أدخلته الإدارة يدوياً — لا نغيّره
    if ex:
        conn.execute("UPDATE departure_permissions SET departure_time = ?, return_time = ? WHERE id = ?", (first_out[:5], ret, ex['id']))
    else:
        conn.execute("INSERT INTO departure_permissions (name, date, departure_time, return_time, reason) VALUES (?, ?, ?, ?, ?)",
                     (emp['name'], date_str, first_out[:5], ret, LEAVE_REASON))


_analysis = {'pending': False, 'lock': threading.Lock()}


def schedule_analysis(delay=4):
    """إعادة التحليل في الخلفية (مرة واحدة لعدة حركات متقاربة) حتى يبقى زر التسجيل سريعاً"""
    with _analysis['lock']:
        if _analysis['pending']:
            return
        _analysis['pending'] = True

    def run():
        time.sleep(delay)
        with _analysis['lock']:
            _analysis['pending'] = False
        conn = db.get_db()
        try:
            _admin().run_full_analysis(conn)
            conn.commit()
        except Exception:
            import logging
            logging.getLogger('attendance').exception('mobile analysis failed')
        finally:
            conn.close()
    threading.Thread(target=run, daemon=True).start()


# =====================================================
#   تطبيق الموظف
# =====================================================
_FAILS = {}


def _blocked(key):
    now = time.time()
    tries = [t for t in _FAILS.get(key, []) if now - t < 900]
    _FAILS[key] = tries
    return len(tries) >= 5


def _account(conn):
    eid = session.get('emp_id')
    if not eid:
        return None, None
    acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (eid,)).fetchone()
    emp = conn.execute("SELECT * FROM employees WHERE id = ?", (eid,)).fetchone()
    if not acc or not emp or not acc['active'] or not emp['active']:
        return None, None
    return acc, emp


DEV_COOKIE = 'att_dev_s'   # نسخة من رمز الهاتف يضعها الخادم (HttpOnly) — سفاري آيفون لا يمسحها بعد 7 أيام كما يمسح ما تكتبه الصفحة


def _token_matches(acc, token):
    return bool(acc['device_hash']) and bool(token) and len(token) >= 32 and secrets.compare_digest(acc['device_hash'], _hash(token))


def _device_ok(acc, token, cfg):
    if cfg.get('mobile_bind_device', '1') != '1':
        return True
    return _token_matches(acc, token) or _token_matches(acc, request.cookies.get(DEV_COOKIE, ''))


@bp.route('/me/login', methods=['GET', 'POST'])
def me_login():
    cfg = settings()
    if request.method == 'GET':
        return render_template('me_login.html', error=None, settings=cfg)
    username = (request.form.get('username') or '').strip().lower()
    password = request.form.get('password') or ''
    token = (request.form.get('device') or '').strip()
    key = f'{_ip()}|{username}'
    if _blocked(key):
        return render_template('me_login.html', error='محاولات كثيرة خاطئة — انتظر 15 دقيقة', settings=cfg), 429
    conn = db.get_db()
    try:
        acc = conn.execute("SELECT * FROM employee_accounts WHERE lower(username) = ?", (username,)).fetchone()
        emp = conn.execute("SELECT * FROM employees WHERE id = ?", (acc['employee_id'],)).fetchone() if acc else None
        if not acc or not check_password_hash(acc['password_hash'], password):
            _FAILS.setdefault(key, []).append(time.time())
            return render_template('me_login.html', error='اسم المستخدم أو كلمة المرور غير صحيحة', settings=cfg)
        if not acc['active'] or not emp or not emp['active']:
            return render_template('me_login.html', error='الحساب موقوف — راجع الإدارة', settings=cfg)
        saved = request.cookies.get(DEV_COOKIE, '')
        if cfg.get('mobile_bind_device', '1') == '1':
            if len(token) < 32 and len(saved) < 32:
                return render_template('me_login.html', error='تعذّر التعرف على الهاتف — فعّل الجافاسكربت وأعد المحاولة', settings=cfg)
            if acc['device_hash']:
                if _token_matches(acc, saved) and not _token_matches(acc, token):
                    token = saved          # سفاري مسح التخزين لكن نسخة الخادم باقية
                elif not _token_matches(acc, token):
                    return render_template('me_login.html', settings=cfg,
                                           error='هذا الحساب مرتبط بهاتف آخر. لا يمكن الدخول إلا من الهاتف المرتبط — اطلب من الإدارة فكّ الربط إن تغيّر هاتفك.')
            if len(token) < 32:
                token = saved
            if not acc['device_hash']:
                conn.execute("UPDATE employee_accounts SET device_hash = ?, device_label = ?, bound_at = ? WHERE id = ?",
                             (_hash(token), _device_label(), now_local().strftime('%Y-%m-%d %H:%M'), acc['id']))
        conn.execute("UPDATE employee_accounts SET last_login_at = ? WHERE id = ?", (now_local().strftime('%Y-%m-%d %H:%M'), acc['id']))
        conn.commit()
    finally:
        conn.close()
    _FAILS.pop(key, None)
    return _start_session(emp, token, redirect(url_for('mobile.me_home')))


def _start_session(emp, token, resp):
    for k in ('emp_id', 'emp_name', 'quick_acc', 'quick_nonce'):
        session.pop(k, None)
    session.permanent = True
    session['emp_id'] = emp['id']
    session['emp_name'] = emp['name']
    if token and len(token) >= 32:
        resp.set_cookie(DEV_COOKIE, token, max_age=3650 * 86400, path='/me', httponly=True,
                        secure=request.is_secure or request.headers.get('X-Forwarded-Proto') == 'https'
                        or request.headers.get('CF-Visitor', '').find('https') >= 0, samesite='Lax')
    return resp


# =====================================================
#   الدخول السريع: الهاتف المرتبط + Face ID / البصمة (بدل اسم المستخدم وكلمة المرور)
# =====================================================
def _quick_account(conn, token):
    """الحساب المرتبط بهذا الهاتف (من رمز الهاتف أو نسخة الخادم)"""
    for t in (token, request.cookies.get(DEV_COOKIE, '')):
        if t and len(t) >= 32:
            acc = conn.execute("SELECT * FROM employee_accounts WHERE device_hash = ?", (_hash(t),)).fetchone()
            if acc:
                emp = conn.execute("SELECT * FROM employees WHERE id = ?", (acc['employee_id'],)).fetchone()
                if acc['active'] and emp and emp['active']:
                    return acc, emp, t
    return None, None, None


@bp.route('/me/quick/options', methods=['POST'])
def me_quick_options():
    """هل يمكن الدخول السريع من هذا الهاتف؟ يرجع اسم الموظف وطريقة التحقق المتاحة"""
    d = request.get_json(silent=True) or {}
    key = f'{_ip()}|quick'
    if _blocked(key):
        return jsonify({'ok': False, 'error': 'محاولات كثيرة — انتظر 15 دقيقة أو ادخل بكلمة المرور'}), 429
    conn = db.get_db()
    try:
        acc, emp, _t = _quick_account(conn, (d.get('device') or '').strip())
        if not acc:
            return jsonify({'ok': False})
        session['quick_acc'] = acc['id']
        nonce = _b64u(secrets.token_bytes(24))
        session['quick_nonce'] = nonce
        out = {'ok': True, 'name': emp['name'], 'nonce': nonce, 'webauthn': None}
        creds = [r['cred_id'] for r in conn.execute("SELECT cred_id FROM webauthn_credentials WHERE employee_id = ?", (emp['id'],))]
        if creds and _webauthn_ready():
            chal = _b64u(secrets.token_bytes(32))
            session['wa_chal'] = chal
            out['webauthn'] = {'challenge': chal, 'rpId': _rp()[0], 'timeout': 60000, 'userVerification': 'required',
                               'allowCredentials': [{'type': 'public-key', 'id': c} for c in creds]}
        return jsonify(out)
    finally:
        conn.close()


@bp.route('/me/quick/login', methods=['POST'])
def me_quick_login():
    import hashlib as _h
    import hmac
    d = request.get_json(silent=True) or {}
    key = f'{_ip()}|quick'
    if _blocked(key):
        return jsonify({'ok': False, 'error': 'محاولات كثيرة — انتظر 15 دقيقة أو ادخل بكلمة المرور'}), 429
    conn = db.get_db()
    try:
        acc, emp, token = _quick_account(conn, (d.get('device') or '').strip())
        nonce = session.pop('quick_nonce', None)
        if not acc or session.pop('quick_acc', None) != acc['id'] or not nonce:
            return jsonify({'ok': False, 'error': 'انتهت صلاحية الطلب — أعد المحاولة'}), 400
        err = ''
        if request.headers.get('X-App'):
            # تطبيق أندرويد: البصمة تمت داخل التطبيق، والطلب موقّع من التطبيق
            ts = request.headers.get('X-App-Ts', '')
            good = hmac.new(APP_KEY.encode(), f"login|{nonce}|{ts}|1".encode(), _h.sha256).hexdigest()
            try:
                fresh = abs(time.time() * 1000 - float(ts)) < 120000
            except ValueError:
                fresh = False
            if not (fresh and hmac.compare_digest(good, request.headers.get('X-App-Sig', ''))):
                err = 'تعذّر التحقق من التطبيق — ادخل بكلمة المرور'
        elif d.get('webauthn'):
            err = verify_assertion(conn, emp['id'], d['webauthn'])
        else:
            err = 'أكّد هويتك بالبصمة أو الوجه'
        if err:
            _FAILS.setdefault(key, []).append(time.time())
            return jsonify({'ok': False, 'error': err}), 403
        conn.execute("UPDATE employee_accounts SET last_login_at = ? WHERE id = ?", (now_local().strftime('%Y-%m-%d %H:%M'), acc['id']))
        conn.commit()
    finally:
        conn.close()
    _FAILS.pop(key, None)
    return _start_session(emp, token, jsonify({'ok': True, 'next': url_for('mobile.me_home')}))


@bp.route('/me/forgot', methods=['GET', 'POST'])
def me_forgot():
    cfg = settings()
    if request.method == 'GET':
        return render_template('me_forgot.html', settings=cfg, done=False, error=None)
    A = _admin()
    if A.forgot_blocked():
        return render_template('me_forgot.html', settings=cfg, done=False, error='طلبات كثيرة — حاول بعد 15 دقيقة'), 429
    username = (request.form.get('username') or '').strip().lower()
    conn = db.get_db()
    try:
        acc = conn.execute("SELECT * FROM employee_accounts WHERE lower(username) = ?", (username,)).fetchone()
        emp = conn.execute("SELECT * FROM employees WHERE id = ?", (acc['employee_id'],)).fetchone() if acc else None
        if acc and emp:
            sent = False
            if (emp['email'] or '').strip():
                token = A.reset_token_create(conn, 'emp', emp['id'])
                conn.commit()
                sent = A.reset_mail(emp['email'].strip(), emp['name'], url_for('mobile.me_reset', token=token, _external=True))
            if not sent:
                A.reset_request_add(conn, 'emp', emp['id'], acc['username'])
            _audit_emp(conn, 'password_forgot', 'employee_accounts', emp['name'], 'رابط بالبريد' if sent else 'طلب للإدارة')
            conn.commit()
    finally:
        conn.close()
    return render_template('me_forgot.html', settings=cfg, done=True, error=None)


@bp.route('/me/reset/<token>', methods=['GET', 'POST'])
def me_reset(token):
    cfg = settings()
    A = _admin()
    conn = db.get_db()
    try:
        row = A.reset_token_get(conn, 'emp', token)
        if not row:
            return render_template('me_reset.html', settings=cfg, error='الرابط غير صالح أو انتهت صلاحيته — اطلب رابطاً جديداً', ok=False, invalid=True)
        if request.method == 'GET':
            return render_template('me_reset.html', settings=cfg, error=None, ok=False, invalid=False)
        pw, pw2 = request.form.get('password') or '', request.form.get('password2') or ''
        if len(pw) < 6 or pw != pw2:
            return render_template('me_reset.html', settings=cfg, error='كلمة المرور 6 أحرف على الأقل، ويجب أن تتطابق الخانتان', ok=False, invalid=False)
        conn.execute("UPDATE employee_accounts SET password_hash = ? WHERE employee_id = ?", (generate_password_hash(pw), row['subject_id']))
        conn.execute("UPDATE password_reset_tokens SET used_at = ? WHERE id = ?", (now_local().strftime('%Y-%m-%d %H:%M:%S'), row['id']))
        conn.execute("UPDATE password_reset_requests SET status = 'done', handled_by = '(رابط البريد)', handled_at = ? "
                     "WHERE kind = 'emp' AND subject_id = ? AND status = 'pending'", (now_local().strftime('%Y-%m-%d %H:%M:%S'), row['subject_id']))
        emp = conn.execute("SELECT name FROM employees WHERE id = ?", (row['subject_id'],)).fetchone()
        _audit_emp(conn, 'password_reset', 'employee_accounts', emp['name'] if emp else '', 'عبر رابط البريد')
        conn.commit()
        return render_template('me_reset.html', settings=cfg, error=None, ok=True, invalid=False)
    finally:
        conn.close()


@bp.route('/me/logout')
def me_logout():
    session.pop('emp_id', None)
    session.pop('emp_name', None)
    return redirect(url_for('mobile.me_login'))


@bp.route('/me')
@bp.route('/me/')
def me_home():
    conn = db.get_db()
    try:
        acc, emp = _account(conn)
    finally:
        conn.close()
    if not acc:
        return redirect(url_for('mobile.me_login'))
    return render_template('me.html', emp=emp, settings=settings())


def _me_json(fn):
    @wraps(fn)
    def wrapped(*a, **kw):
        conn = db.get_db()
        try:
            acc, emp = _account(conn)
            if not acc:
                return jsonify({'ok': False, 'error': 'انتهت الجلسة — سجّل الدخول من جديد', 'login': True}), 401
            cfg = settings()
            token = request.headers.get('X-Device') or ''
            if not _device_ok(acc, token, cfg):
                return jsonify({'ok': False, 'error': 'هذا الحساب مرتبط بهاتف آخر — لا يمكن التسجيل من هذا الجهاز', 'login': True}), 403
            return fn(conn, acc, emp, cfg, *a, **kw)
        finally:
            conn.close()
    return wrapped


@bp.route('/me/api/status')
@_me_json
def me_status(conn, acc, emp, cfg):
    leave_sweep(conn)
    now = now_local()
    d = now.strftime('%Y-%m-%d')
    moves = today_moves(conn, emp['id'], d)
    last = moves[-1]['kind'] if moves else None
    sites = sites_for(conn, emp['dept'])
    month = now.strftime('%Y-%m')
    days = [dict(r) for r in conn.execute(
        "SELECT date, time_in, time_out, status, late_minutes, working_hours FROM report_cache "
        "WHERE name = ? AND date LIKE ? ORDER BY date DESC", (emp['name'], month + '%')).fetchall()]
    recent = [dict(r) for r in conn.execute(
        "SELECT date, time, kind, status, reason, site_name FROM mobile_punches WHERE employee_id = ? "
        "ORDER BY id DESC LIMIT 30", (emp['id'],)).fetchall()]
    for r in recent:
        r['kind_label'] = KINDS.get(r['kind'], r['kind'])
    ov = active_override(conn, emp['id'])
    notes = [dict(r) for r in conn.execute(
        "SELECT id, title, body, created_at, read_at FROM employee_notifications WHERE employee_id = ? ORDER BY id DESC LIMIT 30",
        (emp['id'],)).fetchall()]
    creds = conn.execute("SELECT COUNT(*) FROM webauthn_credentials WHERE employee_id = ?", (emp['id'],)).fetchone()[0]
    return jsonify({
        'mode': emp['attendance_mode'] or 'both', 'mode_text': MODE_TEXT.get(emp['attendance_mode'] or 'both', ''),
        'sources': sorted(MODE_SOURCES.get(emp['attendance_mode'] or 'both', MODE_SOURCES['both'])), 'server_ms': int(time.time() * 1000),
        'override': {'until': ov['ends_at'], 'reason': ov['reason'], 'no_location': bool(ov['no_location'])} if ov else None,
        'notifications': notes, 'unread': sum(1 for n in notes if not n['read_at']),
        'biometric': _bio_mode(acc, cfg), 'bio_enrolled': creds > 0, 'webauthn': _webauthn_ready(),
        'presence': presence_state(conn, emp, acc, cfg), 'presence_check': presence_check_for(conn, emp, acc, cfg),
        'leave': _leave_info(conn, emp, cfg, now),
        'share': __import__('shares').share_info(conn, emp['id']),
        'ok': True, 'now': now.strftime('%Y-%m-%d %H:%M:%S'), 'date': d,
        'name': emp['name'], 'dept': emp['dept'] or '',
        'enabled': cfg.get('mobile_enabled', '1') == '1',
        'moves': [{**m, 'label': KINDS[m['kind']]} for m in moves],
        'last': last, 'allowed': NEXT.get(last, ['in']),
        'anywhere': bool(acc['anywhere']) or cfg.get('mobile_require_geofence', '1') != '1',
        'max_accuracy': float(cfg.get('mobile_max_accuracy') or 150),
        'require_app': require_app(cfg),
        'sites': [{'name': s['name'], 'lat': s['lat'], 'lng': s['lng'], 'radius': s['radius_m']} for s in sites],
        'month': days, 'recent': recent,
        'summary': {
            'days': sum(1 for x in days if x['time_in']),
            'late_days': sum(1 for x in days if (x['late_minutes'] or 0) > 0),
            'late_minutes': sum(int(x['late_minutes'] or 0) for x in days),
            'absent': sum(1 for x in days if x['status'] == 'غائب'),
            'hours': round(sum(float(x['working_hours'] or 0) for x in days), 1),
        },
    })


# =====================================================
#   الخروج المؤقت: من–إلى، حدّ أقصى، موافقة اختيارية، متابعة العودة
# =====================================================
LEAVE_STATUS = {
    'pending': 'بانتظار الموافقة', 'approved': 'معتمد', 'rejected': 'مرفوض', 'cancelled': 'ألغاه الموظف',
    'expired': 'انتهى دون استخدام', 'active': 'خارج الآن', 'overdue': 'تأخر عن العودة',
    'returned': 'عاد في الوقت', 'returned_late': 'عاد متأخراً', 'not_returned': 'لم يعد (مغادرة مبكرة)',
}


def _lcfg(cfg):
    def num(k, d, lo, hi):
        try:
            return max(lo, min(hi, int(float(cfg.get(k) or d))))
        except (TypeError, ValueError):
            return d
    return {'max': num('leave_max_minutes', 60, 0, 600), 'grace': num('leave_grace_min', 10, 0, 120),
            'approval': cfg.get('leave_approval', '0') == '1', 'early': cfg.get('leave_unreturned_early', '1') == '1'}


def _hm(t):
    """HH:MM أو HH:MM:SS ← دقائق منذ منتصف الليل (أو None)"""
    try:
        h, m = [int(x) for x in str(t).split(':')[:2]]
        return h * 60 + m if 0 <= h < 24 and 0 <= m < 60 else None
    except (TypeError, ValueError):
        return None


def _fmt(mins):
    mins = max(0, min(23 * 60 + 59, int(mins)))
    return f'{mins // 60:02d}:{mins % 60:02d}'


def _active_leave(conn, emp_id, date):
    return conn.execute("SELECT * FROM leave_requests WHERE employee_id = ? AND date = ? AND status IN ('active', 'overdue') "
                        "ORDER BY id DESC LIMIT 1", (emp_id, date)).fetchone()


def _leave_check_out(conn, emp, cfg, now, data):
    """قبل تسجيل «خروج مؤقت»: يرجع (رسالة رفض أو '', صف الطلب المعتمد أو None، وقت العودة HH:MM)"""
    lc = _lcfg(cfg)
    date, cur = now.strftime('%Y-%m-%d'), now.hour * 60 + now.minute
    if lc['approval']:
        rows = conn.execute("SELECT * FROM leave_requests WHERE employee_id = ? AND date = ? AND status = 'approved' ORDER BY from_time",
                            (emp['id'], date)).fetchall()
        for r in rows:
            f, t = _hm(r['from_time']), _hm(r['to_time'])
            if f is not None and t is not None and f - 15 <= cur <= t:
                return '', r, r['to_time']
        return 'لا يوجد إذن خروج معتمد لهذا الوقت — اطلب الإذن من «طلب إذن خروج» وانتظر الموافقة', None, None
    to = _hm(data.get('leave_to'))
    if to is None:
        to = cur + (lc['max'] or 60)
    if to <= cur:
        return 'وقت العودة يجب أن يكون بعد الوقت الحالي', None, None
    if lc['max'] and to - cur > lc['max']:
        return f'أقصى مدة للخروج المؤقت {lc["max"]} دقيقة — اختر وقت عودة أقرب', None, None
    return '', None, _fmt(to)


def _leave_after_punch(conn, emp, kind, now, data, req, to_time):
    """بعد قبول الحركة: تحديث سجل الخروج المؤقت"""
    date, t = now.strftime('%Y-%m-%d'), now.strftime('%H:%M')
    ts = now.strftime('%Y-%m-%d %H:%M:%S')
    if kind == 'leave_out':
        if req is not None:
            conn.execute("UPDATE leave_requests SET status = 'active', out_time = ? WHERE id = ?", (t, req['id']))
        else:
            conn.execute("INSERT INTO leave_requests (employee_id, name, dept, date, from_time, to_time, reason, status, out_time, requested_at, source) "
                         "VALUES (?, ?, ?, ?, ?, ?, ?, 'active', ?, ?, 'direct')",
                         (emp['id'], emp['name'], emp['dept'] or '', date, t, to_time, (data.get('leave_reason') or '').strip()[:300], t, ts))
    elif kind in ('leave_back', 'out'):
        lv = _active_leave(conn, emp['id'], date)
        if lv:
            late = (_hm(t) or 0) > (_hm(lv['to_time']) or 0) + _lcfg(settings())['grace']
            conn.execute("UPDATE leave_requests SET status = ?, back_time = ? WHERE id = ?",
                         ('returned_late' if late or lv['status'] == 'overdue' else 'returned', t, lv['id']))


def leave_sweep(conn=None):
    """كل دقيقة: انتهاء الطلبات غير المستخدمة، التأخر عن العودة، ومن لم يعد حتى نهاية الدوام"""
    own = conn is None
    conn = conn or db.get_db()
    try:
        cfg = settings()
        lc = _lcfg(cfg)
        now = now_local()
        if outage_hold(conn, now):
            return      # الخادم كان متوقفاً: لم يكن بإمكان الموظفين تسجيل العودة — ننتظر المهلة
        today, cur = now.strftime('%Y-%m-%d'), now.hour * 60 + now.minute
        # طلبات معتمدة/معلّقة فات وقتها
        for r in conn.execute("SELECT id, date, to_time FROM leave_requests WHERE status IN ('pending', 'approved')").fetchall():
            if r['date'] < today or (r['date'] == today and (_hm(r['to_time']) or 0) < cur):
                conn.execute("UPDATE leave_requests SET status = 'expired' WHERE id = ?", (r['id'],))
        for lv in conn.execute("SELECT * FROM leave_requests WHERE status IN ('active', 'overdue')").fetchall():
            emp = conn.execute("SELECT * FROM employees WHERE id = ?", (lv['employee_id'],)).fetchone()
            if not emp:
                continue
            end = _work_end(conn, emp, cfg, now)
            end_m = end.hour * 60 + end.minute
            day_over = lv['date'] < today or cur > end_m
            if day_over:
                # لم يعد حتى نهاية الدوام: وقت خروجه يُحتسب انصرافاً (مغادرة مبكرة)
                conn.execute("UPDATE leave_requests SET status = 'not_returned' WHERE id = ?", (lv['id'],))
                if lc['early'] and lv['out_time']:
                    has_out = conn.execute("SELECT 1 FROM mobile_punches WHERE employee_id = ? AND date = ? AND kind = 'out' AND status = 'accepted' "
                                           "AND time >= ?", (emp['id'], lv['date'], lv['out_time'])).fetchone()
                    if not has_out:
                        t = lv['out_time'] + ':00'
                        conn.execute("INSERT INTO mobile_punches (employee_id, name, dept, kind, date, time, status, flags, source) "
                                     "VALUES (?, ?, ?, 'out', ?, ?, 'accepted', ?, 'system')",
                                     (emp['id'], emp['name'], emp['dept'] or '', lv['date'], t, 'انصراف محسوب: لم يعد من الخروج المؤقت'))
                        _merge_punch(conn, emp, lv['date'], t, 'out')
                conn.execute("INSERT INTO presence_events (employee_id, name, dept, date, kind, start_ts, end_ts, minutes, details, open) "
                             "VALUES (?, ?, ?, ?, 'leave_not_returned', ?, ?, 0, ?, 0)",
                             (emp['id'], emp['name'], emp['dept'] or '', lv['date'], f"{lv['date']} {lv['out_time'] or lv['from_time']}:00",
                              f"{lv['date']} {_fmt(end_m)}:00", f"خرج {lv['out_time']} وكان يُفترض أن يعود {lv['to_time']}"))
                schedule_analysis(2)
            elif lv['status'] == 'active' and cur > (_hm(lv['to_time']) or 0) + lc['grace']:
                conn.execute("UPDATE leave_requests SET status = 'overdue' WHERE id = ?", (lv['id'],))
                conn.execute("INSERT INTO presence_events (employee_id, name, dept, date, kind, start_ts, end_ts, minutes, details, open) "
                             "VALUES (?, ?, ?, ?, 'leave_overdue', ?, ?, 0, ?, 0)",
                             (emp['id'], emp['name'], emp['dept'] or '', lv['date'], f"{lv['date']} {lv['to_time']}:00",
                              now.strftime('%Y-%m-%d %H:%M:%S'), f"خرج {lv['out_time']} وكان يُفترض أن يعود {lv['to_time']}"))
                conn.execute("INSERT INTO employee_notifications (employee_id, title, body, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
                             (emp['id'], 'تجاوزت وقت العودة', f"كان وقت عودتك من الخروج المؤقت {lv['to_time']} — سجّل «عودة من الخروج» فور وصولك",
                              'النظام', now.strftime('%Y-%m-%d %H:%M:%S')))
                try:
                    push_to_employee(conn, emp['id'], 'تجاوزت وقت العودة', f"كان وقت عودتك {lv['to_time']} — سجّل «عودة من الخروج» فور وصولك")
                except Exception:   # noqa: BLE001
                    pass
        conn.commit()
    finally:
        if own:
            conn.close()


def _leave_info(conn, emp, cfg, now):
    lc = _lcfg(cfg)
    date = now.strftime('%Y-%m-%d')
    rows = [dict(r, status_label=LEAVE_STATUS.get(r['status'], r['status'])) for r in conn.execute(
        "SELECT * FROM leave_requests WHERE employee_id = ? AND date >= ? ORDER BY date, from_time LIMIT 30", (emp['id'], date))]
    active = next((r for r in rows if r['status'] in ('active', 'overdue') and r['date'] == date), None)
    return {'max': lc['max'], 'grace': lc['grace'], 'approval': lc['approval'], 'requests': rows, 'active': active}


@bp.route('/me/api/leave/request', methods=['POST'])
@_me_json
def me_leave_request(conn, acc, emp, cfg):
    """طلب إذن خروج (عند تفعيل الموافقة المسبقة)"""
    lc = _lcfg(cfg)
    d = request.get_json(silent=True) or {}
    now = now_local()
    date = (d.get('date') or now.strftime('%Y-%m-%d'))[:10]
    f, t = _hm(d.get('from')), _hm(d.get('to'))
    reason = (d.get('reason') or '').strip()[:300]
    if f is None or t is None or t <= f:
        return jsonify({'ok': False, 'error': 'حدّد وقت الخروج ووقت العودة (العودة بعد الخروج)'}), 400
    if date < now.strftime('%Y-%m-%d') or (date == now.strftime('%Y-%m-%d') and t <= now.hour * 60 + now.minute):
        return jsonify({'ok': False, 'error': 'الوقت المطلوب انقضى'}), 400
    if lc['max'] and t - f > lc['max']:
        return jsonify({'ok': False, 'error': f'أقصى مدة للخروج المؤقت {lc["max"]} دقيقة'}), 400
    if not reason:
        return jsonify({'ok': False, 'error': 'اكتب سبب الخروج'}), 400
    conn.execute("INSERT INTO leave_requests (employee_id, name, dept, date, from_time, to_time, reason, status, requested_at, source) "
                 "VALUES (?, ?, ?, ?, ?, ?, ?, 'pending', ?, 'request')",
                 (emp['id'], emp['name'], emp['dept'] or '', date, _fmt(f), _fmt(t), reason, now.strftime('%Y-%m-%d %H:%M:%S')))
    conn.commit()
    return jsonify({'ok': True, 'message': 'أُرسل الطلب — يصلك إشعار عند الموافقة'})


@bp.route('/me/api/leave/cancel', methods=['POST'])
@_me_json
def me_leave_cancel(conn, acc, emp, cfg):
    rid = (request.get_json(silent=True) or {}).get('id')
    n = conn.execute("UPDATE leave_requests SET status = 'cancelled' WHERE id = ? AND employee_id = ? AND status IN ('pending', 'approved')",
                     (rid, emp['id'])).rowcount
    conn.commit()
    return jsonify({'ok': bool(n), 'error': '' if n else 'لا يمكن إلغاء هذا الطلب'})


# ----- الإدارة: طلبات الخروج، السجل، التقرير الشهري -----
def _leave_query(args):
    today = now_local().strftime('%Y-%m-%d')
    d_from, d_to = (args.get('from') or today)[:10], (args.get('to') or today)[:10]
    where, params = ["date >= ?", "date <= ?"], [d_from, d_to]
    if args.get('status'):
        where.append("status = ?"); params.append(args['status'])
    return ' AND '.join(where), params, d_from, d_to


@bp.route('/api/leaves')
@admin_required
def api_leaves():
    leave_sweep()
    where, params, d_from, d_to = _leave_query(request.args)
    conn = db.get_db()
    try:
        rows = [dict(r, status_label=LEAVE_STATUS.get(r['status'], r['status'])) for r in conn.execute(
            f"SELECT * FROM leave_requests WHERE {where} ORDER BY date DESC, from_time DESC, id DESC LIMIT 5000", params)]
        rows = [r for r in rows if _employee_scope_ok(r['dept'])]
        bw, bp_, _f, _t = _leave_query({k: v for k, v in request.args.items() if k != 'status'})
        counts = {r['status']: r['n'] for r in conn.execute(f"SELECT status, COUNT(*) AS n FROM leave_requests WHERE {bw} GROUP BY status", bp_)}
        pending = [dict(r) for r in conn.execute(
            "SELECT * FROM leave_requests WHERE status = 'pending' ORDER BY date, from_time").fetchall() if _employee_scope_ok(r['dept'])]
        for r in rows:
            if r['out_time'] and r['back_time']:
                r['minutes'] = max(0, (_hm(r['back_time']) or 0) - (_hm(r['out_time']) or 0))
            r['late_min'] = max(0, (_hm(r['back_time']) or 0) - (_hm(r['to_time']) or 0)) if r['back_time'] else None
        if request.args.get('format') == 'csv':
            import csv
            import io
            L = _admin().L
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow([L(x) for x in ['التاريخ', 'الموظف', 'القسم', 'من', 'إلى', 'خرج', 'عاد', 'المدة (دقيقة)', 'التأخر (دقيقة)', 'الحالة', 'السبب', 'قرار', 'بواسطة']])
            for r in rows:
                w.writerow([r['date'], r['name'], r['dept'], r['from_time'], r['to_time'], r['out_time'] or '', r['back_time'] or '',
                            r.get('minutes') or '', r['late_min'] or '', L(r['status_label']), r['reason'] or '', r['decision_note'] or '', r['decided_by'] or ''])
            return Response('\ufeff' + buf.getvalue(), content_type='text/csv; charset=utf-8',
                            headers={'Content-Disposition': f'attachment; filename=leaves-{d_from}_{d_to}.csv'})
        return jsonify({'ok': True, 'rows': rows, 'counts': counts, 'total': sum(counts.values()), 'statuses': LEAVE_STATUS,
                        'pending': pending})
    finally:
        conn.close()


@bp.route('/api/leaves/decide', methods=['POST'])
@admin_required
def api_leaves_decide():
    d = request.get_json(silent=True) or {}
    ids = [int(i) for i in (d.get('ids') or []) if str(i).isdigit()]
    action, note = d.get('action'), (d.get('note') or '').strip()[:300]
    if not ids or action not in ('approve', 'reject'):
        return jsonify({'ok': False, 'error': 'حدّد طلباً واحداً على الأقل'}), 400
    conn = db.get_db()
    try:
        now = now_local().strftime('%Y-%m-%d %H:%M:%S')
        n = 0
        for r in conn.execute(f"SELECT * FROM leave_requests WHERE id IN ({','.join('?' * len(ids))}) AND status = 'pending'", ids).fetchall():
            if not _employee_scope_ok(r['dept']):
                continue
            st = 'approved' if action == 'approve' else 'rejected'
            conn.execute("UPDATE leave_requests SET status = ?, decided_by = ?, decided_at = ?, decision_note = ? WHERE id = ?",
                         (st, session.get('username', ''), now, note, r['id']))
            title = 'تمت الموافقة على إذن الخروج' if st == 'approved' else 'رُفض طلب إذن الخروج'
            body = f"{r['date']} من {r['from_time']} إلى {r['to_time']}" + (f"\nملاحظة: {note}" if note else '')
            conn.execute("INSERT INTO employee_notifications (employee_id, title, body, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
                         (r['employee_id'], title, body, session.get('username', ''), now))
            try:
                push_to_employee(conn, r['employee_id'], title, body)
            except Exception:   # noqa: BLE001
                pass
            _audit_emp(conn, f'leave_{action}', 'leave_requests', r['name'], f"{r['date']} {r['from_time']}–{r['to_time']}")
            n += 1
        conn.commit()
        return jsonify({'ok': True, 'count': n})
    finally:
        conn.close()


@bp.route('/api/leaves/delete', methods=['POST'])
@admin_required
def api_leaves_delete():
    ids = [int(i) for i in ((request.get_json(silent=True) or {}).get('ids') or []) if str(i).isdigit()]
    if not ids:
        return jsonify({'ok': False, 'error': 'حدّد سجلاً واحداً على الأقل'}), 400
    conn = db.get_db()
    try:
        n = conn.execute(f"DELETE FROM leave_requests WHERE id IN ({','.join('?' * len(ids))})", ids).rowcount
        _audit(conn, 'leave_delete', 'leave_requests', str(n))
        conn.commit()
        return jsonify({'ok': True, 'count': n})
    finally:
        conn.close()


@bp.route('/api/leaves/report')
@admin_required
def api_leaves_report():
    """تقرير شهري لكل موظف: عدد مرات الخروج، المدة، التأخر عن العودة، عدم العودة، الطلبات المرفوضة"""
    month = (request.args.get('month') or now_local().strftime('%Y-%m'))[:7]
    conn = db.get_db()
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM leave_requests WHERE date LIKE ?", (month + '%',))]
        rows = [r for r in rows if _employee_scope_ok(r['dept'])]
        agg = {}
        for r in rows:
            a = agg.setdefault(r['employee_id'], {'employee_id': r['employee_id'], 'name': r['name'], 'dept': r['dept'], 'times': 0, 'minutes': 0,
                                                   'late': 0, 'late_minutes': 0, 'not_returned': 0, 'requested': 0, 'rejected': 0, 'direct': 0})
            if r['source'] == 'request':
                a['requested'] += 1
            if r['status'] == 'rejected':
                a['rejected'] += 1
            if r['out_time']:
                a['times'] += 1
                if r['source'] == 'direct':
                    a['direct'] += 1
                if r['back_time']:
                    a['minutes'] += max(0, (_hm(r['back_time']) or 0) - (_hm(r['out_time']) or 0))
                if r['status'] in ('returned_late', 'overdue'):
                    a['late'] += 1
                    if r['back_time']:
                        a['late_minutes'] += max(0, (_hm(r['back_time']) or 0) - (_hm(r['to_time']) or 0))
                if r['status'] == 'not_returned':
                    a['not_returned'] += 1
        out = sorted(agg.values(), key=lambda x: (-x['times'], x['name'] or ''))
        if request.args.get('format') == 'csv':
            import csv
            import io
            L = _admin().L
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow([L(x) for x in ['الموظف', 'القسم', 'مرات الخروج', 'المدة (دقيقة)', 'تأخر عن العودة', 'دقائق التأخر', 'لم يعد', 'طلبات', 'مرفوضة']])
            for a in out:
                w.writerow([a['name'], a['dept'], a['times'], a['minutes'], a['late'], a['late_minutes'], a['not_returned'], a['requested'], a['rejected']])
            return Response('\ufeff' + buf.getvalue(), content_type='text/csv; charset=utf-8',
                            headers={'Content-Disposition': f'attachment; filename=leave-report-{month}.csv'})
        return jsonify({'ok': True, 'month': month, 'rows': out,
                        'totals': {k: sum(a[k] for a in out) for k in ('times', 'minutes', 'late', 'not_returned', 'requested', 'rejected')}})
    finally:
        conn.close()


@bp.route('/me/api/punch', methods=['POST'])
@_me_json
def me_punch(conn, acc, emp, cfg):
    data = request.get_json(silent=True) or {}
    kind = data.get('kind')
    if kind not in KINDS:
        return jsonify({'ok': False, 'error': 'حركة غير معروفة'}), 400
    if cfg.get('mobile_enabled', '1') != '1':
        return jsonify({'ok': False, 'error': 'التسجيل من الهاتف متوقف حالياً من الإدارة'}), 403

    def num(k):
        try:
            v = float(data.get(k))
            return v if math.isfinite(v) else None
        except (TypeError, ValueError):
            return None
    lat, lng, accuracy = num('lat'), num('lng'), num('accuracy')
    if lat is not None and not (-90 <= lat <= 90 and -180 <= lng <= 180):
        lat = lng = None
    now = now_local()
    d, t = now.strftime('%Y-%m-%d'), now.strftime('%H:%M:%S')
    moves = today_moves(conn, emp['id'], d)
    last = moves[-1]['kind'] if moves else None
    recent = conn.execute("SELECT created_at FROM mobile_punches WHERE employee_id = ? AND status = 'accepted' "
                          "AND created_at >= datetime('now', '-30 seconds')", (emp['id'],)).fetchone()

    if not mode_allows(emp, 'phone'):
        return jsonify({'ok': False, 'error': 'طريقة توثيق دوامك: ' + MODE_TEXT.get(emp['attendance_mode'] or 'both', '') + ' — التسجيل من الهاتف غير متاح لك'}), 403
    override = active_override(conn, emp['id'])
    ok, reason, site, dist = check_location(conn, emp, acc, lat, lng, accuracy, cfg, override)
    flags, source = [], 'browser'
    bio_err = check_biometric(conn, acc, emp, kind, data, cfg)
    if ok and bio_err:
        ok, reason = False, bio_err
    if override:
        flags.append(f"استثناء موقع: {override['reason']}")
    if (lat is not None or request.headers.get('X-App')) and ok:
        bad, more, source = check_integrity(conn, emp, kind, data, lat, lng, accuracy, None if override else site, cfg)
        flags += more
        if bad and not (override and bad.startswith('لم يُقبل التسجيل لاشتباه')):
            ok, reason = False, bad
    elif lat is not None or request.headers.get('X-App'):
        bad, more, source = check_integrity(conn, emp, kind, data, lat, lng, accuracy, None, cfg)
        flags += more
        if data.get('mock') and source == 'app':
            reason = bad or reason
    # ===== ردع التلاعب بالوقت: الحركة تُسجَّل دائماً بوقت الخادم، وساعة الهاتف المختلفة تُعلَّم للمراجعة =====
    clock_note = ''
    try:
        skew_min = abs(time.time() * 1000 - float(data.get('client_ts'))) / 60000
        if skew_min >= 2:
            flags.append(f'ساعة الهاتف مختلفة عن الوقت الصحيح ({int(skew_min)} دقيقة)')
            clock_note = ' — تنبيه: ساعة هاتفك غير مضبوطة؛ سُجّلت الحركة بوقت النظام وأُبلغت الإدارة'
    except (TypeError, ValueError):
        pass
    if data.get('auto_time') is False:
        flags.append('الضبط التلقائي للوقت معطّل في الهاتف')
    if data.get('dev_opts') is True:
        flags.append('خيارات المطوّر مفعّلة في الهاتف')
    if ok and kind not in NEXT.get(last, ['in']):
        ok, reason = False, f'لا يمكن تسجيل «{KINDS[kind]}» الآن — آخر حركة: {KINDS.get(last, "لا شيء")}'
    if ok and recent:
        ok, reason = False, 'سُجّلت حركة قبل لحظات — انتظر نصف دقيقة'
    leave_req, leave_to = None, None
    if ok and kind == 'leave_out':
        lerr, leave_req, leave_to = _leave_check_out(conn, emp, cfg, now, data)
        if lerr:
            ok, reason = False, lerr
    conn.execute(
        "INSERT INTO mobile_punches (employee_id, name, dept, kind, date, time, lat, lng, accuracy, site_id, site_name, "
        "distance_m, status, reason, device_label, ip, flags, source) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (emp['id'], emp['name'], emp['dept'] or '', kind, d, t, lat, lng, accuracy, site['id'] if site else None,
         site['name'] if site else None, round(dist, 1) if dist is not None else None,
         'accepted' if ok else 'rejected', reason or None, acc['device_label'] or _device_label(), _ip(),
         '؛ '.join(flags) or None, source))
    if ok:
        if kind in ('leave_out', 'leave_back', 'out'):
            _leave_after_punch(conn, emp, kind, now, data, leave_req, leave_to)
        if kind in ('in', 'out'):
            _merge_punch(conn, emp, d, t, kind)
        if kind in ('leave_back', 'out') and cfg.get('mobile_leave_deduct', '1') == '1':
            conn.commit()
            _apply_leave(conn, emp, d)
    conn.commit()
    if not ok:
        return jsonify({'ok': False, 'error': reason, 'rejected': True}), 400
    if kind in ('in', 'out', 'leave_back'):
        schedule_analysis()
    if override and (not site or dist is None or dist > (site['radius_m'] or 0)):
        where = ' — باستثناء موقع' + (f' ({dist / 1000:.1f} كم عن «{site["name"]}»)' if site and dist is not None else '')
    else:
        where = f' — {site["name"]}' if site else ''
    if kind == 'leave_out' and leave_to:
        where += f' — وقت العودة المتوقع {leave_to}'
    return jsonify({'ok': True, 'message': f'تم تسجيل {KINDS[kind]} الساعة {t[:5]}' + where + clock_note,
                    'time': t[:5]})


@bp.route('/me/api/password', methods=['POST'])
@_me_json
def me_password(conn, acc, emp, cfg):
    data = request.get_json(silent=True) or {}
    if not check_password_hash(acc['password_hash'], data.get('current') or ''):
        return jsonify({'ok': False, 'error': 'كلمة المرور الحالية غير صحيحة'}), 400
    new = data.get('new') or ''
    if len(new) < 6:
        return jsonify({'ok': False, 'error': 'كلمة المرور الجديدة 6 أحرف على الأقل'}), 400
    conn.execute("UPDATE employee_accounts SET password_hash = ? WHERE id = ?", (generate_password_hash(new), acc['id']))
    conn.commit()
    return jsonify({'ok': True})


# =====================================================
#   الإدارة: حسابات الموظفين
# =====================================================
def _gen_password():
    return ''.join(secrets.choice('23456789') for _ in range(4)) + ''.join(secrets.choice('abcdefghjkmnpqrstuvwxyz') for _ in range(4))


def _unique_username(conn, base):
    base = ''.join(ch for ch in (base or '').lower() if ch.isalnum() or ch in '._-')[:30] or 'emp'
    name, i = base, 1
    while conn.execute("SELECT 1 FROM employee_accounts WHERE lower(username) = ?", (name,)).fetchone():
        i += 1
        name = f'{base}{i}'
    return name


def _audit_emp(conn, action, target, name, value=None):
    try:
        _admin().log_audit(conn, action, target, employee_name=name, new_value=value)
    except Exception:
        pass


def _audit(conn, action, target, value=None):
    try:
        _admin().log_audit(conn, action, target, new_value=value)
    except Exception:
        pass


@bp.route('/api/mobile/accounts', methods=['GET', 'POST'])
@admin_required
def api_accounts():
    conn = db.get_db()
    try:
        if request.method == 'POST':
            data = request.get_json(silent=True) or {}
            emp = conn.execute("SELECT * FROM employees WHERE id = ?", (data.get('employee_id'),)).fetchone()
            if not emp:
                return jsonify({'ok': False, 'error': 'الموظف غير موجود'}), 404
            username = (data.get('username') or '').strip().lower()
            password = data.get('password') or ''
            acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (emp['id'],)).fetchone()
            if username:
                clash = conn.execute("SELECT employee_id FROM employee_accounts WHERE lower(username) = ?", (username,)).fetchone()
                if clash and clash['employee_id'] != emp['id']:
                    return jsonify({'ok': False, 'error': 'اسم المستخدم مستخدم لموظف آخر'}), 400
            if password and len(password) < 6:
                return jsonify({'ok': False, 'error': 'كلمة المرور 6 أحرف على الأقل'}), 400
            if acc:
                if username:
                    conn.execute("UPDATE employee_accounts SET username = ? WHERE id = ?", (username, acc['id']))
                if password:
                    conn.execute("UPDATE employee_accounts SET password_hash = ? WHERE id = ?", (generate_password_hash(password), acc['id']))
                if 'anywhere' in data:
                    conn.execute("UPDATE employee_accounts SET anywhere = ? WHERE id = ?", (1 if data['anywhere'] else 0, acc['id']))
            else:
                if not username or not password:
                    return jsonify({'ok': False, 'error': 'اسم المستخدم وكلمة المرور مطلوبان'}), 400
                conn.execute("INSERT INTO employee_accounts (employee_id, username, password_hash, anywhere) VALUES (?, ?, ?, ?)",
                             (emp['id'], username, generate_password_hash(password), 1 if data.get('anywhere') else 0))
            _audit_emp(conn, 'mobile_account', 'employee_accounts', emp['name'], username)
            conn.commit()
            return jsonify({'ok': True})

        today = now_local().strftime('%Y-%m-%d')
        rows = []
        q = """SELECT e.id, e.name, e.dept, e.device_uid, e.active, e.attendance_mode, e.presence_monitor, a.username, a.active AS acc_active, a.anywhere,
                      a.device_label, a.bound_at, a.last_login_at, a.id AS acc_id, a.bio_mode
               FROM employees e LEFT JOIN employee_accounts a ON a.employee_id = e.id ORDER BY e.dept, e.name"""
        last = {r['employee_id']: dict(r) for r in conn.execute(
            "SELECT employee_id, kind, time, date FROM mobile_punches WHERE id IN "
            "(SELECT MAX(id) FROM mobile_punches WHERE status = 'accepted' GROUP BY employee_id)").fetchall()}
        for r in conn.execute(q).fetchall():
            d = dict(r)
            if not _employee_scope_ok(d['dept']):
                continue
            lm = last.get(d['id'])
            d['last_move'] = (f"{KINDS.get(lm['kind'])} {lm['time'][:5]}" + ('' if lm['date'] == today else f" ({lm['date']})")) if lm else ''
            d['has_account'] = bool(d['acc_id'])
            rows.append(d)
        return jsonify({'ok': True, 'rows': rows})
    finally:
        conn.close()


@bp.route('/api/mobile/accounts/bulk', methods=['POST'])
@admin_required
def api_accounts_bulk():
    data = request.get_json(silent=True) or {}
    ids = [int(i) for i in (data.get('ids') or []) if str(i).isdigit()]
    action = data.get('action')
    if not ids:
        return jsonify({'ok': False, 'error': 'حدّد موظفاً واحداً على الأقل'}), 400
    conn = db.get_db()
    try:
        emps = [dict(r) for r in conn.execute(
            f"SELECT * FROM employees WHERE id IN ({','.join('?' * len(ids))})", ids).fetchall()]
        emps = [e for e in emps if _employee_scope_ok(e['dept'])]
        out, n = [], 0
        for e in emps:
            acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (e['id'],)).fetchone()
            if action == 'create':
                if acc:
                    continue
                username = _unique_username(conn, e['device_uid'] and f"u{e['device_uid']}" or f"emp{e['id']}")
                pw = _gen_password()
                conn.execute("INSERT INTO employee_accounts (employee_id, username, password_hash) VALUES (?, ?, ?)",
                             (e['id'], username, generate_password_hash(pw)))
                out.append({'name': e['name'], 'dept': e['dept'] or '', 'username': username, 'password': pw})
            elif not acc:
                continue
            elif action == 'reset_password':
                pw = _gen_password()
                conn.execute("UPDATE employee_accounts SET password_hash = ? WHERE id = ?", (generate_password_hash(pw), acc['id']))
                out.append({'name': e['name'], 'dept': e['dept'] or '', 'username': acc['username'], 'password': pw})
            elif action == 'unbind':
                conn.execute("UPDATE employee_accounts SET device_hash = NULL, device_label = NULL, bound_at = NULL, qr_secret = NULL WHERE id = ?", (acc['id'],))
            elif action in ('enable', 'disable'):
                conn.execute("UPDATE employee_accounts SET active = ? WHERE id = ?", (1 if action == 'enable' else 0, acc['id']))
            elif action in ('anywhere_on', 'anywhere_off'):
                conn.execute("UPDATE employee_accounts SET anywhere = ? WHERE id = ?", (1 if action == 'anywhere_on' else 0, acc['id']))
            elif action == 'delete':
                conn.execute("DELETE FROM employee_accounts WHERE id = ?", (acc['id'],))
            else:
                return jsonify({'ok': False, 'error': 'إجراء غير معروف'}), 400
            n += 1
            _audit_emp(conn, f'mobile_{action}', 'employee_accounts', e['name'])
        conn.commit()
        return jsonify({'ok': True, 'count': n, 'credentials': out})
    finally:
        conn.close()


@bp.route('/api/employees/bulk-dept', methods=['POST'])
@admin_required
def api_employees_bulk_dept():
    """توزيع الموظفين على الأقسام: نقل المحددين إلى قسم (ومع سجلاتهم السابقة إن طُلب)"""
    data = request.get_json(silent=True) or {}
    ids = [int(i) for i in (data.get('ids') or []) if str(i).isdigit()]
    dept = (data.get('dept') or '').strip()
    if not ids:
        return jsonify({'ok': False, 'error': 'حدّد موظفاً واحداً على الأقل'}), 400
    conn = db.get_db()
    try:
        names = [r['name'] for r in conn.execute(
            f"SELECT name FROM employees WHERE id IN ({','.join('?' * len(ids))})", ids).fetchall()]
        conn.execute(f"UPDATE employees SET dept = ?, synced = 0 WHERE id IN ({','.join('?' * len(ids))})", [dept] + ids)
        if data.get('history', True) and names:
            ph = ','.join('?' * len(names))
            conn.execute(f"UPDATE punches SET dept = ?, synced = 0 WHERE name IN ({ph})", [dept] + names)
            conn.execute(f"UPDATE mobile_punches SET dept = ? WHERE name IN ({ph})", [dept] + names)
        _audit(conn, 'bulk_dept', 'employees', f'{len(names)} ← {dept or "بدون قسم"}')
        conn.commit()
        if data.get('history', True):
            _admin().run_full_analysis(conn)
            conn.commit()
        return jsonify({'ok': True, 'count': len(names)})
    finally:
        conn.close()


# =====================================================
#   الإدارة: مواقع العمل
# =====================================================
def _site_payload(data):
    try:
        lat, lng = float(data.get('lat')), float(data.get('lng'))
        radius = int(float(data.get('radius_m') or 150))
    except (TypeError, ValueError):
        return None, 'أدخل خط العرض والطول ونصف القطر بأرقام صحيحة'
    if not (-90 <= lat <= 90 and -180 <= lng <= 180):
        return None, 'إحداثيات غير صحيحة'
    if not 20 <= radius <= 20000:
        return None, 'نصف القطر بين 20 و 20000 متر'
    name = (data.get('name') or '').strip()
    if not name:
        return None, 'اسم الموقع مطلوب'
    depts = sorted({str(d).strip() for d in (data.get('depts') or []) if str(d).strip()})
    ips = ', '.join(x.strip() for x in str(data.get('ips') or '').replace('\n', ',').split(',') if x.strip())[:500]
    return {'name': name[:120], 'lat': lat, 'lng': lng, 'radius_m': radius, 'depts_json': json.dumps(depts, ensure_ascii=False),
            'active': 1 if data.get('active', True) else 0, 'ips': ips, 'ip_required': 1 if (data.get('ip_required') and ips) else 0}, ''


@bp.route('/api/mobile/sites', methods=['GET', 'POST'])
@admin_required
def api_sites():
    conn = db.get_db()
    try:
        if request.method == 'POST':
            p, err = _site_payload(request.get_json(silent=True) or {})
            if err:
                return jsonify({'ok': False, 'error': err}), 400
            conn.execute("INSERT INTO work_sites (name, lat, lng, radius_m, depts_json, active, ips, ip_required) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                         (p['name'], p['lat'], p['lng'], p['radius_m'], p['depts_json'], p['active'], p['ips'], p['ip_required']))
            _audit(conn, 'site_add', 'work_site', p['name'])
            fixed = recompute_presence_today(conn)
            conn.commit()
            return jsonify({'ok': True, 'recomputed': fixed})
        depts = sorted({r['dept'] for r in conn.execute("SELECT DISTINCT dept FROM employees WHERE dept IS NOT NULL AND dept != ''")})
        counts = {}
        for r in conn.execute("SELECT site_id, COUNT(*) AS n FROM mobile_punches WHERE status = 'accepted' GROUP BY site_id"):
            counts[r['site_id']] = r['n']
        rows = site_rows(conn)
        for s in rows:
            s['moves'] = counts.get(s['id'], 0)
        return jsonify({'ok': True, 'rows': rows, 'depts': depts})
    finally:
        conn.close()


@bp.route('/api/mobile/sites/<int:site_id>', methods=['PUT', 'DELETE'])
@admin_required
def api_site_edit(site_id):
    conn = db.get_db()
    try:
        if request.method == 'DELETE':
            conn.execute("DELETE FROM work_sites WHERE id = ?", (site_id,))
            _audit(conn, 'site_delete', 'work_site', str(site_id))
        else:
            data = request.get_json(silent=True) or {}
            if set(data) == {'active'}:
                conn.execute("UPDATE work_sites SET active = ? WHERE id = ?", (1 if data['active'] else 0, site_id))
            else:
                p, err = _site_payload(data)
                if err:
                    return jsonify({'ok': False, 'error': err}), 400
                conn.execute("UPDATE work_sites SET name = ?, lat = ?, lng = ?, radius_m = ?, depts_json = ?, active = ?, ips = ?, ip_required = ? WHERE id = ?",
                             (p['name'], p['lat'], p['lng'], p['radius_m'], p['depts_json'], p['active'], p['ips'], p['ip_required'], site_id))
            _audit(conn, 'site_edit', 'work_site', str(site_id))
        fixed = recompute_presence_today(conn)
        conn.commit()
        return jsonify({'ok': True, 'recomputed': fixed})
    finally:
        conn.close()


@bp.route('/api/mobile/settings', methods=['GET', 'POST'])
@admin_required
def api_mobile_settings():
    keys = ('mobile_enabled', 'mobile_bind_device', 'mobile_max_accuracy', 'mobile_require_geofence', 'mobile_leave_deduct',
            'mobile_require_app', 'mobile_suspicious', 'mobile_biometric',
            'presence_enabled', 'presence_scope', 'presence_interval_min', 'presence_grace_min',
            'presence_checks_per_day', 'presence_check_window_min', 'presence_depts_json',
            'leave_max_minutes', 'leave_grace_min', 'leave_approval', 'leave_unreturned_early')
    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        old, changed = settings(), []
        for k in keys:
            if k in data:
                v = data[k]
                if k == 'mobile_suspicious':
                    v = 'reject' if v == 'reject' else 'flag'
                elif k == 'mobile_require_app':   # 0 = الكل من أي مكان، android = هواتف أندرويد من التطبيق فقط، 1 = الجميع من التطبيق فقط
                    v = 'android' if v == 'android' else '1' if v in (True, 1, '1', 'true', 'on', 'all') else '0'
                elif k == 'mobile_biometric':
                    v = v if v in ('off', 'optional', 'required') else 'off'
                elif k == 'presence_scope':
                    v = 'all' if v == 'all' else 'selected'
                elif k == 'presence_depts_json':
                    v = json.dumps([str(x) for x in (v if isinstance(v, list) else [])][:200], ensure_ascii=False)
                elif k in ('presence_interval_min', 'presence_grace_min', 'presence_checks_per_day', 'presence_check_window_min',
                           'leave_max_minutes', 'leave_grace_min'):
                    try:
                        v = str(max(0, min(600, int(float(v)))))
                    except (TypeError, ValueError):
                        continue
                elif k == 'mobile_max_accuracy':
                    try:
                        v = str(max(10, min(1000, int(float(v)))))
                    except (TypeError, ValueError):
                        continue
                elif k not in ('mobile_suspicious', 'mobile_require_app', 'mobile_biometric', 'presence_scope', 'presence_depts_json'):
                    v = '1' if v in (True, 1, '1', 'true', 'on') else '0'
                if str(old.get(k, '')) != str(v):
                    changed.append(f"{k}: {old.get(k, '')} → {v}")
                db.set_setting(k, v)
        if changed:
            c = db.get_db()
            _audit(c, 'mobile_settings', 'settings', ' | '.join(changed))
            c.commit()
            c.close()
        return jsonify({'ok': True})
    cfg = settings()
    return jsonify({'ok': True, 'settings': {k: cfg.get(k) for k in keys}})


# =====================================================
#   الإدارة: سجل حركات الهاتف + المتابعة المباشرة + التصدير
# =====================================================
def _punch_query(args):
    today = now_local().strftime('%Y-%m-%d')
    d_from, d_to = args.get('from') or today, args.get('to') or args.get('from') or today
    where, params = ["date BETWEEN ? AND ?"], [d_from, d_to]
    for key, col in (('dept', 'dept'), ('name', 'name'), ('kind', 'kind'), ('status', 'status')):
        if args.get(key):
            where.append(f"{col} = ?")
            params.append(args[key])
    if args.get('flagged') == '1':
        where.append("flags IS NOT NULL AND flags != ''")
    allowed = session.get('allowed_depts') or []
    if allowed:
        where.append(f"dept IN ({','.join('?' * len(allowed))})")
        params += allowed
    return ' AND '.join(where), params, d_from, d_to


@bp.route('/api/mobile/punches')
@admin_required
def api_punches():
    where, params, d_from, d_to = _punch_query(request.args)
    conn = db.get_db()
    try:
        rows = [dict(r) for r in conn.execute(
            f"SELECT * FROM mobile_punches WHERE {where} ORDER BY date DESC, time DESC, id DESC LIMIT 3000", params).fetchall()]
        # الأعداد (بدون فلتر النوع/الحالة حتى تبقى الأزرار صحيحة)
        base_args = {k: v for k, v in request.args.items() if k not in ('kind', 'status', 'flagged')}
        w2, p2, _, _ = _punch_query(base_args)
        counts = {'all': 0, 'rejected': 0, 'flagged': 0, **{k: 0 for k in KINDS}}
        counts['flagged'] = conn.execute(f"SELECT COUNT(*) FROM mobile_punches WHERE {w2} AND flags IS NOT NULL AND flags != ''", p2).fetchone()[0]
        for r in conn.execute(f"SELECT kind, status, COUNT(*) AS n FROM mobile_punches WHERE {w2} GROUP BY kind, status", p2):
            counts['all'] += r['n']
            if r['status'] == 'rejected':
                counts['rejected'] += r['n']
            else:
                counts[r['kind']] = counts.get(r['kind'], 0) + r['n']
        for r in rows:
            r['kind_label'] = KINDS.get(r['kind'], r['kind'])
        names = sorted({r['name'] for r in conn.execute("SELECT DISTINCT name FROM mobile_punches")})
        return jsonify({'ok': True, 'rows': rows, 'counts': counts, 'from': d_from, 'to': d_to, 'names': names})
    finally:
        conn.close()


@bp.route('/api/mobile/punches/delete', methods=['POST'])
@admin_required
def api_mobile_punches_delete():
    """حذف حركات هاتف محددة (مرحلة التجارب أو حركة خاطئة) ثم إعادة بناء دوام الأيام المتأثرة"""
    ids = [int(i) for i in ((request.get_json(silent=True) or {}).get('ids') or []) if str(i).isdigit()]
    if not ids:
        return jsonify({'ok': False, 'error': 'حدّد حركة واحدة على الأقل'}), 400
    conn = db.get_db()
    try:
        ph = ','.join('?' * len(ids))
        rows = [r for r in conn.execute(f"SELECT id, employee_id, name, dept, date FROM mobile_punches WHERE id IN ({ph})", ids)
                if _employee_scope_ok(r['dept'])]
        if not rows:
            return jsonify({'ok': True, 'count': 0})
        conn.execute(f"DELETE FROM mobile_punches WHERE id IN ({','.join('?' * len(rows))})", [r['id'] for r in rows])
        # إعادة بناء صف الهاتف لكل (موظف، يوم) متأثر من الحركات المتبقية
        for emp_id, name, d in {(r['employee_id'], r['name'], r['date']) for r in rows}:
            emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
            conn.execute("DELETE FROM punches WHERE name = ? AND date = ? AND batch_id IN ('mobile', 'kiosk')", (name, d))
            if emp:
                for m in today_moves(conn, emp_id, d):
                    if m['kind'] in ('in', 'out'):
                        _merge_punch(conn, emp, d, m['time'], m['kind'], 'kiosk' if m.get('source') == 'kiosk' else 'mobile')
                try:
                    _apply_leave(conn, emp, d)
                except Exception:
                    pass
            _audit_emp(conn, 'mobile_punch_delete', 'mobile_punches', name, d)
        conn.commit()
        schedule_analysis(1)
        return jsonify({'ok': True, 'count': len(rows)})
    finally:
        conn.close()


@bp.route('/api/mobile/live')
@admin_required
def api_live():
    """آخر حركة لكل موظف اليوم + من لم يسجّل بعد (للمتابعة من أي مكان)"""
    today = now_local().strftime('%Y-%m-%d')
    conn = db.get_db()
    try:
        last = {}
        for r in conn.execute("SELECT * FROM mobile_punches WHERE date = ? AND status = 'accepted' ORDER BY time, id", (today,)):
            last[r['employee_id']] = dict(r)
        emps = [dict(r) for r in conn.execute(
            "SELECT e.id, e.name, e.dept FROM employees e JOIN employee_accounts a ON a.employee_id = e.id "
            "WHERE e.active = 1 AND a.active = 1 ORDER BY e.dept, e.name").fetchall()]
        # الموقع المباشر: من يشارك موقعه الآن (أو مفعّل له التتبع المستمر) يظهر في مكانه الحالي مع مساره،
        # وليس في مكان آخر حركة حضور فقط — وإلا تبقى النقطة ثابتة على الخريطة رغم تحرّك الموظف
        now_s = now_local().strftime('%Y-%m-%d %H:%M:%S')
        live = {}
        for r in conn.execute("SELECT id, employee_id, lat, lng, accuracy, updated_at FROM location_shares WHERE kind = 'live' "
                              "AND stopped_at IS NULL AND ends_at > ? AND lat IS NOT NULL ORDER BY id", (now_s,)):
            live[r['employee_id']] = dict(r)
        out, state = [], {'in': 0, 'out': 0, 'leave': 0, 'none': 0}
        for e in emps:
            if not _employee_scope_ok(e['dept']):
                continue
            m = last.get(e['id'])
            st = 'none' if not m else ('out' if m['kind'] == 'out' else 'leave' if m['kind'] == 'leave_out' else 'in')
            state[st] += 1
            row = {**e, 'state': st, 'kind_label': KINDS.get(m['kind']) if m else '', 'time': m['time'][:5] if m else '',
                   'site': m['site_name'] if m else '', 'lat': m['lat'] if m else None, 'lng': m['lng'] if m else None,
                   'live': False}
            s = live.get(e['id'])
            if s:
                try:
                    age = max(0, int((now_local() - datetime.strptime((s['updated_at'] or '')[:19], '%Y-%m-%d %H:%M:%S')).total_seconds()))
                except ValueError:
                    age = None
                tail = [[p['lat'], p['lng']] for p in conn.execute(
                    "SELECT lat, lng FROM (SELECT lat, lng, ts FROM location_share_points WHERE share_id = ? AND ts >= ? "
                    "ORDER BY ts DESC LIMIT 40) ORDER BY ts", (s['id'], today + ' 00:00:00'))]
                row.update({'live': True, 'share_id': s['id'], 'lat': s['lat'], 'lng': s['lng'], 'accuracy': s['accuracy'],
                            'live_at': (s['updated_at'] or '')[11:19], 'age_s': age, 'tail': tail})
            out.append(row)
        sites = site_rows(conn, True)
        return jsonify({'ok': True, 'date': today, 'rows': out, 'state': state, 'total': len(out),
                        'live': sum(1 for r in out if r['live']), 'server_time': now_s[11:19],
                        'sites': [{'name': s['name'], 'lat': s['lat'], 'lng': s['lng'], 'radius': s['radius_m']} for s in sites]})
    finally:
        conn.close()


@bp.route('/api/mobile/punches/export')
@admin_required
def api_punches_export():
    where, params, d_from, d_to = _punch_query(request.args)
    conn = db.get_db()
    try:
        rows = conn.execute(f"SELECT * FROM mobile_punches WHERE {where} ORDER BY date, time, id", params).fetchall()
    finally:
        conn.close()
    header = ['التاريخ', 'الوقت', 'الموظف', 'القسم', 'الحركة', 'الحالة', 'السبب', 'الموقع', 'المسافة (م)', 'دقة GPS (م)',
              'خط العرض', 'خط الطول', 'الهاتف', 'IP', 'المصدر', 'علامات الاشتباه']
    L = _admin().L
    header = [L(h) for h in header]
    data = [[r['date'], r['time'], r['name'], r['dept'], L(KINDS.get(r['kind'], r['kind'])),
             L('مقبولة' if r['status'] == 'accepted' else 'مرفوضة'), L(r['reason'] or ''), r['site_name'] or '',
             r['distance_m'], r['accuracy'], r['lat'], r['lng'], r['device_label'] or '', r['ip'] or '',
             L('جهاز المدخل' if r['source'] == 'kiosk' else 'التطبيق' if r['source'] == 'app' else 'المتصفح'), L(r['flags'] or '')] for r in rows]
    fname = f'mobile-punches-{d_from}_{d_to}'
    if request.args.get('format') == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = L('حركات الهاتف')
        ws.sheet_view.rightToLeft = _admin()._ui_lang(request) != 'en'
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1B6E8C')
        for row in data:
            ws.append(row)
        for i, w in enumerate((12, 10, 26, 18, 14, 10, 40, 20, 10, 10, 12, 12, 22, 16, 10, 40), 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=fname + '.xlsx')
    s = io.StringIO()
    s.write('﻿')
    w = csv.writer(s)
    w.writerow(header)
    w.writerows(data)
    return send_file(io.BytesIO(s.getvalue().encode('utf-8')), mimetype='text/csv', as_attachment=True, download_name=fname + '.csv')


@bp.route('/api/mobile/myip')
@admin_required
def api_myip():
    """عنوان IP الذي يرى به الخادم هذا الجهاز (لإضافة شبكة Wi-Fi العمل لموقع)"""
    return jsonify({'ok': True, 'ip': _ip()})


@bp.route('/me/app.json')
def me_app_info():
    """للتطبيق: هل التسجيل يتطلب التطبيق؟ (بدون بيانات حساسة)"""
    cfg = settings()
    return jsonify({'ok': True, 'require_app': require_app(cfg)})



# =====================================================
#   البصمة الحيوية (Face ID / بصمة الإصبع) قبل الحركة
#   • في تطبيق أندرويد: نافذة البصمة من النظام نفسه، والنتيجة داخل توقيع الحركة
#   • في المتصفح (آيفون/أندرويد): WebAuthn — مفتاح خاص لا يغادر الهاتف ويتحقق منه الخادم
# =====================================================
import base64 as _b64


def _b64u(b):
    return _b64.urlsafe_b64encode(b).rstrip(b'=').decode()


def _unb64u(s):
    s = (s or '').strip()
    return _b64.urlsafe_b64decode(s + '=' * (-len(s) % 4))


def _webauthn_ready():
    import importlib.util
    return importlib.util.find_spec('cryptography') is not None


def _rp():
    host = (request.host or '').split(':')[0]
    return host, f"{request.scheme}://{request.host}"


def _check_client_data(cdj_b64, kind):
    cd = json.loads(_unb64u(cdj_b64).decode('utf-8'))
    chal = session.pop('wa_chal', None)
    rp_id, origin = _rp()
    if cd.get('type') != kind or not chal or cd.get('challenge') != chal:
        raise ValueError('طلب البصمة منتهي الصلاحية — أعد المحاولة')
    if cd.get('origin') not in (origin, origin.replace('http://', 'https://')):
        raise ValueError('مصدر غير صحيح')
    return cd


def verify_assertion(conn, emp_id, a):
    """يتحقق من توقيع WebAuthn (مفتاح الهاتف) — يرجع '' إن صح أو سبب الخطأ"""
    if not _webauthn_ready():
        return 'التحقق بالبصمة غير متاح على الخادم (ثبّت مكتبة cryptography)'
    try:
        from cryptography.exceptions import InvalidSignature
        from cryptography.hazmat.primitives import hashes, serialization
        from cryptography.hazmat.primitives.asymmetric import ec, padding
        cred = conn.execute("SELECT * FROM webauthn_credentials WHERE employee_id = ? AND cred_id = ?",
                            (emp_id, a.get('id', ''))).fetchone()
        if not cred:
            return 'البصمة غير مسجّلة لهذا الحساب — سجّلها من «حسابي»'
        cdj = _unb64u(a.get('clientDataJSON'))
        _check_client_data(a.get('clientDataJSON'), 'webauthn.get')
        auth = _unb64u(a.get('authenticatorData'))
        rp_id, _origin = _rp()
        if auth[:32] != hashlib.sha256(rp_id.encode()).digest():
            return 'البصمة لا تخص هذا الموقع'
        flags = auth[32]
        if not (flags & 0x01 and flags & 0x04):
            return 'لم يتم التحقق من هوية صاحب الهاتف (البصمة/الوجه)'
        key = serialization.load_der_public_key(_unb64u(cred['public_key']))
        data = auth + hashlib.sha256(cdj).digest()
        sig = _unb64u(a.get('signature'))
        if isinstance(key, ec.EllipticCurvePublicKey):
            key.verify(sig, data, ec.ECDSA(hashes.SHA256()))
        else:
            key.verify(sig, data, padding.PKCS1v15(), hashes.SHA256())
        count = int.from_bytes(auth[33:37], 'big')
        if count and cred['sign_count'] and count <= cred['sign_count']:
            return 'تم رفض البصمة (عداد غير صحيح)'
        conn.execute("UPDATE webauthn_credentials SET sign_count = ? WHERE id = ?", (count, cred['id']))
        return ''
    except InvalidSignature:
        return 'توقيع البصمة غير صحيح'
    except ValueError as e:
        return str(e)
    except Exception as e:   # noqa: BLE001
        return f'تعذّر التحقق من البصمة: {type(e).__name__}'


def _bio_mode(acc, cfg):
    """طريقة بصمة الحضور الفعّالة لهذا الموظف: إعداده الخاص إن وُجد، وإلا الإعداد العام"""
    try:
        per = (acc['bio_mode'] or '').strip()
    except (KeyError, IndexError, TypeError):
        per = ''
    return per if per in ('off', 'optional', 'required') else cfg.get('mobile_biometric', 'off')


def check_biometric(conn, acc, emp, kind, data, cfg):
    mode = _bio_mode(acc, cfg)
    if mode == 'off':
        return ''
    enrolled = conn.execute("SELECT COUNT(*) FROM webauthn_credentials WHERE employee_id = ?", (emp['id'],)).fetchone()[0]
    from_app = bool(request.headers.get('X-App'))
    if from_app:
        if data.get('bio'):
            return ''   # نتيجة البصمة داخل توقيع التطبيق (يتحقق منه check_integrity)
        return 'أكّد هويتك بالبصمة أو الوجه أولاً' if mode == 'required' else ''
    if data.get('webauthn'):
        return verify_assertion(conn, emp['id'], data['webauthn'])
    if mode == 'required':
        return 'سجّل بصمة الوجه/الإصبع من «حسابي» أولاً' if not enrolled else 'أكّد هويتك بالبصمة أو الوجه أولاً'
    return ''


@bp.route('/me/api/webauthn/options', methods=['POST'])
@_me_json
def me_webauthn_options(conn, acc, emp, cfg):
    chal = _b64u(secrets.token_bytes(32))
    session['wa_chal'] = chal
    rp_id, _o = _rp()
    purpose = (request.get_json(silent=True) or {}).get('purpose')
    if purpose == 'register':
        existing = [r['cred_id'] for r in conn.execute("SELECT cred_id FROM webauthn_credentials WHERE employee_id = ?", (emp['id'],))]
        return jsonify({'ok': True, 'publicKey': {
            'challenge': chal, 'rp': {'id': rp_id, 'name': 'تسجيل الدوام'},
            'user': {'id': _b64u(f"emp-{emp['id']}".encode()), 'name': acc['username'], 'displayName': emp['name']},
            'pubKeyCredParams': [{'type': 'public-key', 'alg': -7}, {'type': 'public-key', 'alg': -257}],
            'authenticatorSelection': {'authenticatorAttachment': 'platform', 'userVerification': 'required', 'residentKey': 'discouraged'},
            'timeout': 60000, 'attestation': 'none',
            'excludeCredentials': [{'type': 'public-key', 'id': c} for c in existing]}})
    creds = [r['cred_id'] for r in conn.execute("SELECT cred_id FROM webauthn_credentials WHERE employee_id = ?", (emp['id'],))]
    if not creds:
        return jsonify({'ok': False, 'error': 'سجّل بصمة الوجه/الإصبع من «حسابي» أولاً'}), 400
    return jsonify({'ok': True, 'publicKey': {
        'challenge': chal, 'rpId': rp_id, 'timeout': 60000, 'userVerification': 'required',
        'allowCredentials': [{'type': 'public-key', 'id': c} for c in creds]}})


@bp.route('/me/api/webauthn/register', methods=['POST'])
@_me_json
def me_webauthn_register(conn, acc, emp, cfg):
    d = request.get_json(silent=True) or {}
    if not _webauthn_ready():
        return jsonify({'ok': False, 'error': 'غير متاح على الخادم'}), 400
    try:
        _check_client_data(d.get('clientDataJSON'), 'webauthn.create')
        from cryptography.hazmat.primitives import serialization
        serialization.load_der_public_key(_unb64u(d.get('publicKey')))   # يتأكد أنه مفتاح صالح
    except Exception as e:   # noqa: BLE001
        return jsonify({'ok': False, 'error': f'تعذّر تسجيل البصمة: {e}'}), 400
    conn.execute("INSERT OR REPLACE INTO webauthn_credentials (employee_id, cred_id, public_key, alg, label) VALUES (?, ?, ?, ?, ?)",
                 (emp['id'], d.get('id'), d.get('publicKey'), int(d.get('alg') or -7), _device_label()))
    conn.commit()
    return jsonify({'ok': True})


# =====================================================
#   إشعارات الموظف
# =====================================================
@bp.route('/me/api/notifications/read', methods=['POST'])
@_me_json
def me_notifications_read(conn, acc, emp, cfg):
    conn.execute("UPDATE employee_notifications SET read_at = ? WHERE employee_id = ? AND read_at IS NULL",
                 (now_local().strftime('%Y-%m-%d %H:%M'), emp['id']))
    conn.commit()
    return jsonify({'ok': True})


@bp.route('/api/mobile/notify', methods=['POST'])
@admin_required
def api_notify():
    d = request.get_json(silent=True) or {}
    ids = [int(i) for i in (d.get('ids') or []) if str(i).isdigit()]
    title, body = (d.get('title') or '').strip()[:200], (d.get('body') or '').strip()[:3000]
    if not ids or not title:
        return jsonify({'ok': False, 'error': 'حدّد الموظفين واكتب عنوان الإشعار'}), 400
    conn = db.get_db()
    try:
        emps = [r for r in conn.execute(f"SELECT id, name, dept FROM employees WHERE id IN ({','.join('?' * len(ids))})", ids)
                if _employee_scope_ok(r['dept'])]
        for e in emps:
            conn.execute("INSERT INTO employee_notifications (employee_id, title, body, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
                         (e['id'], title, body, session.get('username', ''), now_local().strftime('%Y-%m-%d %H:%M:%S')))
        _audit(conn, 'notify', 'employee', f'{title} ← {len(emps)} موظف')
        conn.commit()
        return jsonify({'ok': True, 'count': len(emps)})
    finally:
        conn.close()


# =====================================================
#   ملف الموظف (بطاقة الموظف للإدارة)
# =====================================================
RENAME_TABLES = ('punches', 'manual_time_entries', 'date_exceptions', 'late_permissions', 'departure_permissions',
                 'report_cache', 'forced_absences', 'absence_reasons', 'mobile_punches')


@bp.route('/api/mobile/employee/<int:emp_id>')
@admin_required
def api_employee_card(emp_id):
    conn = db.get_db()
    try:
        emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
        if not emp or not _employee_scope_ok(emp['dept']):
            return jsonify({'ok': False, 'error': 'الموظف غير موجود'}), 404
        acc = conn.execute("SELECT username, active, anywhere, device_label, bound_at, last_login_at FROM employee_accounts "
                           "WHERE employee_id = ?", (emp_id,)).fetchone()
        month = now_local().strftime('%Y-%m')
        days = [dict(r) for r in conn.execute(
            "SELECT date, time_in, time_out, status, late_minutes, working_hours FROM report_cache WHERE name = ? "
            "AND date LIKE ? ORDER BY date DESC", (emp['name'], month + '%')).fetchall()]
        moves = [dict(r) for r in conn.execute(
            "SELECT date, time, kind, status, reason, site_name, distance_m, accuracy, source, flags, lat, lng FROM mobile_punches "
            "WHERE employee_id = ? ORDER BY id DESC LIMIT 40", (emp_id,)).fetchall()]
        for m in moves:
            m['kind_label'] = KINDS.get(m['kind'], m['kind'])
        overrides = [dict(r) for r in conn.execute(
            "SELECT * FROM geofence_overrides WHERE employee_id = ? ORDER BY id DESC LIMIT 20", (emp_id,)).fetchall()]
        now = now_local().strftime('%Y-%m-%d %H:%M')
        for o in overrides:
            o['active'] = not o['revoked'] and o['starts_at'] <= now <= o['ends_at']
        notes = [dict(r) for r in conn.execute(
            "SELECT title, body, created_by, created_at, read_at FROM employee_notifications WHERE employee_id = ? ORDER BY id DESC LIMIT 20",
            (emp_id,)).fetchall()]
        sources = {r['src']: r['n'] for r in conn.execute(
            "SELECT CASE WHEN batch_id IN ('mobile', 'kiosk') THEN batch_id ELSE 'device' END AS src, COUNT(*) AS n FROM punches "
            "WHERE name = ? AND date LIKE ? GROUP BY src", (emp['name'], month + '%'))}
        audit = [dict(r) for r in conn.execute(
            "SELECT timestamp, username, action_type, old_value, new_value, reason FROM audit_log WHERE employee_name = ? ORDER BY id DESC LIMIT 60",
            (emp['name'],)).fetchall()]
        return jsonify({'ok': True, 'employee': dict(emp), 'account': dict(acc) if acc else None, 'days': days, 'moves': moves,
                        'overrides': overrides, 'notifications': notes, 'sources': sources, 'audit': audit,
                        'bio': conn.execute("SELECT COUNT(*) FROM webauthn_credentials WHERE employee_id = ?", (emp_id,)).fetchone()[0],
                        'presence': {'monitored': presence_monitored(settings(), emp), 'enabled': _pcfg(settings())['enabled'],
                                     'scope': _pcfg(settings())['scope'],
                                     'events': [dict(r, kind_label=PRESENCE_KINDS.get(r['kind'], r['kind'])) for r in conn.execute(
                                         "SELECT * FROM presence_events WHERE employee_id = ? ORDER BY start_ts DESC LIMIT 60", (emp_id,))]},
                        'summary': {'days': sum(1 for x in days if x['time_in']),
                                    'late_days': sum(1 for x in days if (x['late_minutes'] or 0) > 0),
                                    'late_minutes': sum(int(x['late_minutes'] or 0) for x in days),
                                    'absent': sum(1 for x in days if x['status'] == 'غائب'),
                                    'hours': round(sum(float(x['working_hours'] or 0) for x in days), 1)}})
    finally:
        conn.close()


@bp.route('/api/mobile/employee/<int:emp_id>', methods=['PUT'])
@admin_required
def api_employee_update(emp_id):
    d = request.get_json(silent=True) or {}
    conn = db.get_db()
    try:
        emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
        if not emp or not _employee_scope_ok(emp['dept']):
            return jsonify({'ok': False, 'error': 'الموظف غير موجود'}), 404
        name = (d.get('name') if 'name' in d else emp['name'] or '').strip()
        if not name:
            return jsonify({'ok': False, 'error': 'الاسم مطلوب'}), 400
        if name != emp['name'] and conn.execute("SELECT 1 FROM employees WHERE name = ? AND id != ?", (name, emp_id)).fetchone():
            return jsonify({'ok': False, 'error': 'يوجد موظف آخر بنفس الاسم'}), 400
        mode = d.get('attendance_mode', emp['attendance_mode'] or 'both')
        if mode not in MODE_SOURCES:
            mode = 'both'
        fields = {
            'name': name, 'dept': (d.get('dept', emp['dept']) or '').strip(),
            'device_uid': (d.get('device_uid', emp['device_uid']) or '').strip(),
            'active': 1 if d.get('active', emp['active']) else 0, 'attendance_mode': mode,
            'job_title': (d.get('job_title', emp['job_title']) or '').strip()[:120],
            'phone': (d.get('phone', emp['phone']) or '').strip()[:40],
            'email': (d.get('email', emp['email']) or '').strip()[:120],
            'notes': (d.get('notes', emp['notes']) or '').strip()[:2000],
            'presence_monitor': 1 if d.get('presence_monitor', emp['presence_monitor']) else 0,
        }
        conn.execute("UPDATE employees SET name = ?, dept = ?, device_uid = ?, active = ?, attendance_mode = ?, job_title = ?, "
                     "phone = ?, email = ?, notes = ?, presence_monitor = ?, synced = 0 WHERE id = ?", (*fields.values(), emp_id))
        changed = [k for k in fields if str(fields[k] or '') != str(emp[k] or '')]
        if name != emp['name']:   # تغيير الاسم ينتقل إلى كل السجلات المرتبطة به
            tables = {r['name'] for r in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
            for t in RENAME_TABLES:
                if t in tables and 'name' in {c['name'] for c in conn.execute(f"PRAGMA table_info({t})")}:
                    conn.execute(f"UPDATE {t} SET name = ? WHERE name = ?", (name, emp['name']))
        if fields['dept'] != (emp['dept'] or '') and d.get('history', True):
            conn.execute("UPDATE punches SET dept = ? WHERE name = ?", (fields['dept'], name))
        try:
            _admin().log_audit(conn, 'employee_edit', 'employee', employee_name=name,
                               old_value=', '.join(f'{k}={emp[k]}' for k in changed), new_value=', '.join(f'{k}={fields[k]}' for k in changed))
        except Exception:
            pass
        conn.commit()
        if changed and any(k in changed for k in ('name', 'dept', 'attendance_mode', 'active')):
            schedule_analysis(1)
        return jsonify({'ok': True, 'changed': changed})
    finally:
        conn.close()


@bp.route('/api/mobile/employee/<int:emp_id>/override', methods=['POST'])
@admin_required
def api_override_add(emp_id):
    d = request.get_json(silent=True) or {}
    reason = (d.get('reason') or '').strip()
    if len(reason) < 3:
        return jsonify({'ok': False, 'error': 'اكتب سبب الاستثناء'}), 400
    now = now_local()
    try:
        starts = datetime.strptime((d.get('starts_at') or now.strftime('%Y-%m-%dT%H:%M'))[:16], '%Y-%m-%dT%H:%M')
        ends = datetime.strptime(d['ends_at'][:16], '%Y-%m-%dT%H:%M') if d.get('ends_at') else now + timedelta(hours=float(d.get('hours') or 2))
    except (ValueError, KeyError):
        return jsonify({'ok': False, 'error': 'وقت غير صحيح'}), 400
    if ends <= starts or ends - starts > timedelta(days=31):
        return jsonify({'ok': False, 'error': 'المدة يجب أن تكون بين دقائق و 31 يوماً'}), 400
    conn = db.get_db()
    try:
        emp = conn.execute("SELECT id, name, dept FROM employees WHERE id = ?", (emp_id,)).fetchone()
        if not emp or not _employee_scope_ok(emp['dept']):
            return jsonify({'ok': False, 'error': 'الموظف غير موجود'}), 404
        conn.execute("INSERT INTO geofence_overrides (employee_id, starts_at, ends_at, no_location, reason, created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                     (emp_id, starts.strftime('%Y-%m-%d %H:%M'), ends.strftime('%Y-%m-%d %H:%M'), 1 if d.get('no_location') else 0,
                      reason[:300], session.get('username', ''), now_local().strftime('%Y-%m-%d %H:%M:%S')))
        try:
            _admin().log_audit(conn, 'geofence_override', 'employee', employee_name=emp['name'],
                               new_value=f"{starts:%Y-%m-%d %H:%M} ← {ends:%Y-%m-%d %H:%M}" + (' (بدون موقع)' if d.get('no_location') else ''), reason=reason)
        except Exception:
            pass
        if d.get('notify', True):
            conn.execute("INSERT INTO employee_notifications (employee_id, title, body, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
                         (emp_id, 'سُمح لك بالتسجيل من خارج موقع العمل',
                          f"حتى {ends:%Y-%m-%d %H:%M}" + (' — حتى بدون تحديد الموقع' if d.get('no_location') else '') + f"\nالسبب: {reason}",
                          session.get('username', ''), now_local().strftime('%Y-%m-%d %H:%M:%S')))
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


@bp.route('/api/mobile/override/<int:ov_id>', methods=['DELETE'])
@admin_required
def api_override_revoke(ov_id):
    conn = db.get_db()
    try:
        ov = conn.execute("SELECT o.*, e.name, e.dept FROM geofence_overrides o JOIN employees e ON e.id = o.employee_id WHERE o.id = ?",
                          (ov_id,)).fetchone()
        if not ov or not _employee_scope_ok(ov['dept']):
            return jsonify({'ok': False, 'error': 'غير موجود'}), 404
        conn.execute("UPDATE geofence_overrides SET revoked = 1 WHERE id = ?", (ov_id,))
        try:
            _admin().log_audit(conn, 'geofence_override_revoke', 'employee', employee_name=ov['name'], reason=ov['reason'])
        except Exception:
            pass
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


@bp.route('/api/mobile/employees/bulk', methods=['POST'])
@admin_required
def api_employees_bulk():
    """تعليق/تفعيل الموظفين وطريقة توثيق الدوام لعدة موظفين"""
    d = request.get_json(silent=True) or {}
    ids = [int(i) for i in (d.get('ids') or []) if str(i).isdigit()]
    action, value = d.get('action'), d.get('value')
    if not ids:
        return jsonify({'ok': False, 'error': 'حدّد موظفاً واحداً على الأقل'}), 400
    conn = db.get_db()
    try:
        emps = [r for r in conn.execute(f"SELECT id, name, dept FROM employees WHERE id IN ({','.join('?' * len(ids))})", ids)
                if _employee_scope_ok(r['dept'])]
        for e in emps:
            if action == 'suspend':
                conn.execute("UPDATE employees SET active = 0, synced = 0 WHERE id = ?", (e['id'],))
            elif action == 'activate':
                conn.execute("UPDATE employees SET active = 1, synced = 0 WHERE id = ?", (e['id'],))
            elif action == 'mode' and value in MODE_SOURCES:
                conn.execute("UPDATE employees SET attendance_mode = ? WHERE id = ?", (value, e['id']))
            elif action == 'presence' and value in ('on', 'off'):
                conn.execute("UPDATE employees SET presence_monitor = ? WHERE id = ?", (1 if value == 'on' else 0, e['id']))
            elif action == 'bio' and value in ('', 'off', 'optional', 'required'):
                conn.execute("UPDATE employee_accounts SET bio_mode = ? WHERE employee_id = ?", (value, e['id']))
            else:
                return jsonify({'ok': False, 'error': 'إجراء غير معروف'}), 400
            _audit_emp(conn, f'employees_{action}', 'employee', e['name'], value or '')
        conn.commit()
        schedule_analysis(1)
        return jsonify({'ok': True, 'count': len(emps)})
    finally:
        conn.close()


# =====================================================
#   أجهزة المداخل (QR): تابلت/هاتف ثابت عند الباب يقرأ رمز الموظف المتغيّر
#   - رمز الموظف يُولَّد داخل هاتفه دون إنترنت: HMAC(سرّ الحساب، رقم الموظف:رقم الفترة) ويتغيّر كل 30 ثانية
#   - وقت الحركة = وقت جهاز المدخل المتزامن مع الخادم (لا وقت هاتف الموظف)
#   - جهاز المدخل يحفظ الحركات عند انقطاع الإنترنت ويرسلها لاحقاً
# =====================================================
QR_STEP = 30          # ثانية لكل رمز
QR_WINDOW = 2         # يُقبل الرمز ±فترتين (دقيقة) حول وقت جهاز المدخل
KIOSK_DEDUP_SEC = 60  # مسح متكرر للموظف نفسه خلال دقيقة يُتجاهل


def _qr_code(secret, emp_id, step):
    import hmac
    return hmac.new(secret.encode(), f'{emp_id}:{step}'.encode(), hashlib.sha256).hexdigest()[:10]


def _qr_secret(conn, acc):
    if acc['qr_secret']:
        return acc['qr_secret']
    sec = secrets.token_hex(20)
    conn.execute("UPDATE employee_accounts SET qr_secret = ? WHERE id = ?", (sec, acc['id']))
    conn.commit()
    return sec


@bp.route('/me/api/qr')
@_me_json
def me_qr_info(conn, acc, emp, cfg):
    """سرّ الرمز المتغيّر لهذا الهاتف (يُحفظ في الهاتف ليعمل الرمز دون إنترنت) + وقت الخادم لضبط ساعة التطبيق"""
    if not mode_allows(emp, 'qr'):
        return jsonify({'ok': False, 'error': 'طريقة توثيق دوامك: ' + MODE_TEXT.get(emp['attendance_mode'] or 'both', '') + ' — رمز QR غير متاح لك'}), 403
    return jsonify({'ok': True, 'emp_id': emp['id'], 'name': emp['name'], 'secret': _qr_secret(conn, acc),
                    'step': QR_STEP, 'server_ms': int(time.time() * 1000)})


@bp.route('/me/qr')
def me_qr_page():
    """صفحة رمز الموظف — تعمل دون إنترنت بعد أول فتح (تُخزَّن في الهاتف)"""
    return render_template('me_qr.html', settings=settings())


@bp.route('/me/api/time')
def me_time():
    return jsonify({'ok': True, 'server_ms': int(time.time() * 1000)})


# ---------- إدارة أجهزة المداخل ----------
@bp.route('/api/kiosks', methods=['GET', 'POST'])
@admin_required
def api_kiosks():
    conn = db.get_db()
    try:
        if request.method == 'POST':
            d = request.get_json(silent=True) or {}
            name = (d.get('name') or '').strip()[:80]
            if not name:
                return jsonify({'ok': False, 'error': 'اكتب اسماً للجهاز (مثلاً: المدخل الرئيسي)'}), 400
            site = d.get('site_id') or None
            cur = conn.execute("INSERT INTO qr_kiosks (name, site_id) VALUES (?, ?)", (name, site))
            code = _kiosk_new_pair(conn, cur.lastrowid)
            _audit(conn, 'kiosk_add', 'kiosk', name)
            conn.commit()
            return jsonify({'ok': True, 'id': cur.lastrowid, 'pair_code': code})
        now = now_local().strftime('%Y-%m-%d %H:%M:%S')
        rows = []
        for r in conn.execute("SELECT k.*, s.name AS site_name FROM qr_kiosks k LEFT JOIN work_sites s ON s.id = k.site_id ORDER BY k.id"):
            x = dict(r)
            x.pop('token_hash', None)
            x['paired'] = bool(r['token_hash'])
            x['pair_valid'] = bool(r['pair_code'] and (r['pair_expires'] or '') > now)
            if not x['pair_valid']:
                x['pair_code'] = None
            x['today'] = conn.execute("SELECT COUNT(*) FROM mobile_punches WHERE source = 'kiosk' AND device_label = ? AND date = ?",
                                      (r['name'], now[:10])).fetchone()[0]
            rows.append(x)
        return jsonify({'ok': True, 'rows': rows})
    finally:
        conn.close()


def _kiosk_new_pair(conn, kid):
    code = ''.join(secrets.choice('0123456789') for _ in range(6))
    exp = (now_local() + timedelta(minutes=30)).strftime('%Y-%m-%d %H:%M:%S')
    conn.execute("UPDATE qr_kiosks SET pair_code = ?, pair_expires = ? WHERE id = ?", (code, exp, kid))
    return code


@bp.route('/api/kiosks/bulk', methods=['POST'])
@admin_required
def api_kiosks_bulk():
    d = request.get_json(silent=True) or {}
    ids = [int(i) for i in (d.get('ids') or []) if str(i).isdigit()]
    action = d.get('action')
    if not ids:
        return jsonify({'ok': False, 'error': 'حدّد جهازاً واحداً على الأقل'}), 400
    conn = db.get_db()
    try:
        out = {}
        for kid in ids:
            k = conn.execute("SELECT * FROM qr_kiosks WHERE id = ?", (kid,)).fetchone()
            if not k:
                continue
            if action == 'pair':
                out[kid] = _kiosk_new_pair(conn, kid)
            elif action == 'unpair':
                conn.execute("UPDATE qr_kiosks SET token_hash = NULL WHERE id = ?", (kid,))
            elif action in ('enable', 'disable'):
                conn.execute("UPDATE qr_kiosks SET active = ? WHERE id = ?", (1 if action == 'enable' else 0, kid))
            elif action == 'delete':
                conn.execute("DELETE FROM qr_kiosks WHERE id = ?", (kid,))
            elif action == 'edit':
                conn.execute("UPDATE qr_kiosks SET name = ?, site_id = ? WHERE id = ?",
                             ((d.get('name') or k['name']).strip()[:80], d.get('site_id') or None, kid))
            else:
                return jsonify({'ok': False, 'error': 'إجراء غير معروف'}), 400
            _audit(conn, f'kiosk_{action}', 'kiosk', k['name'])
        conn.commit()
        return jsonify({'ok': True, 'codes': out})
    finally:
        conn.close()


# ---------- جهاز المدخل نفسه ----------
def _kiosk_auth(conn):
    tok = request.headers.get('X-Kiosk', '')
    if len(tok) < 32:
        return None
    k = conn.execute("SELECT * FROM qr_kiosks WHERE token_hash = ?", (hashlib.sha256(tok.encode()).hexdigest(),)).fetchone()
    if not k or not k['active']:
        return None
    conn.execute("UPDATE qr_kiosks SET last_seen = ?, last_ip = ? WHERE id = ?", (now_local().strftime('%Y-%m-%d %H:%M:%S'), _ip(), k['id']))
    return k


@bp.route('/kiosk')
def kiosk_page():
    return render_template('kiosk.html', settings=settings())


@bp.route('/kiosk/api/pair', methods=['POST'])
def kiosk_pair():
    code = ((request.get_json(silent=True) or {}).get('code') or '').strip()
    key = f'{_ip()}|kiosk'
    if _blocked(key):
        return jsonify({'ok': False, 'error': 'محاولات كثيرة — انتظر 15 دقيقة'}), 429
    conn = db.get_db()
    try:
        now = now_local().strftime('%Y-%m-%d %H:%M:%S')
        k = conn.execute("SELECT * FROM qr_kiosks WHERE pair_code = ? AND pair_expires > ?", (code, now)).fetchone() if len(code) == 6 else None
        if not k:
            _FAILS.setdefault(key, []).append(time.time())
            return jsonify({'ok': False, 'error': 'رمز الإقران غير صحيح أو انتهت صلاحيته — أنشئ رمزاً جديداً من لوحة الإدارة'}), 400
        tok = secrets.token_urlsafe(40)
        conn.execute("UPDATE qr_kiosks SET token_hash = ?, pair_code = NULL, pair_expires = NULL WHERE id = ?",
                     (hashlib.sha256(tok.encode()).hexdigest(), k['id']))
        _audit(conn, 'kiosk_paired', 'kiosk', k['name'])
        conn.commit()
        return jsonify({'ok': True, 'token': tok, 'name': k['name']})
    finally:
        conn.close()


@bp.route('/kiosk/api/sync')
def kiosk_sync():
    """وقت الخادم + قائمة الموظفين وأسرار رموزهم (للتحقق دون إنترنت)"""
    conn = db.get_db()
    try:
        k = _kiosk_auth(conn)
        if not k:
            return jsonify({'ok': False, 'error': 'الجهاز غير مقترن أو معطّل', 'unpaired': True}), 401
        roster = []
        for r in conn.execute("SELECT a.*, e.name, e.dept FROM employee_accounts a JOIN employees e ON e.id = a.employee_id "
                              "WHERE a.active = 1 AND e.active = 1"):
            roster.append({'id': r['employee_id'], 'name': r['name'], 'dept': r['dept'] or '', 'secret': _qr_secret(conn, r)})
        conn.commit()
        return jsonify({'ok': True, 'server_ms': int(time.time() * 1000), 'name': k['name'], 'step': QR_STEP,
                        'window': QR_WINDOW, 'dedup': KIOSK_DEDUP_SEC, 'roster': roster, 'org': settings().get('org_name') or ''})
    finally:
        conn.close()


@bp.route('/kiosk/api/scans', methods=['POST'])
def kiosk_scans():
    """يستقبل مسحاً واحداً (مباشر) أو دفعة محفوظة أثناء انقطاع الإنترنت"""
    recs = (request.get_json(silent=True) or {}).get('records') or []
    conn = db.get_db()
    try:
        k = _kiosk_auth(conn)
        if not k:
            return jsonify({'ok': False, 'error': 'الجهاز غير مقترن أو معطّل', 'unpaired': True}), 401
        site = conn.execute("SELECT * FROM work_sites WHERE id = ?", (k['site_id'],)).fetchone() if k['site_id'] else None
        now_ms = time.time() * 1000
        results = []
        touched = set()
        for r in sorted(recs[:500], key=lambda x: float(x.get('ts') or 0)):
            uid = str(r.get('uid') or '')[:64]
            res = {'uid': uid, 'ok': False}
            results.append(res)
            if not uid or conn.execute("SELECT 1 FROM kiosk_scans_seen WHERE uid = ?", (uid,)).fetchone():
                res.update(ok=True, dup=True)
                continue
            try:
                emp_id, step, ts = int(r.get('emp_id')), int(r.get('step')), float(r.get('ts'))
            except (TypeError, ValueError):
                res['error'] = 'بيانات غير صالحة'
                continue
            conn.execute("INSERT INTO kiosk_scans_seen (uid, kiosk_id) VALUES (?, ?)", (uid, k['id']))
            acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (emp_id,)).fetchone()
            emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
            if not acc or not emp or not acc['active'] or not emp['active'] or not acc['qr_secret']:
                res['error'] = 'الموظف غير معروف أو حسابه موقوف'
                continue
            res['name'] = emp['name']
            if not mode_allows(emp, 'qr'):
                res['error'] = 'طريقة توثيق هذا الموظف لا تشمل رمز QR'
                continue
            import hmac as _hm
            if not _hm.compare_digest(_qr_code(acc['qr_secret'], emp_id, step), str(r.get('code') or '')):
                res['error'] = 'رمز غير صحيح'
                continue
            flags = []
            # الوقت: وقت جهاز المدخل المتزامن، ولا يُقبل وقت في المستقبل أو أقدم من 7 أيام
            if ts > now_ms + 120000 or ts < now_ms - 7 * 86400000:
                res['error'] = 'وقت جهاز المدخل غير صحيح — يحتاج اتصالاً بالإنترنت لضبطه'
                continue
            if abs(step - int(ts / 1000 // QR_STEP)) > QR_WINDOW:
                res['error'] = 'الرمز منتهي — اطلب من الموظف فتح الرمز من جديد'
                continue
            if r.get('clock') != 'synced':
                flags.append('ساعة جهاز المدخل لم تُضبط من الخادم قبل هذا المسح')
            if r.get('offline'):
                flags.append('سُجّل أثناء انقطاع إنترنت جهاز المدخل')
            try:
                from zoneinfo import ZoneInfo
                when = datetime.fromtimestamp(ts / 1000, ZoneInfo(os.environ.get('ATT_TZ', 'Asia/Amman'))).replace(tzinfo=None)
            except Exception:
                when = datetime.fromtimestamp(ts / 1000)
            d, t = when.strftime('%Y-%m-%d'), when.strftime('%H:%M:%S')
            last = conn.execute("SELECT time FROM mobile_punches WHERE employee_id = ? AND date = ? AND status = 'accepted' "
                                "ORDER BY time DESC LIMIT 1", (emp_id, d)).fetchone()
            if last and abs((when - datetime.strptime(f"{d} {last['time']}", '%Y-%m-%d %H:%M:%S')).total_seconds()) < KIOSK_DEDUP_SEC:
                res.update(ok=True, dup=True, error='')
                continue
            moves = today_moves(conn, emp_id, d)
            kind = 'in' if not any(m['kind'] == 'in' for m in moves) else 'out'
            conn.execute(
                "INSERT INTO mobile_punches (employee_id, name, dept, kind, date, time, site_id, site_name, status, device_label, ip, flags, source) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'accepted', ?, ?, ?, 'kiosk')",
                (emp_id, emp['name'], emp['dept'] or '', kind, d, t, site['id'] if site else None, site['name'] if site else k['name'],
                 k['name'], _ip(), '؛ '.join(flags) or None))
            _merge_punch(conn, emp, d, t, kind, 'kiosk')
            touched.add(d)
            res.update(ok=True, kind=kind, kind_label=KINDS[kind], time=t[:5], date=d)
        conn.commit()
        if touched:
            schedule_analysis(2)
        return jsonify({'ok': True, 'results': results, 'server_ms': int(time.time() * 1000)})
    finally:
        conn.close()


# =====================================================
#   مراقبة الحضور أثناء الدوام
#   - تطبيق أندرويد يرسل الموقع في الخلفية كل N دقيقة بعد الحضور وحتى الانصراف/نهاية الدوام
#   - فحوص حضور عشوائية (لكل الهواتف): تأكيد الوجود خلال مهلة قصيرة
#   - التنبيهات: خارج النطاق، الموقع مطفأ، انقطاع الإشارة، فحص فائت/خارج النطاق
# =====================================================
PRESENCE_KINDS = {
    'outside': 'خارج نطاق العمل أثناء الدوام', 'no_location': 'الموقع مطفأ أو وهمي أثناء الدوام',
    'no_signal': 'انقطاع إشارة الهاتف أثناء الدوام', 'missed_check': 'لم يردّ على فحص الحضور',
    'failed_check': 'ردّ على فحص الحضور من خارج النطاق',
    'leave_overdue': 'تأخر عن العودة من الخروج المؤقت', 'leave_not_returned': 'لم يعد من الخروج المؤقت',
}


def _pcfg(cfg):
    def num(k, d, lo, hi):
        try:
            return max(lo, min(hi, int(float(cfg.get(k) or d))))
        except (TypeError, ValueError):
            return d
    try:
        depts = set(json.loads(cfg.get('presence_depts_json') or '[]'))
    except ValueError:
        depts = set()
    return {'enabled': cfg.get('presence_enabled', '0') == '1', 'scope': cfg.get('presence_scope', 'selected'), 'depts': depts,
            'interval': num('presence_interval_min', 10, 2, 60), 'grace': num('presence_grace_min', 20, 5, 240),
            'checks': num('presence_checks_per_day', 0, 0, 5), 'window': num('presence_check_window_min', 10, 3, 60)}


def presence_monitored(cfg, emp, acc=None):
    p = _pcfg(cfg)
    if not p['enabled'] or not emp or not emp['active'] or (acc is not None and not acc['active']):
        return False
    if not mode_allows(emp, 'phone'):   # لا يُراقَب من لا يسجّل من الهاتف
        return False
    # الكل، أو موظف محدد، أو موظف في قسم محدد (يشمل من يُضاف للقسم لاحقاً)
    return p['scope'] == 'all' or bool(emp['presence_monitor']) or (emp['dept'] or '') in p['depts']


def _work_end(conn, emp, cfg, now):
    row = conn.execute("SELECT departure_time FROM dept_attendance_settings WHERE dept = ?", (emp['dept'] or '',)).fetchone()
    t = (row['departure_time'] if row and row['departure_time'] else '') or cfg.get('departure_time') or '15:00'
    try:
        h, m = [int(x) for x in t.split(':')[:2]]
    except ValueError:
        h, m = 15, 0
    return now.replace(hour=h, minute=m, second=0, microsecond=0)


def _duty(conn, emp, now):
    """('on' | 'leave' | 'off', وقت الحضور) حسب حركات اليوم (التطبيق وجهاز المدخل)"""
    moves = today_moves(conn, emp['id'], now.strftime('%Y-%m-%d'))
    if not moves:
        # بصم على جهاز البصمة فقط (يصل مباشرة عبر الإنترنت أو بعد السحب): يُعدّ في الدوام حتى نهايته
        r = conn.execute("SELECT MIN(time_in) AS t FROM punches WHERE name = ? AND date = ? AND time_in IS NOT NULL AND time_in != '' "
                         "AND COALESCE(batch_id, '') NOT IN ('mobile', 'kiosk')", (emp['name'], now.strftime('%Y-%m-%d'))).fetchone()
        if r and r['t']:
            return 'on', r['t']
        return 'off', None
    first_in = next((m['time'] for m in moves if m['kind'] == 'in'), None)
    last = moves[-1]['kind']
    return ('off' if last == 'out' else 'leave' if last == 'leave_out' else 'on'), first_in


def presence_state(conn, emp, acc, cfg, now=None):
    now = now or now_local()
    p = _pcfg(cfg)
    st = {'active': False, 'interval_s': p['interval'] * 60, 'end_ms': 0, 'reason': '', 'monitored': presence_monitored(cfg, emp, acc)}
    if not st['monitored']:
        st['reason'] = 'disabled'
        return st
    duty, _ = _duty(conn, emp, now)
    end = _work_end(conn, emp, cfg, now) + timedelta(minutes=15)
    st['end_ms'] = int(time.time() * 1000 + (end - now).total_seconds() * 1000)
    if duty == 'off' or now > end:
        st['reason'] = 'off_duty'
        return st
    st['active'] = True
    st['leave'] = duty == 'leave'
    return st


def _exempt(conn, emp, acc, cfg, now):
    if acc['anywhere'] or cfg.get('mobile_require_geofence', '1') != '1':
        return 'يسجّل من أي مكان'
    if active_override(conn, emp['id']):
        return 'استثناء موقع ساري'
    if _duty(conn, emp, now)[0] == 'leave':
        lv = _active_leave(conn, emp['id'], now.strftime('%Y-%m-%d'))
        if not lv or lv['status'] != 'overdue':      # تجاوز وقت العودة ينهي الإعفاء
            return 'خروج مؤقت مسجّل'
    return ''


def _inside(conn, emp, lat, lng, accuracy):
    best = None
    for s in sites_for(conn, emp['dept']):
        d = haversine_m(lat, lng, s['lat'], s['lng'])
        slack = min(accuracy or 0, 50)
        margin = d - slack - s['radius_m']
        if best is None or margin < best[0]:
            best = (margin, d, s)
    if best is None:
        return None, None, None
    return (1 if best[0] <= 0 else 0), best[1], best[2]


def recompute_presence_today(conn):
    """بعد تعديل مواقع العمل: إعادة حساب «داخل/خارج» لقراءات اليوم بالمواقع الجديدة،
    وحذف تنبيهات «خارج الموقع» التي لم تعد صحيحة (حتى لا يبقى الموظف خارجاً حتى القراءة التالية)"""
    date = now_local().strftime('%Y-%m-%d')
    emps = {}
    changed = 0
    for r in conn.execute("SELECT id, employee_id, lat, lng, accuracy, inside FROM presence_pings "
                          "WHERE date = ? AND status = 'ok' AND lat IS NOT NULL AND lng IS NOT NULL", (date,)).fetchall():
        emp = emps.get(r['employee_id'])
        if emp is None:
            emp = emps[r['employee_id']] = conn.execute("SELECT * FROM employees WHERE id = ?", (r['employee_id'],)).fetchone() or False
        if not emp:
            continue
        inside, dist, _s = _inside(conn, emp, r['lat'], r['lng'], r['accuracy'])
        if inside != r['inside']:
            changed += 1
        conn.execute("UPDATE presence_pings SET inside = ?, distance_m = ? WHERE id = ?",
                     (inside, round(dist, 1) if dist is not None else None, r['id']))
    for ev in conn.execute("SELECT id, employee_id, start_ts, end_ts, open FROM presence_events WHERE date = ? AND kind = 'outside'", (date,)).fetchall():
        still = conn.execute("SELECT COUNT(*) FROM presence_pings WHERE employee_id = ? AND date = ? AND ts BETWEEN ? AND ? "
                             "AND status = 'ok' AND inside = 0", (ev['employee_id'], date, ev['start_ts'], ev['end_ts'])).fetchone()[0]
        if not still:
            conn.execute("DELETE FROM presence_events WHERE id = ?", (ev['id'],))
        elif ev['open']:
            last = conn.execute("SELECT inside FROM presence_pings WHERE employee_id = ? AND date = ? AND status = 'ok' "
                                "ORDER BY ts DESC, id DESC LIMIT 1", (ev['employee_id'], date)).fetchone()
            if last and last['inside'] == 1:
                conn.execute("UPDATE presence_events SET open = 0 WHERE id = ?", (ev['id'],))
    return changed


def _event_upsert(conn, emp, date, kind, start_ts, end_ts, details=''):
    ev = conn.execute("SELECT * FROM presence_events WHERE employee_id = ? AND date = ? AND kind = ? AND open = 1 ORDER BY id DESC LIMIT 1",
                      (emp['id'], date, kind)).fetchone()
    mins = int((datetime.strptime(end_ts, '%Y-%m-%d %H:%M:%S') - datetime.strptime(start_ts, '%Y-%m-%d %H:%M:%S')).total_seconds() // 60)
    if ev:
        conn.execute("UPDATE presence_events SET end_ts = ?, minutes = ?, details = COALESCE(NULLIF(?, ''), details) WHERE id = ?",
                     (end_ts, max(mins, 0), details, ev['id']))
        return False
    conn.execute("INSERT INTO presence_events (employee_id, name, dept, date, kind, start_ts, end_ts, minutes, details, open) "
                 "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 1)", (emp['id'], emp['name'], emp['dept'] or '', date, kind, start_ts, end_ts, max(mins, 0), details))
    return True


_ALERT_TEXT = {
    'outside': ('أنت خارج موقع العمل', 'رصد النظام أنك خارج موقع العمل منذ {m} دقيقة أثناء الدوام. عُد إلى موقع العمل، أو سجّل «خروج مؤقت» إن كان خروجك بإذن.'),
    'no_location': ('خدمة الموقع متوقفة', 'تعذّر التحقق من وجودك في موقع العمل منذ {m} دقيقة لأن خدمة الموقع (GPS) متوقفة أو غير مسموح بها. شغّلها من إعدادات الهاتف.'),
    'no_signal': ('انقطع اتصال هاتفك بالنظام', 'لم تصل قراءات من هاتفك منذ {m} دقيقة أثناء الدوام. افتح تطبيق الحضور وتأكد من اتصال الإنترنت وتشغيل الموقع.'),
}


def _alert_employee(conn, emp, kind, minutes, now):
    """إشعار الموظف عند فتح تنبيه تواجد جديد (داخل التطبيق + إشعار الهاتف إن كان مفعّلاً)"""
    if kind not in _ALERT_TEXT:
        return
    title, body = _ALERT_TEXT[kind]
    body = body.format(m=max(int(minutes or 0), 1))
    conn.execute("INSERT INTO employee_notifications (employee_id, title, body, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
                 (emp['id'], title, body, 'النظام', now.strftime('%Y-%m-%d %H:%M:%S')))
    try:
        push_to_employee(conn, emp['id'], title, body)
    except Exception:   # noqa: BLE001
        pass


def _close_events(conn, emp, date, kinds):
    conn.execute(f"UPDATE presence_events SET open = 0 WHERE employee_id = ? AND date = ? AND open = 1 AND kind IN ({','.join('?' * len(kinds))})",
                 [emp['id'], date] + list(kinds))


def record_ping(conn, emp, acc, cfg, lat, lng, accuracy, status, now=None, source='app'):
    """يحفظ قراءة موقع ويحدّث التنبيهات. يرجع (داخل؟، المسافة)"""
    now = now or now_local()
    p = _pcfg(cfg)
    date, ts = now.strftime('%Y-%m-%d'), now.strftime('%Y-%m-%d %H:%M:%S')
    # انقطاع الإشارة يُقاس بين قراءات الخلفية (أندرويد) فقط؛ صفحة المتصفح لا ترسل إلا وهي مفتوحة
    prev = conn.execute("SELECT ts FROM presence_pings WHERE employee_id = ? AND date = ? AND source = 'app' ORDER BY ts DESC, id DESC LIMIT 1",
                        (emp['id'], date)).fetchone() if source == 'app' else None
    exempt = _exempt(conn, emp, acc, cfg, now)
    inside, dist = None, None
    if status == 'ok' and lat is not None and lng is not None:
        inside, dist, _site = _inside(conn, emp, lat, lng, accuracy)
    if exempt and status == 'ok':
        status = 'exempt'
    conn.execute("INSERT INTO presence_pings (employee_id, date, ts, lat, lng, accuracy, inside, distance_m, status, source) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                 (emp['id'], date, ts, lat, lng, accuracy, inside, round(dist, 1) if dist is not None else None, status, source))
    # 1) انقطاع الإشارة: فجوة طويلة بين قراءتين
    open_ns = conn.execute("SELECT id, start_ts FROM presence_events WHERE employee_id = ? AND date = ? AND kind = 'no_signal' AND open = 1",
                           (emp['id'], date)).fetchone() if source == 'app' else None
    if open_ns:      # عادت الإشارة بعد انقطاع رصده الفحص الدوري: نغلق التنبيه بمدته الفعلية
        mins = int((now - datetime.strptime(open_ns['start_ts'], '%Y-%m-%d %H:%M:%S')).total_seconds() // 60)
        conn.execute("UPDATE presence_events SET end_ts = ?, minutes = ?, open = 0 WHERE id = ?", (ts, mins, open_ns['id']))
    elif prev:
        gap = (now - datetime.strptime(prev['ts'], '%Y-%m-%d %H:%M:%S')).total_seconds() / 60
        if gap >= max(p['grace'], p['interval'] * 3):
            conn.execute("INSERT INTO presence_events (employee_id, name, dept, date, kind, start_ts, end_ts, minutes, details, open) "
                         "VALUES (?, ?, ?, ?, 'no_signal', ?, ?, ?, ?, 0)",
                         (emp['id'], emp['name'], emp['dept'] or '', date, prev['ts'], ts, int(gap), 'لم تصل قراءات من الهاتف (مطفأ، بلا إنترنت، أو أُغلق التطبيق)'))
    # 2) سلسلة القراءات «السيئة» الحالية: خارج النطاق أو بلا موقع
    bad = 'outside' if (status == 'ok' and inside == 0) else 'no_location' if status in ('location_off', 'mock') else None
    if not bad:
        _close_events(conn, emp, date, ('outside', 'no_location'))
        return inside, dist
    rows = conn.execute("SELECT ts, inside, status FROM presence_pings WHERE employee_id = ? AND date = ? ORDER BY ts DESC, id DESC LIMIT 500",
                        (emp['id'], date)).fetchall()
    start = ts
    for r in rows:
        k = 'outside' if (r['status'] == 'ok' and r['inside'] == 0) else 'no_location' if r['status'] in ('location_off', 'mock') else None
        if k != bad:
            break
        start = r['ts']
    if (now - datetime.strptime(start, '%Y-%m-%d %H:%M:%S')).total_seconds() / 60 >= p['grace']:
        det = f'آخر مسافة عن موقع العمل: {int(dist)} م' if (bad == 'outside' and dist is not None) else ('موقع وهمي' if status == 'mock' else 'خدمة الموقع مطفأة')
        if _event_upsert(conn, emp, date, bad, start, ts, det):
            _alert_employee(conn, emp, bad, (now - datetime.strptime(start, '%Y-%m-%d %H:%M:%S')).total_seconds() / 60, now)
    return inside, dist


def _schedule_checks(conn, emp, cfg, now):
    p = _pcfg(cfg)
    date = now.strftime('%Y-%m-%d')
    if p['checks'] <= 0 or conn.execute("SELECT 1 FROM presence_checks WHERE employee_id = ? AND date = ?", (emp['id'], date)).fetchone():
        return
    duty, first_in = _duty(conn, emp, now)
    if duty == 'off' or not first_in:
        return
    start = max(now + timedelta(minutes=10), datetime.strptime(f'{date} {first_in[:8]}', '%Y-%m-%d %H:%M:%S') + timedelta(minutes=30))
    end = _work_end(conn, emp, cfg, now) - timedelta(minutes=p['window'] + 10)
    span = (end - start).total_seconds()
    if span < 600:
        return
    import random
    rnd = random.SystemRandom()
    for off in sorted(rnd.uniform(0, span) for _ in range(p['checks'])):
        due = start + timedelta(seconds=off)
        conn.execute("INSERT INTO presence_checks (employee_id, date, due_at, expires_at) VALUES (?, ?, ?, ?)",
                     (emp['id'], date, due.strftime('%Y-%m-%d %H:%M:%S'), (due + timedelta(minutes=p['window'])).strftime('%Y-%m-%d %H:%M:%S')))


def _expire_checks(conn, now, emp_id=None):
    ts = now.strftime('%Y-%m-%d %H:%M:%S')
    q = "SELECT c.*, e.name, e.dept FROM presence_checks c JOIN employees e ON e.id = c.employee_id WHERE c.status = 'pending' AND c.expires_at < ?"
    params = [ts]
    if emp_id:
        q += " AND c.employee_id = ?"
        params.append(emp_id)
    for c in conn.execute(q, params).fetchall():
        # فحص فات وقته بعد انصراف الموظف لا يُحتسب
        out_time = conn.execute("SELECT MIN(time) FROM mobile_punches WHERE employee_id = ? AND date = ? AND kind = 'out' AND status = 'accepted' "
                                "AND time >= ?", (c['employee_id'], c['date'], c['due_at'][11:19])).fetchone()[0]
        if out_time and out_time <= c['expires_at'][11:19]:
            conn.execute("UPDATE presence_checks SET status = 'cancelled' WHERE id = ?", (c['id'],))
            continue
        conn.execute("UPDATE presence_checks SET status = 'missed' WHERE id = ?", (c['id'],))
        conn.execute("INSERT INTO presence_events (employee_id, name, dept, date, kind, start_ts, end_ts, minutes, details, open) "
                     "VALUES (?, ?, ?, ?, 'missed_check', ?, ?, 0, ?, 0)",
                     (c['employee_id'], c['name'], c['dept'] or '', c['date'], c['due_at'], c['expires_at'], 'لم يؤكد وجوده خلال المهلة'))


def presence_check_for(conn, emp, acc, cfg, now=None):
    """الفحص المطلوب الآن من هذا الموظف (أو None)"""
    now = now or now_local()
    if not presence_monitored(cfg, emp, acc):
        return None
    _schedule_checks(conn, emp, cfg, now)
    _expire_checks(conn, now, emp['id'])
    ts = now.strftime('%Y-%m-%d %H:%M:%S')
    c = conn.execute("SELECT * FROM presence_checks WHERE employee_id = ? AND status = 'pending' AND due_at <= ? AND expires_at >= ? ORDER BY due_at LIMIT 1",
                     (emp['id'], ts, ts)).fetchone()
    conn.commit()
    if not c:
        return None
    left = int((datetime.strptime(c['expires_at'], '%Y-%m-%d %H:%M:%S') - now).total_seconds())
    return {'id': c['id'], 'expires_at': c['expires_at'], 'seconds_left': max(left, 0)}


def _device_notes(conn, emp_id):
    """إشعارات لم تُعرض بعد على الهاتف كإشعارات أندرويد (آخر يومين)، وتُعلَّم كمعروضة"""
    rows = conn.execute("SELECT id, title, body FROM employee_notifications WHERE employee_id = ? AND device_shown_at IS NULL "
                        "AND read_at IS NULL AND created_at >= datetime('now', '-2 days') ORDER BY id LIMIT 10", (emp_id,)).fetchall()
    if rows:
        conn.execute(f"UPDATE employee_notifications SET device_shown_at = datetime('now') WHERE id IN ({','.join('?' * len(rows))})",
                     [r['id'] for r in rows])
    return [{'id': r['id'], 'title': r['title'], 'body': r['body'] or ''} for r in rows]


def _app_signed():
    """تحقق توقيع تطبيق أندرويد (بدون جلسة): يرجع البيانات أو None"""
    import hashlib as _h
    import hmac as _hm
    raw = request.get_data() or b''
    ts = request.headers.get('X-App-Ts', '')
    good = _hm.new(APP_KEY.encode(), raw + b'|' + ts.encode(), _h.sha256).hexdigest()
    try:
        fresh = abs(time.time() * 1000 - float(ts)) < 600000
    except ValueError:
        fresh = False
    if not (fresh and _hm.compare_digest(good, request.headers.get('X-App-Sig', ''))):
        return None
    try:
        return json.loads(raw.decode('utf-8') or '{}')
    except ValueError:
        return None


@bp.route('/me/api/notes/pull', methods=['POST'])
def me_notes_pull():
    """يستدعيها تطبيق أندرويد في الخلفية كل ربع ساعة تقريباً ليعرض الإشعارات الجديدة كإشعارات الهاتف"""
    d = _app_signed()
    if d is None:
        return jsonify({'ok': False, 'error': 'توقيع غير صحيح'}), 403
    conn = db.get_db()
    try:
        acc, emp, _t = _quick_account(conn, (d.get('device') or request.headers.get('X-Device') or '').strip())
        if not acc:
            return jsonify({'ok': False, 'stop': True}), 401
        notes = _device_notes(conn, emp['id'])
        conn.commit()
        return jsonify({'ok': True, 'notes': notes})
    finally:
        conn.close()


@bp.route('/me/api/presence/ping', methods=['POST'])
def me_presence_ping():
    """قراءة موقع من خدمة الخلفية في تطبيق أندرويد (بدون جلسة: رمز الهاتف + توقيع التطبيق)"""
    import hashlib as _h
    import hmac as _hm
    raw = request.get_data() or b''
    ts = request.headers.get('X-App-Ts', '')
    good = _hm.new(APP_KEY.encode(), raw + b'|' + ts.encode(), _h.sha256).hexdigest()
    try:
        fresh = abs(time.time() * 1000 - float(ts)) < 600000
    except ValueError:
        fresh = False
    if not (fresh and _hm.compare_digest(good, request.headers.get('X-App-Sig', ''))):
        return jsonify({'ok': False, 'continue': False, 'error': 'توقيع غير صحيح'}), 403
    d = json.loads(raw.decode('utf-8') or '{}')
    conn = db.get_db()
    try:
        acc, emp, _t = _quick_account(conn, (d.get('device') or request.headers.get('X-Device') or '').strip())
        if not acc:
            return jsonify({'ok': False, 'continue': False, 'error': 'الهاتف غير مرتبط بحساب'}), 401
        cfg = settings()
        now = now_local()
        st = presence_state(conn, emp, acc, cfg, now)
        import shares   # مشاركة الموقع المباشر بمبادرة الموظف تستخدم خدمة الخلفية نفسها
        share = shares.active_share(conn, emp['id'], now)
        if not st['active'] and not share:
            return jsonify({'ok': True, 'continue': False, 'reason': st['reason']})
        lat, lng, accu = d.get('lat'), d.get('lng'), d.get('accuracy')
        chk = None
        if st['active']:
            status = 'mock' if d.get('mock') else 'location_off' if (d.get('location_off') or lat is None) else 'ok'
            record_ping(conn, emp, acc, cfg, lat, lng, accu, status, now)
            chk = presence_check_for(conn, emp, acc, cfg, now)
        interval, remaining = st['interval_s'], (max(0, st['end_ms'] - int(time.time() * 1000)) if st['active'] else 0)
        if share:
            if d.get('mock'):      # موقع وهمي: تُوقف المشاركة وتُعلَّم للإدارة
                shares.stop_for_mock(conn, share)
                share = None
            else:
                shares.add_point(conn, share, lat, lng, accu, now)
        if share:
            interval = shares.share_interval(conn, share)
            share_left = int((datetime.strptime(share['ends_at'], '%Y-%m-%d %H:%M:%S') - now).total_seconds() * 1000)
            remaining = max(remaining, share_left)
        elif not st['active']:
            conn.commit()
            return jsonify({'ok': True, 'continue': False, 'reason': 'mock'})
        notes = _device_notes(conn, emp['id'])
        conn.commit()
        return jsonify({'ok': True, 'continue': True, 'interval_s': interval, 'end_ms': st['end_ms'],
                        'remaining_ms': max(0, remaining), 'check': chk, 'notes': notes, 'share': bool(share)})
    finally:
        conn.close()


@bp.route('/me/api/share/backfill', methods=['POST'])
def me_share_backfill():
    """نقاط مسار حفظها هاتف الموظف أثناء انقطاع الاتصال بالخادم (توقفه أو انقطاع الإنترنت) — تُضاف للمسار بأوقاتها الأصلية"""
    import hashlib as _h
    import hmac as _hm
    raw = request.get_data() or b''
    ts = request.headers.get('X-App-Ts', '')
    good = _hm.new(APP_KEY.encode(), raw + b'|' + ts.encode(), _h.sha256).hexdigest()
    try:
        fresh = abs(time.time() * 1000 - float(ts)) < 600000
    except ValueError:
        fresh = False
    if not (fresh and _hm.compare_digest(good, request.headers.get('X-App-Sig', ''))):
        return jsonify({'ok': False, 'error': 'توقيع غير صحيح'}), 403
    try:
        d = json.loads(raw.decode('utf-8') or '{}')
    except ValueError:
        return jsonify({'ok': False, 'error': 'بيانات غير صالحة'}), 400
    conn = db.get_db()
    try:
        acc, emp, _t = _quick_account(conn, (d.get('device') or request.headers.get('X-Device') or '').strip())
        if not acc:
            return jsonify({'ok': False, 'error': 'الهاتف غير مرتبط بحساب'}), 401
        import shares
        try:
            from zoneinfo import ZoneInfo
            tz = ZoneInfo(os.environ.get('ATT_TZ', 'Asia/Amman'))
        except Exception:
            tz = None
        now_ms = time.time() * 1000
        added, cache = 0, {}
        pts = [x for x in (d.get('points') or [])[:500] if isinstance(x, dict)]
        for x in sorted(pts, key=lambda x: float(x.get('t') or 0)):
            try:
                t, lat, lng = float(x['t']), float(x['lat']), float(x['lng'])
            except (KeyError, TypeError, ValueError):
                continue
            if not (now_ms - 86400000 < t <= now_ms + 60000) or not (-90 <= lat <= 90 and -180 <= lng <= 180):
                continue
            when = datetime.fromtimestamp(t / 1000, tz).replace(tzinfo=None) if tz else datetime.fromtimestamp(t / 1000)
            wts = when.strftime('%Y-%m-%d %H:%M:%S')
            # المشاركة التي كانت سارية في ذلك الوقت
            sh = conn.execute("SELECT * FROM location_shares WHERE employee_id = ? AND kind = 'live' AND started_at <= ? "
                              "AND ends_at >= ? AND (stopped_at IS NULL OR stopped_at >= ?) ORDER BY id DESC LIMIT 1",
                              (emp['id'], wts, wts, wts)).fetchone()
            if not sh:
                continue
            if x.get('mock'):
                shares._flag(conn, sh['id'], 'موقع وهمي (Fake GPS) في نقاط أُرسلت بعد انقطاع الاتصال')
                continue
            if conn.execute("SELECT 1 FROM location_share_points WHERE share_id = ? AND ts = ?", (sh['id'], wts)).fetchone():
                continue
            acc_m = x.get('accuracy')
            conn.execute("INSERT INTO location_share_points (share_id, ts, lat, lng, accuracy) VALUES (?, ?, ?, ?, ?)",
                         (sh['id'], wts, lat, lng, acc_m))
            conn.execute("UPDATE location_shares SET points = points + 1 WHERE id = ?", (sh['id'],))
            # الموقع الحالي يتحدّث فقط إن كانت النقطة أحدث مما لدى الخادم
            conn.execute("UPDATE location_shares SET lat = ?, lng = ?, accuracy = ?, updated_at = ? WHERE id = ? "
                         "AND (updated_at IS NULL OR updated_at < ?)", (lat, lng, acc_m, wts, sh['id'], wts))
            cache[sh['id']] = True
            added += 1
        conn.commit()
        return jsonify({'ok': True, 'added': added})
    finally:
        conn.close()


@bp.route('/me/api/presence/web-ping', methods=['POST'])
@_me_json
def me_presence_web_ping(conn, acc, emp, cfg):
    """قراءة موقع من صفحة التطبيق وهي مفتوحة (آيفون/المتصفح)"""
    st = presence_state(conn, emp, acc, cfg)
    if not st['active']:
        return jsonify({'ok': True, 'active': False})
    d = request.get_json(silent=True) or {}
    lat, lng = d.get('lat'), d.get('lng')
    record_ping(conn, emp, acc, cfg, lat, lng, d.get('accuracy'), 'ok' if lat is not None else 'location_off', source='web')
    conn.commit()
    return jsonify({'ok': True, 'active': True})


@bp.route('/me/api/presence/confirm', methods=['POST'])
@_me_json
def me_presence_confirm(conn, acc, emp, cfg):
    d = request.get_json(silent=True) or {}
    now = now_local()
    chk = presence_check_for(conn, emp, acc, cfg, now)
    if not chk:
        return jsonify({'ok': False, 'error': 'لا يوجد فحص حضور مطلوب الآن'}), 400
    lat, lng, accu = d.get('lat'), d.get('lng'), d.get('accuracy')
    ok, reason, site, dist = check_location(conn, emp, acc, lat, lng, accu, cfg, active_override(conn, emp['id']))
    status = 'confirmed' if ok else 'outside'
    conn.execute("UPDATE presence_checks SET status = ?, responded_at = ?, lat = ?, lng = ?, distance_m = ? WHERE id = ?",
                 (status, now.strftime('%Y-%m-%d %H:%M:%S'), lat, lng, round(dist, 1) if dist is not None else None, chk['id']))
    if not ok:
        conn.execute("INSERT INTO presence_events (employee_id, name, dept, date, kind, start_ts, end_ts, minutes, details, open) "
                     "VALUES (?, ?, ?, ?, 'failed_check', ?, ?, 0, ?, 0)",
                     (emp['id'], emp['name'], emp['dept'] or '', now.strftime('%Y-%m-%d'), now.strftime('%Y-%m-%d %H:%M:%S'),
                      now.strftime('%Y-%m-%d %H:%M:%S'), reason))
    if lat is not None:
        record_ping(conn, emp, acc, cfg, lat, lng, accu, 'ok', now, source='web')
    conn.commit()
    return jsonify({'ok': ok, 'message': 'شكراً — تم تأكيد وجودك' if ok else '', 'error': '' if ok else reason})


# ===== حماية من التنبيهات الكاذبة عند توقف الخادم أو انقطاع إنترنته =====
# إذا لم تصل الهواتفُ الخادمَ لأن الخادم نفسه كان متوقفاً (أو انقطع إنترنت المدرسة) فهذا ليس «انقطاع هاتف»:
# • توقف الخادم: «نبضة» تُسجَّل كل دقيقة من المجدول؛ فجوة ≥ 3 دقائق عند الإقلاع = كان متوقفاً.
# • انقطاع الإنترنت والخادم يعمل: كل الهواتف النشطة (2 أو أكثر) صامتة معاً = الخلل عند الخادم لا عند الهواتف.
# في الحالتين تُمنح الهواتف مهلة كاملة لتتصل من جديد قبل أي تنبيه، وتُمدَّد فحوص الحضور المعلّقة، ويُسجَّل الحدث في سجل التدقيق.
OUTAGE_MIN = 3


def _kv_get(conn, key):
    r = conn.execute("SELECT value FROM app_settings WHERE key = ?", (key,)).fetchone()
    return r['value'] if r else ''


def _kv_set(conn, key, value):
    conn.execute("INSERT INTO app_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value", (key, value))


def _outage_recovered(conn, p, now, since_ts, why):
    """بعد عودة الاتصال: مهلة للهواتف + تمديد الفحوص المعلّقة + سجل تدقيق"""
    hold = max(p['grace'], p['interval'] * 3) + 2
    until = (now + timedelta(minutes=hold)).strftime('%Y-%m-%d %H:%M:%S')
    _kv_set(conn, 'presence_grace_until', until)
    conn.execute("UPDATE presence_checks SET expires_at = ? WHERE status = 'pending' AND expires_at < ?", (until, until))
    text = (f"{why}: من {since_ts[:16]} إلى {now.strftime('%Y-%m-%d %H:%M')} — "
            f"أُوقفت تنبيهات الانقطاع حتى {until[11:16]} ريثما تتصل الهواتف")
    try:   # يُستدعى من المجدول (بلا جلسة مستخدم) فنكتب سجل التدقيق مباشرة باسم «النظام»
        conn.execute("INSERT INTO audit_log (username, action_type, target_type, new_value, timestamp) VALUES (?, ?, ?, ?, ?)",
                     ('النظام', 'server_outage', 'system', text, now.strftime('%Y-%m-%d %H:%M:%S')))
    except Exception:   # noqa: BLE001
        pass


def outage_hold(conn, now=None):
    """هل التنبيهات الزمنية (انقطاع الهاتف، فحص فائت، تأخر العودة) معلّقة الآن بسبب توقف الخادم؟"""
    now_s = (now or now_local()).strftime('%Y-%m-%d %H:%M:%S')
    until = _kv_get(conn, 'presence_grace_until')
    return bool(_kv_get(conn, 'presence_net_silent')) or bool(until and now_s < until)


def server_heartbeat(conn, p, now):
    """يُستدعى كل دقيقة من المجدول فقط: يكشف أن الخادم كان متوقفاً (فجوة في النبضات)"""
    ts = now.strftime('%Y-%m-%d %H:%M:%S')
    last = _kv_get(conn, 'server_heartbeat')
    if last:
        try:
            gap = (now - datetime.strptime(last[:19], '%Y-%m-%d %H:%M:%S')).total_seconds() / 60
        except ValueError:
            gap = 0
        if gap >= OUTAGE_MIN:
            _outage_recovered(conn, p, now, last, 'توقف الخادم')
    _kv_set(conn, 'server_heartbeat', ts)


def presence_sweep(conn=None, tick=False):
    """فحص دوري (كل دقيقة من المجدول، وعند فتح لوحة المتابعة):
    - هاتف أرسل قراءات اليوم ثم انقطع (أُطفئ، أُغلق التطبيق، بلا إنترنت) أثناء الدوام ← تنبيه «انقطاع الإشارة» مفتوح
    - الفحوص العشوائية التي انتهت مهلتها ← «لم يردّ»
    - إنهاء التنبيهات المفتوحة لمن انصرف أو انتهى دوامه"""
    own = conn is None
    conn = conn or db.get_db()
    try:
        cfg = settings()
        p = _pcfg(cfg)
        now = now_local()
        if tick:
            server_heartbeat(conn, p, now)
        hold = outage_hold(conn, now)
        if not hold:
            _expire_checks(conn, now)
        if not p['enabled']:
            conn.commit()
            return
        date, ts = now.strftime('%Y-%m-%d'), now.strftime('%Y-%m-%d %H:%M:%S')
        thr = max(p['grace'], p['interval'] * 3)
        if p['checks'] > 0:
            # جدولة الفحوص العشوائية لكل من سجّل حضوره اليوم (حتى لو كان التطبيق مغلقاً) ثم إرسال إشعار عند حلول موعدها
            for r in conn.execute("SELECT DISTINCT employee_id FROM mobile_punches WHERE date = ? AND status = 'accepted' AND kind = 'in'", (date,)).fetchall():
                emp = conn.execute("SELECT * FROM employees WHERE id = ?", (r['employee_id'],)).fetchone()
                acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (r['employee_id'],)).fetchone()
                if emp and acc and presence_state(conn, emp, acc, cfg, now)['active']:
                    _schedule_checks(conn, emp, cfg, now)
            _notify_due_checks(conn, now)
        last = {r['employee_id']: r['ts'] for r in conn.execute(
            "SELECT employee_id, MAX(ts) AS ts FROM presence_pings WHERE date = ? AND source = 'app' GROUP BY employee_id", (date,))}
        # انقطاع جماعي: كل الهواتف النشطة صامتة معاً ← الخلل في الخادم/الإنترنت لا في الهواتف
        active_gaps = []
        for emp_id, last_ts in last.items():
            emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
            acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (emp_id,)).fetchone()
            if emp and acc and presence_state(conn, emp, acc, cfg, now)['active']:
                active_gaps.append(((now - datetime.strptime(last_ts, '%Y-%m-%d %H:%M:%S')).total_seconds() / 60, last_ts))
        silent = [g for g in active_gaps if g[0] >= max(OUTAGE_MIN * 2, p['interval'] * 2)]
        since = _kv_get(conn, 'presence_net_silent')
        if len(active_gaps) >= 2 and len(silent) == len(active_gaps):
            if not since:
                _kv_set(conn, 'presence_net_silent', max(g[1] for g in active_gaps))
            conn.commit()
            return      # لا تنبيهات أثناء الانقطاع الجماعي
        if since:       # عاد الاتصال: هاتف واحد على الأقل وصل
            _kv_set(conn, 'presence_net_silent', '')
            _outage_recovered(conn, p, now, since, 'انقطع اتصال كل الهواتف بالخادم (إنترنت الخادم أو النفق)')
            hold = True
        if hold:
            conn.commit()
            return      # مهلة بعد عودة الخادم: ننتظر أن تتصل الهواتف قبل الحكم بالانقطاع
        for emp_id, last_ts in last.items():
            emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
            acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (emp_id,)).fetchone()
            if not emp or not acc:
                continue
            st = presence_state(conn, emp, acc, cfg, now)
            if not st['active']:
                # انصرف أو انتهى الدوام: نغلق ما بقي مفتوحاً
                conn.execute("UPDATE presence_events SET open = 0 WHERE employee_id = ? AND date = ? AND open = 1", (emp_id, date))
                continue
            gap = (now - datetime.strptime(last_ts, '%Y-%m-%d %H:%M:%S')).total_seconds() / 60
            if gap >= thr and not _exempt(conn, emp, acc, cfg, now):
                if _event_upsert(conn, emp, date, 'no_signal', last_ts, ts, 'لم تصل قراءات من الهاتف (مطفأ، بلا إنترنت، أو أُغلق التطبيق)'):
                    _alert_employee(conn, emp, 'no_signal', gap, now)
        # في الدوام ولم يصل من هاتفه أي قراءة إطلاقاً اليوم (التطبيق لم يُفتح بعد الحضور، أو بصم على الجهاز فقط)
        cand = {r['employee_id'] for r in conn.execute("SELECT DISTINCT employee_id FROM mobile_punches WHERE date = ? AND status = 'accepted'", (date,))}
        cand |= {r['id'] for r in conn.execute("SELECT DISTINCT e.id FROM employees e JOIN punches p ON p.name = e.name "
                                               "WHERE p.date = ? AND COALESCE(p.batch_id, '') NOT IN ('mobile', 'kiosk')", (date,))}
        for emp_id in cand - set(last):
            emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
            acc = conn.execute("SELECT * FROM employee_accounts WHERE employee_id = ?", (emp_id,)).fetchone()
            if not emp or not acc or not presence_state(conn, emp, acc, cfg, now)['active'] or _exempt(conn, emp, acc, cfg, now):
                continue
            first_in = _duty(conn, emp, now)[1]
            try:
                since = datetime.strptime(f"{date} {str(first_in)[:8]}", '%Y-%m-%d %H:%M:%S')
            except (TypeError, ValueError):
                continue
            gap = (now - since).total_seconds() / 60
            if gap >= max(p['grace'], p['interval'] * 3):
                if _event_upsert(conn, emp, date, 'no_signal', since.strftime('%Y-%m-%d %H:%M:%S'), ts,
                                 'لم تصل أي قراءة من هاتفه منذ الحضور (لم يُفتح التطبيق، أو الموقع غير مسموح)'):
                    _alert_employee(conn, emp, 'no_signal', gap, now)
        conn.commit()
    finally:
        if own:
            conn.close()


# ----- إشعارات المتصفح (Web Push) لفحوص الحضور — آيفون (من الشاشة الرئيسية، iOS 16.4+) وأندرويد -----
def _vapid():
    """مفاتيح VAPID تُنشأ مرة واحدة وتُحفظ في الإعدادات"""
    cfg = settings()
    if cfg.get('vapid_private') and cfg.get('vapid_public'):
        return cfg['vapid_private'], cfg['vapid_public']
    try:
        from py_vapid import Vapid01
        from cryptography.hazmat.primitives import serialization
    except ImportError:
        return None, None
    v = Vapid01()
    v.generate_keys()
    priv = v.private_pem().decode()
    pub = _b64u(v.public_key.public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint))
    db.set_setting('vapid_private', priv)
    db.set_setting('vapid_public', pub)
    return priv, pub


def push_to_employee(conn, emp_id, title, body, url='/me'):
    try:
        from pywebpush import WebPushException, webpush
    except ImportError:
        return 0
    priv, _pub = _vapid()
    if not priv:
        return 0
    sent = 0
    for sub in conn.execute("SELECT * FROM push_subscriptions WHERE employee_id = ?", (emp_id,)).fetchall():
        try:
            webpush({'endpoint': sub['endpoint'], 'keys': {'p256dh': sub['p256dh'], 'auth': sub['auth']}},
                    json.dumps({'title': title, 'body': body, 'url': url}, ensure_ascii=False),
                    vapid_private_key=priv, vapid_claims={'sub': 'mailto:noreply@iajaward.org'}, timeout=10)
            sent += 1
        except WebPushException as e:      # اشتراك منتهٍ: نحذفه
            if getattr(e, 'response', None) is not None and e.response.status_code in (404, 410):
                conn.execute("DELETE FROM push_subscriptions WHERE id = ?", (sub['id'],))
        except Exception:   # noqa: BLE001
            pass
    return sent


def _notify_due_checks(conn, now):
    ts = now.strftime('%Y-%m-%d %H:%M:%S')
    for c in conn.execute("SELECT * FROM presence_checks WHERE status = 'pending' AND notified = 0 AND due_at <= ? AND expires_at >= ?", (ts, ts)).fetchall():
        conn.execute("UPDATE presence_checks SET notified = 1 WHERE id = ?", (c['id'],))
        conn.commit()
        mins = max(1, int((datetime.strptime(c['expires_at'], '%Y-%m-%d %H:%M:%S') - now).total_seconds() // 60))
        push_to_employee(conn, c['employee_id'], 'فحص الحضور', f'افتح التطبيق وأكّد وجودك في موقع العمل خلال {mins} دقيقة')
    conn.commit()


@bp.route('/me/api/push/key')
def me_push_key():
    _priv, pub = _vapid()
    return jsonify({'ok': bool(pub), 'key': pub or ''})


@bp.route('/me/api/push/subscribe', methods=['POST'])
@_me_json
def me_push_subscribe(conn, acc, emp, cfg):
    d = request.get_json(silent=True) or {}
    keys = d.get('keys') or {}
    if not d.get('endpoint') or not keys.get('p256dh') or not keys.get('auth'):
        return jsonify({'ok': False, 'error': 'اشتراك غير صالح'}), 400
    conn.execute("INSERT OR REPLACE INTO push_subscriptions (employee_id, endpoint, p256dh, auth) VALUES (?, ?, ?, ?)",
                 (emp['id'], d['endpoint'][:1000], keys['p256dh'][:200], keys['auth'][:100]))
    conn.commit()
    return jsonify({'ok': True})


# ----- الإدارة -----
def _presence_live(conn, cfg, now):
    """حالة المراقبة لكل موظف اليوم: آخر قراءة، داخل/خارج، انقطاع"""
    p = _pcfg(cfg)
    date = now.strftime('%Y-%m-%d')
    last = {r['employee_id']: dict(r) for r in conn.execute(
        "SELECT * FROM presence_pings WHERE id IN (SELECT MAX(id) FROM presence_pings WHERE date = ? GROUP BY employee_id)", (date,))}
    open_ev = {}
    for r in conn.execute("SELECT employee_id, kind FROM presence_events WHERE date = ? AND open = 1", (date,)):
        open_ev.setdefault(r['employee_id'], []).append(r['kind'])
    counts = {r['employee_id']: r['n'] for r in conn.execute("SELECT employee_id, COUNT(*) AS n FROM presence_events WHERE date = ? GROUP BY employee_id", (date,))}
    out = {}
    for emp_id, pg in last.items():
        mins = (now - datetime.strptime(pg['ts'], '%Y-%m-%d %H:%M:%S')).total_seconds() / 60
        out[emp_id] = {'last_ping': pg['ts'][11:16], 'inside': pg['inside'], 'status': pg['status'], 'distance': pg['distance_m'],
                       'stale': mins >= max(p['grace'], p['interval'] * 3), 'open': open_ev.get(emp_id, []), 'events': counts.get(emp_id, 0)}
    for emp_id, n in counts.items():
        out.setdefault(emp_id, {'events': n, 'open': open_ev.get(emp_id, [])})
    return out


@bp.route('/api/mobile/presence/alerts')
@admin_required
def api_presence_alerts():
    """تنبيهات التواجد المفتوحة الآن (لشريط التنبيه في لوحة التحكم)"""
    conn = db.get_db()
    try:
        presence_sweep(conn)
        date = now_local().strftime('%Y-%m-%d')
        rows = [dict(r) for r in conn.execute("SELECT employee_id, name, dept, kind, start_ts, minutes FROM presence_events "
                                               "WHERE date = ? AND open = 1 ORDER BY start_ts", (date,))]
        rows = [r for r in rows if _employee_scope_ok(r['dept'])]
        by_kind = {}
        for r in rows:
            by_kind[r['kind']] = by_kind.get(r['kind'], 0) + 1
        return jsonify({'ok': True, 'count': len({r['employee_id'] for r in rows}), 'by_kind': by_kind, 'rows': rows[:50]})
    finally:
        conn.close()


@bp.route('/api/mobile/presence/live')
@admin_required
def api_presence_live():
    conn = db.get_db()
    try:
        cfg = settings()
        now = now_local()
        presence_sweep(conn)
        return jsonify({'ok': True, 'enabled': _pcfg(cfg)['enabled'], 'rows': _presence_live(conn, cfg, now)})
    finally:
        conn.close()


def _presence_query(args):
    today = now_local().strftime('%Y-%m-%d')
    d_from, d_to = (args.get('from') or today)[:10], (args.get('to') or today)[:10]
    where, params = ["date >= ?", "date <= ?"], [d_from, d_to]
    if args.get('kind'):
        where.append("kind = ?"); params.append(args['kind'])
    if args.get('emp'):
        where.append("employee_id = ?"); params.append(int(args['emp']))
    return ' AND '.join(where), params, d_from, d_to


@bp.route('/api/mobile/presence/events')
@admin_required
def api_presence_events():
    where, params, d_from, d_to = _presence_query(request.args)
    conn = db.get_db()
    try:
        _expire_checks(conn, now_local())
        conn.commit()
        rows = [dict(r) for r in conn.execute(f"SELECT * FROM presence_events WHERE {where} ORDER BY start_ts DESC, id DESC LIMIT 3000", params)]
        rows = [r for r in rows if _employee_scope_ok(r['dept'])]
        base_w, base_p, _f, _t = _presence_query({k: v for k, v in request.args.items() if k != 'kind'})
        counts = {r['kind']: r['n'] for r in conn.execute(f"SELECT kind, COUNT(*) AS n FROM presence_events WHERE {base_w} GROUP BY kind", base_p)}
        for r in rows:
            r['kind_label'] = PRESENCE_KINDS.get(r['kind'], r['kind'])
        if request.args.get('format') == 'csv':
            import csv
            import io
            L = _admin().L
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow([L(x) for x in ['التاريخ', 'الموظف', 'القسم', 'التنبيه', 'من', 'إلى', 'المدة (دقيقة)', 'التفاصيل']])
            for r in rows:
                w.writerow([r['date'], r['name'], r['dept'], L(r['kind_label']), r['start_ts'][11:16], (r['end_ts'] or '')[11:16], r['minutes'], L(r['details'] or '')])
            return Response('﻿' + buf.getvalue(), content_type='text/csv; charset=utf-8',
                            headers={'Content-Disposition': f'attachment; filename=presence-{d_from}_{d_to}.csv'})
        return jsonify({'ok': True, 'rows': rows, 'counts': counts, 'total': sum(counts.values()), 'kinds': PRESENCE_KINDS})
    finally:
        conn.close()


@bp.route('/api/mobile/presence/events/delete', methods=['POST'])
@admin_required
def api_presence_events_delete():
    ids = [int(i) for i in ((request.get_json(silent=True) or {}).get('ids') or []) if str(i).isdigit()]
    if not ids:
        return jsonify({'ok': False, 'error': 'حدّد تنبيهاً واحداً على الأقل'}), 400
    conn = db.get_db()
    try:
        n = conn.execute(f"DELETE FROM presence_events WHERE id IN ({','.join('?' * len(ids))})", ids).rowcount
        _audit(conn, 'presence_events_delete', 'presence', f'{n}')
        conn.commit()
        return jsonify({'ok': True, 'count': n})
    finally:
        conn.close()
